"""Archivos que van y vienen entre la plataforma y la app del generador."""
import hashlib
import hmac
import io
import json
import re
import unicodedata

from django.conf import settings
from django.db import transaction
from django.utils import timezone

from notas.models import AsignacionDocente, Curso, Docente, Materia, PonderacionAreaMateria

from . import nucleo
from .models import Franja, Horario, Jornada, PlanHoras, Sesion


def _firma(colegio_id, exportado):
    msg = f'{colegio_id}|{exportado}'.encode()
    return hmac.new(settings.SECRET_KEY.encode(), msg, hashlib.sha256).hexdigest()[:32]


def planes_sin_docente(colegio, cursos=None):
    """Las materias con intensidad horaria pero sin docente (PlanHoras sin asignación académica).

    En el archivo de datos van como asignaciones con id NEGATIVO (-id del plan) y sin docente."""
    con_asig = set(AsignacionDocente.objects.filter(colegio=colegio).values_list('curso_id', 'materia_id'))
    qs = PlanHoras.objects.filter(colegio=colegio, horas__gt=0).select_related('materia', 'curso')
    if cursos is not None:
        qs = qs.filter(curso_id__in=cursos)
    return [p for p in qs if (p.curso_id, p.materia_id) not in con_asig]


def clave_sesion(s, planes=None):
    """El id con que el generador conoce la clase: la asignación o, sin docente, -id del plan."""
    if s.asignacion_id:
        return s.asignacion_id
    return (planes or {}).get((s.curso_id, s.materia_id), 0)


def mapa_planes(colegio):
    return {(p.curso_id, p.materia_id): -p.id for p in PlanHoras.objects.filter(colegio=colegio)}


def datos_del_colegio(colegio, con_firma=True, sin_docente=True):
    """El contenido del archivo de datos (.mch), como diccionario.

    sin_docente: incluye también las materias que tienen intensidad horaria pero todavía no
    tienen docente (se arman igual; el docente se pone después)."""
    jornadas = list(Jornada.objects.filter(colegio=colegio).prefetch_related('franjas', 'cursos'))
    cursos_en_jornada = {c.id for j in jornadas for c in j.cursos.all()}
    asignaciones = list(AsignacionDocente.objects.filter(colegio=colegio, curso_id__in=cursos_en_jornada)
                        .select_related('materia', 'docente__user', 'curso'))
    planes = planes_sin_docente(colegio, cursos_en_jornada) if sin_docente else []
    exportado = timezone.now().isoformat(timespec='seconds')
    areas = dict(PonderacionAreaMateria.objects.filter(colegio=colegio).values_list('materia_id', 'area__nombre'))
    d = {
        'tipo': nucleo.TIPO_DATOS, 'version': nucleo.VERSION, 'exportado': exportado,
        'colegio': {'id': colegio.id, 'nombre': colegio.nombre},
        'jornadas': [{'id': j.id, 'nombre': j.nombre, 'dias': sorted(j.dias),
                      'cursos': [c.id for c in j.cursos.all()],
                      'franjas': [{'id': f.id, 'orden': f.orden, 'inicio': f.inicio.strftime('%H:%M'),
                                   'fin': f.fin.strftime('%H:%M'), 'descanso': f.descanso, 'etiqueta': f.etiqueta}
                                  for f in sorted(j.franjas.all(), key=lambda x: (x.orden, x.inicio))]}
                     for j in jornadas],
        'cursos': [{'id': c.id, 'nombre': c.nombre, 'grado': c.grado,
                    'sede': c.sede.nombre if c.sede_id else ''}
                   for c in Curso.objects.filter(id__in=cursos_en_jornada).select_related('sede').order_by('orden')],
        'docentes': [{'id': x.id, 'nombre': x.user.get_full_name() or x.user.username}
                     for x in Docente.objects.filter(colegio=colegio).select_related('user')],
        'materias': [{'id': m.id, 'nombre': m.nombre, 'abreviatura': m.abreviatura or '', 'area': areas.get(m.id, '')}
                     for m in Materia.objects.filter(id__in={a.materia_id for a in asignaciones} |
                                                     {p.materia_id for p in planes}).order_by('nombre')],
        'asignaciones': [{'id': a.id, 'curso': a.curso_id, 'materia': a.materia_id, 'docente': a.docente_id,
                          'ih': a.intensidad_horaria_semanal} for a in asignaciones] +
                        [{'id': -p.id, 'curso': p.curso_id, 'materia': p.materia_id, 'docente': None, 'ih': p.horas}
                         for p in planes],
    }
    d['disponibilidad'] = disponibilidad_para_archivo(colegio, jornadas)
    if con_firma:
        d['firma'] = _firma(colegio.id, exportado)
    return d


def disponibilidad_para_archivo(colegio, jornadas=None):
    """{docente_id: {'no_puede', 'prefiere_no', 'dia_libre', 'nota'}} de lo que marcaron en la plataforma."""
    from .models import Disponibilidad
    jornadas = jornadas if jornadas is not None else list(Jornada.objects.filter(colegio=colegio).prefetch_related('franjas'))
    validas = {f'{d}-{f.id}' for j in jornadas for d in j.dias for f in j.franjas.all() if not f.descanso}
    out = {}
    for x in Disponibilidad.objects.filter(colegio=colegio):
        no = [k for k in x.lista('no') if k in validas]
        pref = [k for k in x.lista('pref') if k in validas]
        if no or pref or x.dia_libre or x.nota:
            out[str(x.docente_id)] = {'no_puede': no, 'prefiere_no': pref, 'dia_libre': x.dia_libre, 'nota': x.nota}
    return out


def archivo_datos(colegio, sin_docente=True):
    return json.dumps(datos_del_colegio(colegio, sin_docente=sin_docente), ensure_ascii=False, indent=1).encode('utf-8')


# ---------------------------------------------------------------------------
# Importar
# ---------------------------------------------------------------------------
def _crear(colegio, nombre, origen, filas, informe):
    """filas: [(id, dia, franja_id)] ya validadas. id > 0 es una asignación; id < 0 es una materia
    sin docente (-id del PlanHoras)."""
    asignaciones = {a.id: a for a in AsignacionDocente.objects.filter(colegio=colegio, id__in={f[0] for f in filas if f[0] > 0})}
    planes = {-p.id: p for p in PlanHoras.objects.filter(colegio=colegio, id__in={-f[0] for f in filas if f[0] < 0})}
    nuevas = []
    for a, d, f in filas:
        if a in asignaciones:
            x = asignaciones[a]
            nuevas.append(Sesion(curso_id=x.curso_id, asignacion_id=a, materia_id=x.materia_id, dia=d, franja_id=f))
        elif a in planes:
            x = planes[a]
            nuevas.append(Sesion(curso_id=x.curso_id, asignacion=None, materia_id=x.materia_id, dia=d, franja_id=f))
    with transaction.atomic():
        h = Horario.objects.create(colegio=colegio, nombre=nombre[:80], origen=origen, informe=informe)
        for x in nuevas:
            x.horario = h
        Sesion.objects.bulk_create(nuevas, ignore_conflicts=True)
    return h


def importar_resultado(archivo, colegio):
    """Lee el .mchr de la app. Devuelve (horario, avisos)."""
    texto = archivo.read().decode('utf-8-sig', errors='replace')
    r = nucleo.leer_json(texto, nucleo.TIPO_RESULTADO)
    origen = r.get('origen', {})
    if origen.get('colegio') != colegio.id:
        raise nucleo.ArchivoInvalido('Este horario se hizo con los datos de otro colegio.')
    if origen.get('firma') != _firma(colegio.id, origen.get('exportado')):
        raise nucleo.ArchivoInvalido('El archivo no salió de los datos exportados por esta plataforma.')
    actual = nucleo.Datos(datos_del_colegio(colegio, con_firma=False))
    sesiones, avisos = [], []
    for s in r.get('sesiones', []):
        try:
            a, d, f = int(s['asignacion']), int(s['dia']), int(s['franja'])
        except (KeyError, TypeError, ValueError):
            continue
        if a not in actual.asignaciones or f not in actual.franjas:
            avisos.append('Algunas clases eran de asignaciones u horas que ya no existen y se dejaron por fuera.')
            continue
        sesiones.append({'asignacion': a, 'curso': actual.asignaciones[a]['curso'], 'dia': d, 'franja': f})
    avisos = sorted(set(avisos))
    informe = dict(r.get('informe') or {})
    informe['choques'] = nucleo.revisar(actual, sesiones)
    h = _crear(colegio, r.get('nombre') or 'Horario generado', 'GENERADO',
               [(s['asignacion'], s['dia'], s['franja']) for s in sesiones], informe)
    return h, avisos


def _norm(t):
    t = unicodedata.normalize('NFKD', str(t or '')).encode('ascii', 'ignore').decode().upper()
    return re.sub(r'[^A-Z0-9]', '', t)


DIA_POR_NOMBRE = {'LUNES': 1, 'LU': 1, 'LUN': 1, 'MARTES': 2, 'MA': 2, 'MAR': 2, 'MIERCOLES': 3, 'MI': 3, 'MIE': 3,
                  'JUEVES': 4, 'JU': 4, 'JUE': 4, 'VIERNES': 5, 'VI': 5, 'VIE': 5, 'SABADO': 6, 'SA': 6, 'SAB': 6}


def importar_excel(archivo, colegio):
    """Excel de otra aplicación: columnas CURSO, DIA, HORA, MATERIA (y DOCENTE opcional).

    HORA puede ser el número de la hora de clase (1, 2, 3…) o la hora de inicio (07:00).
    """
    from openpyxl import load_workbook
    try:
        ws = load_workbook(io.BytesIO(archivo.read()), data_only=True, read_only=True).worksheets[0]
    except Exception:
        raise nucleo.ArchivoInvalido('No se pudo abrir el Excel (guárdelo como .xlsx).')
    filas = [list(f) for f in ws.iter_rows(values_only=True)]
    if not filas:
        raise nucleo.ArchivoInvalido('El Excel está vacío.')
    titulos = [_norm(t) for t in filas[0]]
    col = {}
    for clave, nombres in (('curso', ('CURSO', 'GRUPO')), ('dia', ('DIA',)), ('hora', ('HORA', 'BLOQUE', 'FRANJA', 'PERIODO')),
                           ('materia', ('MATERIA', 'ASIGNATURA', 'AREA')), ('docente', ('DOCENTE', 'PROFESOR'))):
        for i, t in enumerate(titulos):
            if t in nombres:
                col[clave] = i
                break
    faltan = [c for c in ('curso', 'dia', 'hora', 'materia') if c not in col]
    if faltan:
        raise nucleo.ArchivoInvalido('Al Excel le faltan las columnas: ' + ', '.join(f.upper() for f in faltan) + '.')
    cursos = {_norm(c.nombre): c for c in Curso.objects.filter(colegio=colegio)}
    jornada_de = {}
    for j in Jornada.objects.filter(colegio=colegio).prefetch_related('franjas', 'cursos'):
        clase = [f for f in sorted(j.franjas.all(), key=lambda x: (x.orden, x.inicio)) if not f.descanso]
        for c in j.cursos.all():
            jornada_de[c.id] = clase
    materias = {}
    for m in Materia.objects.filter(colegio=colegio):
        materias.setdefault(_norm(m.nombre), m)
        if m.abreviatura:
            materias.setdefault(_norm(m.abreviatura), m)
    planes = {(p.curso_id, p.materia_id): p for p in PlanHoras.objects.filter(colegio=colegio)}
    sin_docente = {}
    asig = {}
    for a in AsignacionDocente.objects.filter(colegio=colegio).select_related('materia', 'docente__user'):
        asig.setdefault((a.curso_id, _norm(a.materia.nombre)), []).append(a)
        if a.materia.abreviatura:
            asig.setdefault((a.curso_id, _norm(a.materia.abreviatura)), []).append(a)
    validas, errores = [], []
    for i, f in enumerate(filas[1:], start=2):
        v = lambda k: f[col[k]] if col.get(k) is not None and col[k] < len(f) else None
        if not any(x not in (None, '') for x in f):
            continue
        curso = cursos.get(_norm(v('curso')))
        if not curso:
            errores.append(f'Fila {i}: no hay un curso «{v("curso")}».')
            continue
        dia_txt = str(v('dia') or '').strip()
        dia = int(dia_txt) if dia_txt.isdigit() else DIA_POR_NOMBRE.get(_norm(dia_txt))
        clase = jornada_de.get(curso.id)
        if not clase:
            errores.append(f'Fila {i}: el curso {curso.nombre} no está en ninguna jornada.')
            continue
        hora = v('hora')
        franja = None
        if hasattr(hora, 'strftime'):
            hora = hora.strftime('%H:%M')
        hora = str(hora or '').strip()
        if hora.replace('.0', '').isdigit():
            n = int(float(hora))
            franja = clase[n - 1] if 1 <= n <= len(clase) else None
        elif re.match(r'^\d{1,2}:\d{2}', hora):
            hh = hora[:5].zfill(5)
            franja = next((x for x in clase if x.inicio.strftime('%H:%M') == hh), None)
        if not dia or not franja:
            errores.append(f'Fila {i}: día u hora no válidos ({v("dia")}, {v("hora")}).')
            continue
        candidatas = asig.get((curso.id, _norm(v('materia'))), [])
        if len(candidatas) > 1 and v('docente'):
            candidatas = [a for a in candidatas if _norm(a.docente.user.get_full_name()) == _norm(v('docente'))] or candidatas
        if not candidatas:
            # Sin docente: se arma igual con la materia (queda pendiente asignarle docente)
            m = materias.get(_norm(v('materia')))
            if not m:
                errores.append(f'Fila {i}: no hay una materia «{v("materia")}» en el colegio.')
                continue
            sin_docente.setdefault((curso.id, m.id), []).append((dia, franja.id))
            continue
        validas.append((candidatas[0].id, dia, franja.id))
    for (cid, mid), horas in sin_docente.items():
        plan = planes.get((cid, mid))
        if plan is None:
            plan = PlanHoras.objects.create(colegio=colegio, curso_id=cid, materia_id=mid, horas=len(horas))
        elif plan.horas < len(horas):
            plan.horas = len(horas)
            plan.save(update_fields=['horas'])
        validas += [(-plan.id, d, f) for d, f in horas]
    if sin_docente:
        errores.insert(0, f'{len(sin_docente)} materia(s) entraron sin docente: asígnelo en la asignación académica '
                          'y la clase se amarra sola.')
    if not validas:
        raise nucleo.ArchivoInvalido('No se pudo leer ninguna clase. ' + (errores[0] if errores else ''))
    actual = nucleo.Datos(datos_del_colegio(colegio, con_firma=False))
    sesiones = [{'asignacion': a, 'dia': d, 'franja': f} for a, d, f in validas]
    informe = {'choques': nucleo.revisar(actual, sesiones), 'errores_importacion': errores[:50]}
    h = _crear(colegio, f'Importado {timezone.localdate():%d/%m/%Y}', 'IMPORTADO', validas, informe)
    return h, errores
