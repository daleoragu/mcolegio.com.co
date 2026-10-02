# pruebas/exports.py
"""Genera el informe en Excel con las mismas hojas que producía el HTML."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .utils import hhmm

VINO = 'FF460F10'
MORADO = 'FF9E1982'
BLANCO = 'FFFFFFFF'


def _encabezado(hoja, titulos, color=VINO):
    hoja.append(titulos)
    relleno = PatternFill('solid', fgColor=color)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color=BLANCO, size=10)
        celda.fill = relleno
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    hoja.freeze_panes = 'A2'


def _anchos(hoja, anchos):
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho


def construir_libro(origen, sesiones, resumen):
    libro = Workbook()

    # --- Hoja 1: resultados por estudiante ---
    hoja = libro.active
    hoja.title = 'Resultados'
    _encabezado(hoja, [
        'Apellidos', 'Nombres', 'Grupo', 'Sede', 'Estudiante en plataforma', 'Curso',
        'Fecha', 'Aciertos', 'Total', '% acierto', 'Franja', 'Tiempo', 'Cambios', 'Intento n.º',
    ])
    for s in sesiones:
        hoja.append([
            s.apellidos, s.nombres, s.grupo, s.sede,
            str(s.estudiante) if s.estudiante else 'Sin vincular',
            str(s.curso) if s.curso else '',
            s.fecha.strftime('%Y-%m-%d %H:%M'),
            s.aciertos, s.total, s.porcentaje, s.franja[0],
            hhmm(s.segundos), s.cambios, s.vez,
        ])
    _anchos(hoja, [22, 22, 10, 18, 26, 12, 17, 9, 8, 10, 11, 9, 9, 11])

    # --- Hoja 2: análisis por pregunta ---
    hoja2 = libro.create_sheet('Por pregunta')
    _encabezado(hoja2, [
        'N.º', 'Competencia', 'Aprendizaje evaluado', 'Clave',
        'Aciertos', 'Presentaron', '% acierto', 'Distractor más marcado', 'Veces',
    ], MORADO)
    for p in resumen['preguntas']:
        hoja2.append([
            p['numero'], p['competencia'], p['aprendizaje'], p['clave'],
            p['ok'], p['n'], p['pct'], p['distractor'], p['distractor_n'],
        ])
    _anchos(hoja2, [8, 13, 62, 8, 10, 12, 11, 24, 8])

    # --- Hoja 3: competencias ---
    hoja3 = libro.create_sheet('Competencias')
    _encabezado(hoja3, ['Competencia', 'Respuestas correctas', 'Respuestas totales', '% acierto'], MORADO)
    for c in resumen['competencias']:
        hoja3.append([c['sigla'], c['ok'], c['n'], c['pct']])
    _anchos(hoja3, [18, 22, 20, 12])

    # --- Hoja 4: respuesta por respuesta ---
    hoja4 = libro.create_sheet('Detalle')
    _encabezado(hoja4, [
        'Apellidos', 'Nombres', 'Grupo', 'Pregunta', 'Competencia',
        'Marcó', 'Clave', '¿Acertó?', 'Segundos', 'Cambios de opción',
    ], MORADO)
    for s in sesiones:
        for d in s.detalles.all():
            hoja4.append([
                s.apellidos, s.nombres, s.grupo, d.numero, d.competencia,
                d.marcada, d.clave, 'Sí' if d.acierto else 'No', d.segundos, d.intentos,
            ])
    _anchos(hoja4, [22, 22, 10, 10, 13, 9, 9, 10, 10, 18])

    # --- Hoja 5: resumen ---
    hoja5 = libro.create_sheet('Resumen')
    hoja5.append(['Prueba', getattr(origen, 'titulo', '')])
    hoja5.append(['Presentaciones', resumen['n']])
    hoja5.append(['Promedio del grupo (%)', resumen['promedio']])
    hoja5.append(['Mejor resultado (%)', resumen['mejor']])
    hoja5.append(['Resultado más bajo (%)', resumen['peor']])
    hoja5.append(['Tiempo promedio', hhmm(resumen['tiempo_promedio'])])
    hoja5.append([])
    hoja5.append(['Franja', 'Estudiantes', '%'])
    for f in resumen['franjas']:
        hoja5.append([f['nombre'], f['n'], f['pct']])
    for celda in ('A1', 'A2', 'A3', 'A4', 'A5', 'A6', 'A8'):
        hoja5[celda].font = Font(bold=True)
    _anchos(hoja5, [26, 18, 10])

    memoria = BytesIO()
    libro.save(memoria)
    return memoria.getvalue()
