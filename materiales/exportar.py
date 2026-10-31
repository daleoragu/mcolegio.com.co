"""PDF y Word de un documento de Mis materiales, con el encabezado del colegio o del docente."""
import io
import re
import tempfile
from contextlib import contextmanager
from html.parser import HTMLParser
from pathlib import Path

from django.template.loader import render_to_string
from django.utils import timezone

from . import bloques as B
from .models import Imagen

MODOS = {
    # letra (diferencia), margen en mm, interlineado, columnas forzadas, alto máximo de imagen (mm)
    'normal': {'dif': 0, 'margen': 16, 'linea': 1.4, 'columnas': None, 'img': 120, 'sep': 10},
    'compacto': {'dif': -0.5, 'margen': 11, 'linea': 1.25, 'columnas': None, 'img': 85, 'sep': 6},
    'maximo': {'dif': -1, 'margen': 8, 'linea': 1.15, 'columnas': 2, 'img': 60, 'sep': 4},
}


def _leer(archivo):
    if not archivo:
        return None
    try:
        with archivo.open('rb') as f:
            return f.read()
    except Exception:
        return None


@contextmanager
def imagenes_en_disco(doc, encabezado=None):
    """Escribe en una carpeta temporal las imágenes del documento (y el logo del docente) para que
    WeasyPrint las lea del disco y no tenga que pedirlas por internet. Devuelve {id: ruta}."""
    usadas = B.imagenes_usadas(doc.contenido)
    with tempfile.TemporaryDirectory(prefix='material-') as carpeta:
        rutas = {}
        for img in Imagen.objects.filter(pk__in=usadas, propietario=doc.propietario):
            datos = _leer(img.archivo)
            if datos:
                ruta = Path(carpeta) / f'{img.pk}{Path(img.archivo.name).suffix or ".jpg"}'
                ruta.write_bytes(datos)
                rutas[img.pk] = ruta
        if encabezado is not None and getattr(encabezado, 'logo', None):
            datos = _leer(encabezado.logo)
            if datos:
                ruta = Path(carpeta) / f'logo{Path(encabezado.logo.name).suffix or ".png"}'
                ruta.write_bytes(datos)
                rutas['logo'] = ruta
        yield rutas


def _numerar(bloques):
    n = 0
    for b in bloques:
        if b['tipo'] == 'pregunta':
            n += 1
            b['numero'] = n
    return bloques


def contexto(doc, modo=None, docente=''):
    contenido = B.limpiar(doc.contenido, B.imagenes_usadas(doc.contenido))
    opciones = contenido['opciones']
    modo = modo if modo in MODOS else opciones['modo']
    m = MODOS[modo]
    bloques = _numerar([dict(b) for b in contenido['bloques']])
    for b in bloques:
        if b['tipo'] == 'pregunta':
            b['disposicion'] = B.disposicion_opciones(b['opciones'], modo)
            b['letras'] = list(zip('ABCDEF', b['opciones']))
    return {
        'doc': doc, 'bloques': bloques, 'opciones': opciones, 'modo': modo,
        'letra': max(8.5, opciones['letra'] + m['dif']), 'margen': m['margen'], 'interlineado': m['linea'],
        'columnas': m['columnas'] or opciones['columnas'], 'alto_imagen': m['img'], 'separacion': m['sep'],
        'docente': docente, 'fecha': timezone.localdate(),
    }


def _html(request, doc, modo, docente, colegio, encabezado, rutas):
    ctx = contexto(doc, modo, docente)
    ctx.update({'colegio': colegio, 'encabezado': encabezado,
                'rutas': {k: v.as_uri() for k, v in rutas.items()}})
    return render_to_string('materiales/documento_pdf.html', ctx, request=request)


def pdf(request, doc, colegio=None, encabezado=None, docente='', modo=None):
    from weasyprint import HTML
    with imagenes_en_disco(doc, encabezado) as rutas:
        html = _html(request, doc, modo, docente, colegio, encabezado, rutas)
        return HTML(string=html, base_url=request.build_absolute_uri('/')).write_pdf()


def paginas(request, doc, colegio=None, encabezado=None, docente=''):
    """Cuántas páginas sale el documento en cada modo, para mostrar cuánto papel se ahorra."""
    from weasyprint import HTML
    salida = {}
    with imagenes_en_disco(doc, encabezado) as rutas:
        for modo in MODOS:
            html = _html(request, doc, modo, docente, colegio, encabezado, rutas)
            salida[modo] = len(HTML(string=html, base_url=request.build_absolute_uri('/')).render().pages)
    return salida


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------
class _Trozos(HTMLParser):
    """Convierte el HTML limpio en párrafos de trozos con su formato: [[(texto, {'b','i'...}), ...], ...]."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parrafos = [[]]
        self.formato = []
        self.vinieta = []

    def _nuevo(self):
        if self.parrafos[-1]:
            self.parrafos.append([])

    def handle_starttag(self, tag, attrs):
        if tag in ('p', 'br'):
            self._nuevo()
        elif tag == 'li':
            self._nuevo()
            self.parrafos[-1].append(('• ', frozenset()))
        elif tag in ('b', 'i', 'u', 'sub', 'sup'):
            self.formato.append(tag)

    def handle_endtag(self, tag):
        if tag in ('p', 'li'):
            self._nuevo()
        elif tag in self.formato:
            self.formato.remove(tag)

    def handle_data(self, data):
        if data:
            self.parrafos[-1].append((data, frozenset(self.formato)))


def _parrafos(html):
    p = _Trozos()
    p.feed(html or '')
    return [x for x in p.parrafos if any(t.strip() for t, _ in x)]


def _escribir(parrafo, trozos, tam=None):
    from docx.shared import Pt
    for texto, f in trozos:
        r = parrafo.add_run(texto)
        r.bold = 'b' in f or None
        r.italic = 'i' in f or None
        r.underline = 'u' in f or None
        if 'sub' in f:
            r.font.subscript = True
        if 'sup' in f:
            r.font.superscript = True
        if tam:
            r.font.size = Pt(tam)


def _sombrear(celda, color):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc = celda._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tc.append(shd)


def docx(doc, colegio=None, encabezado=None, docente='', modo=None):
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    ctx = contexto(doc, modo, docente)
    letra = ctx['letra']
    w = Document()
    estilo = w.styles['Normal']
    estilo.font.name = 'Arial'
    estilo.font.size = Pt(letra)
    estilo.paragraph_format.space_after = Pt(2 if ctx['modo'] != 'normal' else 4)
    sec = w.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    margen = Cm(ctx['margen'] / 10)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = margen
    util = sec.page_width - 2 * margen

    # Encabezado
    if colegio is not None:
        from puntoexacto.cuadernillo import encabezado_docx
        encabezado_docx(w, colegio, util)
    elif encabezado is not None and (encabezado.lineas or encabezado.logo):
        t = w.add_table(rows=1, cols=2)
        logo = _leer(encabezado.logo) if encabezado.logo else None
        if logo:
            try:
                t.rows[0].cells[0].paragraphs[0].add_run().add_picture(io.BytesIO(logo), height=Cm(1.8))
            except Exception:
                pass
        for i, linea in enumerate(encabezado.lineas):
            p = t.rows[0].cells[1].paragraphs[0] if i == 0 else t.rows[0].cells[1].add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(linea)
            r.bold = i == 0
            r.font.size = Pt(11 if i == 0 else 9)

    titulo = w.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = titulo.add_run(doc.titulo)
    r.bold = True
    r.font.size = Pt(letra + 3)
    if ctx['opciones']['datos']:
        datos = [x for x in (doc.area, doc.grado and f'Grado {doc.grado}', doc.periodo, docente and f'Docente: {docente}') if x]
        if datos:
            p = w.add_paragraph(' · '.join(datos))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if ctx['opciones']['nombre_estudiante']:
        w.add_paragraph('Nombre: ______________________________________   Curso: ________   Fecha: ____________')

    if ctx['columnas'] == 2:
        nueva = w.add_section(WD_SECTION.CONTINUOUS)
        cols = nueva._sectPr.xpath('./w:cols')
        col = cols[0] if cols else OxmlElement('w:cols')
        col.set(qn('w:num'), '2')
        col.set(qn('w:space'), '400')
        if not cols:
            nueva._sectPr.append(col)

    with imagenes_en_disco(doc) as rutas:
        for b in ctx['bloques']:
            t = b['tipo']
            if t == 'titulo':
                p = w.add_paragraph()
                r = p.add_run(b['texto'])
                r.bold = True
                r.font.size = Pt(letra + (3 if b['nivel'] == 1 else 1.5))
            elif t == 'seccion':
                p = w.add_paragraph()
                r = p.add_run(b['titulo'])
                r.bold = True
                r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x68)
                r.font.size = Pt(letra + 1.5)
                if b['indicacion']:
                    q = w.add_paragraph()
                    rr = q.add_run(b['indicacion'])
                    rr.italic = True
            elif t == 'texto':
                for trozos in _parrafos(b['html']):
                    _escribir(w.add_paragraph(), trozos)
            elif t == 'imagen' and b['imagen'] in rutas:
                p = w.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                try:
                    ancho = util * b['ancho'] / 100 / (2 if ctx['columnas'] == 2 else 1)
                    p.add_run().add_picture(str(rutas[b['imagen']]), width=int(ancho))
                except Exception:
                    pass
                if b['pie']:
                    q = w.add_paragraph(b['pie'])
                    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif t == 'tabla':
                tabla = w.add_table(rows=len(b['filas']), cols=len(b['filas'][0]))
                tabla.style = 'Table Grid'
                for i, fila in enumerate(b['filas']):
                    for j, valor in enumerate(fila):
                        celda = tabla.cell(i, j)
                        celda.text = valor
                        if i == 0 and b['encabezado']:
                            for r in celda.paragraphs[0].runs:
                                r.bold = True
                            _sombrear(celda, 'E8EEF9')
                w.add_paragraph()
            elif t == 'recuadro':
                tabla = w.add_table(rows=1, cols=1)
                tabla.style = 'Table Grid'
                celda = tabla.cell(0, 0)
                _sombrear(celda, {'recuerda': 'FFF6DF', 'reto': 'FDECEE', 'vocabulario': 'E6F4EA'}.get(b['estilo'], 'EEF2F8'))
                r = celda.paragraphs[0].add_run(b['titulo'])
                r.bold = True
                for trozos in _parrafos(b['html']):
                    _escribir(celda.add_paragraph(), trozos)
                w.add_paragraph()
            elif t == 'lineas':
                for _ in range(b['n']):
                    w.add_paragraph('_' * (45 if ctx['columnas'] == 2 else 90))
            elif t == 'pregunta':
                p = w.add_paragraph()
                r = p.add_run(f'{b["numero"]}. ')
                r.bold = True
                p.add_run(b['enunciado'])
                if b.get('imagen') in rutas:
                    try:
                        w.add_paragraph().add_run().add_picture(str(rutas[b['imagen']]), width=int(util * 0.4))
                    except Exception:
                        pass
                if b['disposicion'] == 'fila':
                    w.add_paragraph('     '.join(f'{l}. {o}' for l, o in b['letras']))
                elif b['disposicion'] == 'dos':
                    pares = b['letras']
                    for k in range(0, len(pares), 2):
                        w.add_paragraph('\t'.join(f'{l}. {o}' for l, o in pares[k:k + 2]))
                else:
                    for l, o in b['letras']:
                        q = w.add_paragraph(f'{l}. {o}')
                        q.paragraph_format.left_indent = Cm(0.6)
                for _ in range(b['lineas']):
                    w.add_paragraph('_' * (45 if ctx['columnas'] == 2 else 90))
            elif t == 'separador':
                w.add_paragraph('—' * 30).alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif t == 'salto':
                w.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    salida = io.BytesIO()
    w.save(salida)
    return salida.getvalue()


def nombre_archivo(doc, ext):
    base = re.sub(r'[^\w\s-]', '', doc.titulo).strip()[:60] or 'material'
    return f'{base}.{ext}'
