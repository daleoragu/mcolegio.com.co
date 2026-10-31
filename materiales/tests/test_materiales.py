"""Mis materiales: crear desde plantillas, guardar, imágenes, PDF, Word, páginas y publicar en el aula."""
import io
import json
import shutil
import tempfile

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, SimpleTestCase, override_settings

from aula.models import Aula, Material
from materiales import bloques as B
from materiales.models import Documento, EncabezadoDocente, Imagen
from notas.tests.base import CLAVE, ColegioDePrueba

MEDIA = tempfile.mkdtemp(prefix='materiales-')


def png(nombre='dibujo.png', ancho=2400, alto=1200):
    from PIL import Image
    salida = io.BytesIO()
    Image.new('RGB', (ancho, alto), (200, 220, 240)).save(salida, format='PNG')
    return SimpleUploadedFile(nombre, salida.getvalue(), content_type='image/png')


class Bloques(SimpleTestCase):

    def test_limpia_y_acota(self):
        sucio = {'bloques': [
            {'tipo': 'texto', 'html': '<p onclick="x()">Hola <b>mundo</b><script>alert(1)</script><a href="http://x">enlace</a></p>'},
            {'tipo': 'virus', 'html': 'x'},
            {'tipo': 'imagen', 'imagen': 99, 'ancho': 500},
            {'tipo': 'imagen', 'imagen': 7, 'ancho': 500},
            {'tipo': 'pregunta', 'enunciado': '¿2+2?', 'opciones': ['3', '4', '', '5'], 'correcta': 1},
            {'tipo': 'lineas', 'n': 999},
        ], 'opciones': {'columnas': 2, 'modo': 'raro', 'letra': 40}}
        limpio = B.limpiar(sucio, {7})
        tipos = [b['tipo'] for b in limpio['bloques']]
        self.assertEqual(tipos, ['texto', 'imagen', 'pregunta', 'lineas'])
        texto = limpio['bloques'][0]['html']
        self.assertNotIn('script', texto)
        self.assertNotIn('onclick', texto)
        self.assertNotIn('href', texto)
        self.assertIn('<b>mundo</b>', texto)
        self.assertEqual(limpio['bloques'][1]['ancho'], 100)
        self.assertEqual(limpio['bloques'][2]['opciones'], ['3', '4', '5'])
        self.assertEqual(limpio['bloques'][3]['n'], 30)
        self.assertEqual(limpio['opciones'], {'columnas': 2, 'letra': 14, 'modo': 'normal',
                                              'nombre_estudiante': True, 'datos': True})

    def test_opciones_cortas_en_una_linea(self):
        self.assertEqual(B.disposicion_opciones(['3', '2', '4', '1'], 'normal'), 'fila')
        self.assertEqual(B.disposicion_opciones(['Triángulo 1', 'Triángulo 2'], 'normal'), 'lista')
        self.assertEqual(B.disposicion_opciones(['Triángulo 1', 'Triángulo 2'], 'compacto'), 'fila')
        self.assertEqual(B.disposicion_opciones(['Porque hacía mucho sol en la mañana'] * 4, 'compacto'), 'dos')
        self.assertEqual(B.disposicion_opciones(['x' * 80] * 4, 'maximo'), 'lista')


@override_settings(MEDIA_ROOT=MEDIA)
class EnElColegio(ColegioDePrueba):

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.a.modulos_activos = list(cls.a.modulos_activos) + ['salon']
        cls.a.save(update_fields=['modulos_activos'])

    def setUp(self):
        self.profe = self.cliente(self.u_docente)

    def crear(self, tipo='guia_momentos', **extra):
        r = self.profe.post('/materiales/nuevo/', dict(tipo=tipo, titulo='Fracciones', **extra))
        doc = Documento.objects.latest('pk')
        self.assertRedirects(r, f'/materiales/{doc.pk}/', fetch_redirect_response=False)
        return doc

    def test_crear_guardar_imagen_y_exportar(self):
        self.assertContains(self.profe.get('/materiales/'), 'Guía por momentos')
        doc = self.crear()
        self.assertEqual(doc.colegio, self.a)
        secciones = [b['titulo'] for b in doc.contenido['bloques'] if b['tipo'] == 'seccion']
        self.assertEqual(secciones[0], '1. Exploración')
        self.assertEqual(self.profe.get(f'/materiales/{doc.pk}/').status_code, 200)

        # Imagen pegada: se reduce a 1600 px y cuenta en la nube
        r = self.profe.post('/materiales/imagen/', {'imagen': png()})
        img = Imagen.objects.get()
        self.assertTrue(r.json()['ok'])
        self.assertEqual((img.ancho, img.alto), (1600, 800))
        self.assertGreater(img.peso, 0)
        self.assertEqual(self.profe.post('/materiales/imagen/', {'imagen': SimpleUploadedFile('a.png', b'no')}).status_code, 400)
        # La imagen de otro docente no se puede usar
        otro = User.objects.create_user('otro_profe', password=CLAVE)
        ajena = Imagen.objects.create(propietario=otro, archivo='x.jpg', peso=10)

        contenido = {'bloques': [
            {'id': 'a', 'tipo': 'titulo', 'texto': 'Fracciones equivalentes', 'nivel': 1},
            {'id': 'b', 'tipo': 'texto', 'html': '<p>Una <b>fracción</b> es parte de un todo.</p>'},
            {'id': 'c', 'tipo': 'imagen', 'imagen': img.pk, 'ancho': 50, 'pie': 'Una pizza'},
            {'id': 'd', 'tipo': 'imagen', 'imagen': ajena.pk},
            {'id': 'e', 'tipo': 'pregunta', 'enunciado': '¿Cuál es equivalente a 1/2?', 'opciones': ['2/4', '1/3', '3/4', '2/3'], 'correcta': 0},
            {'id': 'f', 'tipo': 'tabla', 'filas': [['Fracción', 'Decimal'], ['1/2', '0,5']], 'encabezado': True},
            {'id': 'g', 'tipo': 'recuadro', 'estilo': 'recuerda', 'titulo': 'Recuerda', 'html': '<p>Multiplica arriba y abajo.</p>'},
            {'id': 'h', 'tipo': 'lineas', 'n': 3},
        ], 'opciones': {'columnas': 2, 'modo': 'compacto', 'letra': 11}}
        r = self.profe.post(f'/materiales/{doc.pk}/guardar/', json.dumps({
            'titulo': 'Fracciones 4.º', 'area': 'Matemáticas', 'grado': '4.º', 'periodo': 'Periodo 3', 'contenido': contenido}),
            content_type='application/json')
        self.assertTrue(r.json()['ok'])
        doc.refresh_from_db()
        self.assertEqual(doc.titulo, 'Fracciones 4.º')
        self.assertEqual([b['id'] for b in doc.contenido['bloques']], ['a', 'b', 'c', 'e', 'f', 'g', 'h'])

        r = self.profe.get(f'/materiales/{doc.pk}/pdf/')
        self.assertEqual(r['Content-Type'], 'application/pdf')
        self.assertTrue(r.content.startswith(b'%PDF'))
        r = self.profe.get(f'/materiales/{doc.pk}/paginas/')
        self.assertEqual(set(r.json()), {'normal', 'compacto', 'maximo'})
        self.assertGreaterEqual(r.json()['normal'], r.json()['maximo'])

        r = self.profe.get(f'/materiales/{doc.pk}/word/')
        from docx import Document
        w = Document(io.BytesIO(r.content))
        textos = '\n'.join(p.text for p in w.paragraphs)
        self.assertIn('Fracciones 4.º', textos)
        self.assertIn('Matemáticas', textos)
        self.assertIn('A. 2/4', textos)
        self.assertEqual(len(w.inline_shapes), 1)

    def test_publicar_en_el_aula_y_permisos(self):
        doc = self.crear('taller')
        self.profe.get(f'/aula/asignacion/{self.asig.id}/')
        aula = Aula.objects.get(asignacion=self.asig)
        self.assertContains(self.profe.get(f'/materiales/{doc.pk}/'), 'Publicar en un aula')
        self.profe.post(f'/materiales/{doc.pk}/accion/', {'accion': 'publicar', 'aula': aula.pk, 'modo': 'compacto'})
        m = Material.objects.get(aula=aula)
        self.assertEqual(m.titulo, 'Fracciones')
        self.assertTrue(m.nombre_archivo.endswith('.pdf'))
        with m.archivo.open('rb') as f:
            self.assertTrue(f.read(4) == b'%PDF')
        # El estudiante no entra; otro docente no ve el documento
        self.assertEqual(self.cliente(self.estudiantes[0].user).get('/materiales/').status_code, 403)
        from notas.models import Docente
        otro = Docente.objects.create(colegio=self.a, user=User.objects.create_user('docente2', password=CLAVE))
        self.assertEqual(self.cliente(otro.user).get(f'/materiales/{doc.pk}/').status_code, 404)
        self.assertEqual(self.cliente(otro.user).get('/materiales/').status_code, 200)
        # Sin el módulo, no hay Mis materiales
        self.a.modulos_activos = [x for x in self.a.modulos_activos if x != 'salon']
        self.a.save(update_fields=['modulos_activos'])
        self.assertEqual(self.profe.get('/materiales/').status_code, 403)

    def test_formato_del_colegio(self):
        rectora = self.cliente(self.rectora)
        r = rectora.post('/materiales/nuevo/', {'tipo': 'escuela_nueva', 'titulo': 'Formato institucional'})
        modelo = Documento.objects.get(titulo='Formato institucional')
        # El docente no marca formatos del colegio
        doc = self.crear('libre')
        self.profe.post(f'/materiales/{doc.pk}/accion/', {'accion': 'formato'})
        doc.refresh_from_db()
        self.assertFalse(doc.plantilla_colegio)
        rectora.post(f'/materiales/{modelo.pk}/accion/', {'accion': 'formato'})
        self.assertContains(self.profe.get('/materiales/'), 'Formato institucional')
        nuevo = self.crear('institucional', base=modelo.pk)
        self.assertEqual(nuevo.tipo, 'institucional')
        self.assertEqual([b['tipo'] for b in nuevo.contenido['bloques']], [b['tipo'] for b in modelo.contenido['bloques']])
        self.assertNotEqual(nuevo.contenido['bloques'][0]['id'], modelo.contenido['bloques'][0]['id'])
        # Duplicar y eliminar
        self.profe.post(f'/materiales/{nuevo.pk}/accion/', {'accion': 'duplicar'})
        self.assertTrue(Documento.objects.filter(titulo='Fracciones (copia)').exists())
        self.profe.post(f'/materiales/{nuevo.pk}/accion/', {'accion': 'eliminar'})
        self.assertFalse(Documento.objects.filter(pk=nuevo.pk).exists())


@override_settings(MEDIA_ROOT=MEDIA)
class PlanIndividual(ColegioDePrueba):

    def test_encabezado_propio_y_nube_llena(self):
        from salon_digital.models import Cuenta
        u = User.objects.create_user('libre', password=CLAVE, first_name='Lina', last_name='Libre')
        cuenta = Cuenta.objects.create(usuario=u, slug='lina', nombre_publico='Lina Libre', estado='verificada')
        c = Client(HTTP_HOST='localhost')
        c.force_login(u)
        self.assertContains(c.get('/materiales/'), 'Mi encabezado')
        c.post('/materiales/encabezado/', {'linea1': 'Prof. Lina Libre', 'linea2': 'Matemáticas', 'logo': png('logo.png', 300, 300)})
        enc = EncabezadoDocente.objects.get(usuario=u)
        self.assertEqual(enc.lineas, ['Prof. Lina Libre', 'Matemáticas'])
        self.assertTrue(enc.logo)
        c.post('/materiales/nuevo/', {'tipo': 'evaluacion', 'titulo': 'Quiz'})
        doc = Documento.objects.get(propietario=u)
        self.assertIsNone(doc.colegio)
        r = c.get(f'/materiales/{doc.pk}/pdf/')
        self.assertTrue(r.content.startswith(b'%PDF'))
        r = c.get(f'/materiales/{doc.pk}/word/')
        from docx import Document
        self.assertIn('Prof. Lina Libre', '\n'.join(cell.text for t in Document(io.BytesIO(r.content)).tables
                                                   for row in t.rows for cell in row.cells))
        # Prueba gratis: 200 MB; si se llena, no deja subir más
        cuenta.plan = 'prueba'
        cuenta.save(update_fields=['plan'])
        Imagen.objects.create(propietario=u, archivo='lleno.jpg', peso=200 * 1024 * 1024)
        r = c.post('/materiales/imagen/', {'imagen': png()})
        self.assertEqual(r.status_code, 400)
        self.assertIn('nube está llena', r.json()['error'])
        # Una persona sin cuenta de docente no entra
        nadie = User.objects.create_user('nadie', password=CLAVE)
        c2 = Client(HTTP_HOST='localhost')
        c2.force_login(nadie)
        self.assertEqual(c2.get('/materiales/').status_code, 403)
