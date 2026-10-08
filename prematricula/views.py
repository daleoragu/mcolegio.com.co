# -*- coding: utf-8 -*-
"""Prematrícula · pantallas de la familia (sin usuario) y de la secretaría."""
import io
import os
import zipfile

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import user_passes_test
from django.core.mail import send_mail
from django.db import transaction
from django.db.models import Count, Q
from django.forms import modelformset_factory
from django.http import FileResponse, Http404, HttpResponse, HttpResponseNotFound
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from notas.models import Curso, Estudiante
from notas.permisos import es_admin_usuario

from . import logica
from .forms import ConfiguracionForm, PreguntaForm, RequisitoForm, SolicitudForm
from .models import Documento, PreguntaEncuesta, Requisito, Solicitud


def _colegio_o_404(request):
    if not getattr(request, 'colegio', None):
        raise Http404('Colegio no configurado')
    return request.colegio


def _correo_configurado():
    return getattr(settings, 'EMAIL_HOST', '') not in ('', 'localhost')


def _avisar_familia(request, solicitud, asunto, cuerpo):
    """Correo al acudiente con el enlace para consultar la solicitud."""
    if not (solicitud.acudiente_correo and _correo_configurado()):
        return False
    enlace = request.build_absolute_uri(
        reverse('prematricula:estado', args=[logica.token_consulta(solicitud)]))
    texto = (f'{solicitud.colegio.nombre}\n\n{cuerpo}\n\nRadicado: {solicitud.radicado}\n'
             f'Consulte su solicitud aquí: {enlace}\n')
    try:
        send_mail(asunto, texto, settings.DEFAULT_FROM_EMAIL, [solicitud.acudiente_correo])
        return True
    except Exception:
        return False


# ===========================================================================
# Familia (portal, sin iniciar sesión)
# ===========================================================================

def inicio(request):
    colegio = _colegio_o_404(request)
    conf = logica.configuracion_de(colegio, crear=False)
    return render(request, 'prematricula/inicio.html', {
        'colegio': colegio, 'conf': conf, 'abierta': bool(conf and conf.esta_abierta()),
    })


def verificar_antiguo(request):
    """El antiguo se identifica con documento y fecha de nacimiento."""
    colegio = _colegio_o_404(request)
    conf = logica.configuracion_de(colegio, crear=False)
    if not (conf and conf.esta_abierta() and conf.permite_antiguos):
        return redirect('prematricula:inicio')
    error = ''
    if request.method == 'POST':
        from datetime import date
        try:
            nacimiento = date.fromisoformat(request.POST.get('fecha_nacimiento') or '')
        except ValueError:
            nacimiento = None
        est = logica.buscar_antiguo(colegio, request.POST.get('documento'), nacimiento) if nacimiento else None
        if est is None:
            error = ('No encontramos un estudiante del colegio con ese documento y esa fecha de '
                     'nacimiento. Revise los datos o comuníquese con secretaría.')
        elif Solicitud.objects.filter(colegio=colegio, ano_lectivo=conf.ano_lectivo,
                                      estudiante_antiguo=est).exclude(estado='RECHAZADA').exists():
            error = 'Este estudiante ya tiene una solicitud para ese año. Consulte su estado con el radicado.'
        else:
            request.session['pm_antiguo'] = est.id
            return redirect('prematricula:formulario', tipo='antiguo')
    return render(request, 'prematricula/verificar.html', {'colegio': colegio, 'conf': conf, 'error': error})


def formulario(request, tipo):
    colegio = _colegio_o_404(request)
    conf = logica.configuracion_de(colegio, crear=False)
    if not (conf and conf.esta_abierta()):
        return redirect('prematricula:inicio')
    tipo = 'ANTIGUO' if tipo == 'antiguo' else 'NUEVO'
    antiguo = None
    if tipo == 'ANTIGUO':
        antiguo = Estudiante.objects.filter(id=request.session.get('pm_antiguo'), colegio=colegio).first()
        if antiguo is None or not conf.permite_antiguos:
            return redirect('prematricula:verificar')

    inicial = logica.datos_iniciales_antiguo(antiguo) if antiguo else {}
    if request.method == 'POST':
        form = SolicitudForm(request.POST, request.FILES, colegio=colegio, tipo=tipo, initial=inicial)
        if form.is_valid():
            with transaction.atomic():
                s = form.save(commit=False)
                s.colegio, s.tipo, s.ano_lectivo = colegio, tipo, conf.ano_lectivo
                s.estudiante_antiguo = antiguo
                s.acepta_datos = True
                s.encuesta = form.encuesta()
                s.radicado = logica.nuevo_radicado(colegio, conf.ano_lectivo)
                s.save()
                for requisito, archivo in form.archivos():
                    Documento.objects.create(solicitud=s, requisito=requisito, archivo=archivo,
                                             nombre_original=archivo.name[:200])
            request.session.pop('pm_antiguo', None)
            _avisar_familia(request, s, f'Prematrícula recibida · {s.radicado}',
                            f'Recibimos la solicitud de prematrícula de {s.nombres} {s.apellidos} '
                            f'para {s.nombre_grado} de {s.ano_lectivo}. Le avisaremos cuando la revisemos.')
            return redirect('prematricula:estado', token=logica.token_consulta(s))
        messages.error(request, 'Revise los campos marcados en rojo.')
    else:
        form = SolicitudForm(colegio=colegio, tipo=tipo, initial=inicial)
    from .forms import CAMPOS_ESTUDIANTE, CAMPOS_FAMILIA, CAMPOS_PROCEDENCIA, CAMPOS_SALUD
    return render(request, 'prematricula/formulario.html', {
        'colegio': colegio, 'conf': conf, 'form': form, 'tipo': tipo, 'antiguo': antiguo,
        'secciones': [
            ('Estudiante', form.campos(CAMPOS_ESTUDIANTE + CAMPOS_PROCEDENCIA)),
            ('Salud y emergencia', form.campos(CAMPOS_SALUD)),
            ('Familia', form.campos(CAMPOS_FAMILIA)),
        ],
        'aspira': form.campos(['grado', 'sede', 'jornada']),
        'documentos': form.campos_documentos(),
        'encuesta': form.campos_encuesta(),
        'tratamiento': conf.tratamiento_datos or logica.TRATAMIENTO_GENERAL,
    })


def consultar(request):
    colegio = _colegio_o_404(request)
    error = ''
    if request.method == 'POST':
        radicado = (request.POST.get('radicado') or '').strip().upper()
        documento = (request.POST.get('documento') or '').strip()
        s = Solicitud.objects.filter(colegio=colegio, radicado=radicado, numero_documento=documento).first()
        if s:
            return redirect('prematricula:estado', token=logica.token_consulta(s))
        error = 'No hay una solicitud con ese radicado y ese documento.'
    return render(request, 'prematricula/consultar.html', {'colegio': colegio, 'error': error})


def estado(request, token):
    colegio = _colegio_o_404(request)
    s = logica.solicitud_de_token(token, colegio)
    if s is None:
        return render(request, 'prematricula/consultar.html',
                      {'colegio': colegio, 'error': 'El enlace no es válido. Consulte con su radicado.'},
                      status=404)
    puede_subir = s.estado in ('RECIBIDA', 'REVISION', 'PENDIENTE')
    if request.method == 'POST' and puede_subir:
        subidos = 0
        for r in s.requisitos():
            f = request.FILES.get(f'doc_{r.id}')
            if not (f and r.pide_archivo):
                continue
            ext = os.path.splitext(f.name)[1].lower()
            if ext not in {'.pdf', '.jpg', '.jpeg', '.png', '.webp', '.heic'} or f.size > 6 * 1024 * 1024:
                messages.error(request, f'«{f.name}» no se recibió: debe ser PDF o foto de máximo 6 MB.')
                continue
            Documento.objects.create(solicitud=s, requisito=r, archivo=f, nombre_original=f.name[:200])
            s.checklist.pop(str(r.id), None)        # vuelve a quedar por revisar
            subidos += 1
        if subidos:
            if s.estado == 'PENDIENTE':
                s.estado = 'REVISION'
            s.save()
            messages.success(request, f'Se recibieron {subidos} documento(s).')
        return redirect('prematricula:estado', token=token)
    filas = [(r, d, c) for r, d, c in s.estado_requisitos()]
    return render(request, 'prematricula/estado.html', {
        'colegio': colegio, 's': s, 'filas': filas, 'puede_subir': puede_subir,
        'enlace': request.build_absolute_uri(),
    })


# ===========================================================================
# Secretaría / administración
# ===========================================================================

admin_requerido = user_passes_test(es_admin_usuario)


@admin_requerido
def panel(request):
    colegio = _colegio_o_404(request)
    conf = logica.configuracion_de(colegio)
    ano = int(request.GET.get('ano') or conf.ano_lectivo)
    qs = Solicitud.objects.filter(colegio=colegio, ano_lectivo=ano).select_related('sede', 'curso_asignado')
    filtros = {k: request.GET.get(k, '') for k in ('estado', 'grado', 'tipo', 'q')}
    if filtros['estado']:
        qs = qs.filter(estado=filtros['estado'])
    if filtros['grado'] != '':
        qs = qs.filter(grado=filtros['grado'])
    if filtros['tipo']:
        qs = qs.filter(tipo=filtros['tipo'])
    if filtros['q']:
        q = filtros['q']
        qs = qs.filter(Q(nombres__icontains=q) | Q(apellidos__icontains=q) |
                       Q(numero_documento__icontains=q) | Q(radicado__icontains=q))
    todas = Solicitud.objects.filter(colegio=colegio, ano_lectivo=ano)
    por_estado = dict(todas.values_list('estado').annotate(n=Count('id')))
    por_grado = list(todas.values('grado').annotate(n=Count('id'),
                                                   aprobadas=Count('id', filter=Q(estado='APROBADA')))
                     .order_by('grado'))
    from notas.models.perfiles import NOMBRE_GRADO
    for g in por_grado:
        g['nombre'] = NOMBRE_GRADO.get(g['grado'], g['grado'])
    pendientes_activar = Estudiante.objects.filter(
        colegio=colegio, is_active=False,
        id__in=todas.filter(estado='APROBADA', tipo='NUEVO').values('estudiante_creado')).count()
    anos = sorted(set(Solicitud.objects.filter(colegio=colegio).values_list('ano_lectivo', flat=True))
                  | {conf.ano_lectivo}, reverse=True)
    return render(request, 'prematricula/panel.html', {
        'colegio': colegio, 'conf': conf, 'ano': ano, 'anos': anos, 'solicitudes': qs,
        'filtros': filtros, 'estados': Solicitud.ESTADOS, 'tipos': Solicitud.TIPOS,
        'por_estado': [(c, n, por_estado.get(c, 0)) for c, n in Solicitud.ESTADOS],
        'total': todas.count(), 'por_grado': por_grado, 'pendientes_activar': pendientes_activar,
        'url_portal': request.build_absolute_uri(reverse('prematricula:inicio')),
    })


@admin_requerido
def detalle(request, solicitud_id):
    colegio = _colegio_o_404(request)
    s = get_object_or_404(Solicitud.objects.select_related('sede', 'curso_asignado', 'estudiante_antiguo__curso'),
                          id=solicitud_id, colegio=colegio)
    if request.method == 'POST':
        accion = request.POST.get('accion') or 'guardar'
        ids = {str(r.id) for r in s.requisitos()}
        s.checklist = {i: (request.POST.get(f'req_{i}') == 'on') for i in ids
                       if request.POST.get(f'req_{i}') == 'on' or i in s.checklist}
        curso_id = request.POST.get('curso_asignado') or ''
        s.curso_asignado = Curso.objects.filter(id=curso_id, colegio=colegio).first() if curso_id.isdigit() else None
        s.nota_interna = (request.POST.get('nota_interna') or '').strip()
        s.mensaje_familia = (request.POST.get('mensaje_familia') or '').strip()
        if s.estado == 'RECIBIDA':
            s.estado = 'REVISION'
        s.revisada_por = request.user
        s.save()

        if accion == 'pendiente':
            s.estado = 'PENDIENTE'
            s.save(update_fields=['estado'])
            faltan = ', '.join(r.nombre for r in s.faltantes())
            avisado = _avisar_familia(
                request, s, f'Prematrícula {s.radicado}: falta información',
                (s.mensaje_familia or 'Revisamos su solicitud y falta información.')
                + (f'\nPendiente: {faltan}.' if faltan else '')
                + '\nPuede subir los documentos desde el enlace de consulta.')
            messages.warning(request, 'Solicitud marcada «Falta algo».'
                             + (' Se le avisó a la familia por correo.' if avisado else ''))
        elif accion == 'rechazar':
            s.estado = 'RECHAZADA'
            s.decidida = timezone.now()
            s.save(update_fields=['estado', 'decidida'])
            _avisar_familia(request, s, f'Prematrícula {s.radicado}',
                            s.mensaje_familia or 'Su solicitud de prematrícula no fue aprobada. '
                                                 'Comuníquese con secretaría para más información.')
            messages.info(request, 'Solicitud marcada como no aprobada.')
        elif accion == 'aprobar':
            faltan = s.faltantes()
            if faltan and request.POST.get('aprobar_igual') != 'on':
                messages.error(request, 'Faltan requisitos obligatorios: ' + ', '.join(r.nombre for r in faltan)
                               + '. Márquelos como cumplidos o confirme «aprobar aunque falten».')
                return redirect('prematricula:detalle', solicitud_id=s.id)
            try:
                est, nuevo, clave = logica.aprobar(s, request.user)
            except logica.NoSePuedeAprobar as e:
                messages.error(request, str(e))
                return redirect('prematricula:detalle', solicitud_id=s.id)
            cuerpo = (s.mensaje_familia or f'¡Bienvenidos! La matrícula de {s.nombres} {s.apellidos} para '
                      f'{s.nombre_grado} de {s.ano_lectivo} fue aprobada.')
            if nuevo:
                cuerpo += (f'\nCurso: {s.curso_asignado}.\nUsuario de la plataforma: {nuevo.username}\n'
                           f'Contraseña inicial: {clave} (cámbiela al entrar).')
            _avisar_familia(request, s, f'Matrícula aprobada · {s.radicado}', cuerpo)
            if nuevo:
                messages.success(request, f'Matrícula aprobada. Estudiante creado en {s.curso_asignado}: '
                                          f'usuario {nuevo.username}, contraseña inicial {clave}.'
                                 + ('' if est.is_active else f' Queda inactivo hasta que empiece {s.ano_lectivo}.'))
            else:
                messages.success(request, 'Matrícula renovada: se actualizó la ficha del estudiante.')
        else:
            messages.success(request, 'Revisión guardada.')
        return redirect('prematricula:detalle', solicitud_id=s.id)

    cursos = Curso.objects.filter(colegio=colegio).select_related('sede').order_by('grado', 'orden', 'nombre')
    sugeridos = [c for c in cursos if c.grado == s.grado and (not s.sede_id or c.sede_id == s.sede_id)]
    respuestas = []
    for p in PreguntaEncuesta.objects.filter(colegio=colegio):
        v = s.encuesta.get(str(p.id))
        if v not in (None, '', []):
            respuestas.append((p.texto, ', '.join(v) if isinstance(v, list) else v))
    return render(request, 'prematricula/detalle.html', {
        'colegio': colegio, 's': s, 'filas': s.estado_requisitos(), 'cursos': cursos,
        'sugeridos': sugeridos, 'respuestas': respuestas,
        'otros_docs': s.documentos.filter(requisito__isnull=True),
    })


@admin_requerido
def documento(request, documento_id):
    """Los documentos no quedan con enlace público: se sirven solo a la secretaría."""
    colegio = _colegio_o_404(request)
    d = get_object_or_404(Documento, id=documento_id, solicitud__colegio=colegio)
    try:
        archivo = d.archivo.open('rb')
    except (FileNotFoundError, OSError):
        raise Http404('El archivo ya no está.')
    return FileResponse(archivo, filename=d.nombre_original or os.path.basename(d.archivo.name))


@admin_requerido
def documentos_zip(request, solicitud_id):
    colegio = _colegio_o_404(request)
    s = get_object_or_404(Solicitud, id=solicitud_id, colegio=colegio)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as z:
        for i, d in enumerate(s.documentos.select_related('requisito'), start=1):
            try:
                with d.archivo.open('rb') as f:
                    contenido = f.read()
            except (FileNotFoundError, OSError):
                continue
            ext = os.path.splitext(d.archivo.name)[1]
            nombre = (d.requisito.nombre if d.requisito else 'Documento')[:60]
            z.writestr(f'{i:02d} {nombre}{ext}', contenido)
    respuesta = HttpResponse(buffer.getvalue(), content_type='application/zip')
    respuesta['Content-Disposition'] = f'attachment; filename="{s.radicado} {s.apellidos}.zip"'
    return respuesta


@admin_requerido
def exportar(request):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    colegio = _colegio_o_404(request)
    conf = logica.configuracion_de(colegio)
    ano = int(request.GET.get('ano') or conf.ano_lectivo)
    solicitudes = list(Solicitud.objects.filter(colegio=colegio, ano_lectivo=ano)
                       .select_related('sede', 'curso_asignado').order_by('grado', 'apellidos'))
    preguntas = list(PreguntaEncuesta.objects.filter(colegio=colegio))
    campos = [('radicado', 'Radicado'), ('get_estado_display', 'Estado'), ('get_tipo_display', 'Tipo'),
              ('nombre_grado', 'Grado'), ('sede', 'Sede'), ('jornada', 'Jornada'), ('curso_asignado', 'Curso asignado'),
              ('apellidos', 'Apellidos'), ('nombres', 'Nombres'), ('tipo_documento', 'Tipo doc.'),
              ('numero_documento', 'Documento'), ('fecha_nacimiento', 'Nacimiento'),
              ('lugar_nacimiento', 'Lugar de nacimiento'), ('get_genero_display', 'Sexo'),
              ('direccion', 'Dirección'), ('barrio', 'Barrio'), ('telefono', 'Teléfono'),
              ('colegio_procedencia', 'Colegio de procedencia'), ('ultimo_grado', 'Último grado'),
              ('eps', 'EPS'), ('rh', 'RH'), ('alergias', 'Salud'), ('necesita_apoyo', 'Requiere apoyo'),
              ('apoyo_detalle', 'Apoyo'), ('emergencia_nombre', 'Emergencia'), ('emergencia_telefono', 'Tel. emergencia'),
              ('acudiente_nombre', 'Acudiente'), ('acudiente_documento', 'Doc. acudiente'),
              ('acudiente_parentesco', 'Parentesco'), ('acudiente_celular', 'Celular acudiente'),
              ('acudiente_correo', 'Correo acudiente'), ('madre_nombre', 'Madre'), ('madre_celular', 'Cel. madre'),
              ('padre_nombre', 'Padre'), ('padre_celular', 'Cel. padre'), ('creada', 'Recibida')]
    libro = openpyxl.Workbook()
    h = libro.active
    h.title = f'Prematrícula {ano}'
    h.append([t for _, t in campos] + ['Requisitos pendientes'] + [p.texto for p in preguntas])
    for c in h[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='193661')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for s in solicitudes:
        fila = []
        for campo, _ in campos:
            v = getattr(s, campo)
            v = v() if callable(v) else v
            if hasattr(v, 'tzinfo') and getattr(v, 'tzinfo', None):
                v = timezone.localtime(v).replace(tzinfo=None)
            if isinstance(v, bool):
                v = 'Sí' if v else 'No'
            fila.append(v if isinstance(v, (int, float, str)) or v is None or hasattr(v, 'year') else str(v))
        fila.append(', '.join(r.nombre for r in s.faltantes()))
        for p in preguntas:
            v = s.encuesta.get(str(p.id), '')
            fila.append(', '.join(v) if isinstance(v, list) else v)
        h.append(fila)
    h.freeze_panes = 'C2'
    for i in range(1, h.max_column + 1):
        h.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 18
    respuesta = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    respuesta['Content-Disposition'] = f'attachment; filename="prematricula-{ano}.xlsx"'
    libro.save(respuesta)
    return respuesta


@admin_requerido
def encuesta(request):
    colegio = _colegio_o_404(request)
    conf = logica.configuracion_de(colegio)
    ano = int(request.GET.get('ano') or conf.ano_lectivo)
    solicitudes = Solicitud.objects.filter(colegio=colegio, ano_lectivo=ano)
    return render(request, 'prematricula/encuesta.html', {
        'colegio': colegio, 'ano': ano, 'total': solicitudes.count(),
        'resumen': logica.resumen_encuesta(colegio, solicitudes),
    })


@admin_requerido
def configuracion(request):
    colegio = _colegio_o_404(request)
    conf = logica.configuracion_de(colegio)
    Requisitos = modelformset_factory(Requisito, form=RequisitoForm, extra=1, can_delete=True)
    Preguntas = modelformset_factory(PreguntaEncuesta, form=PreguntaForm, extra=1, can_delete=True)
    req_qs = Requisito.objects.filter(colegio=colegio)
    preg_qs = PreguntaEncuesta.objects.filter(colegio=colegio)
    if request.method == 'POST':
        form = ConfiguracionForm(request.POST, instance=conf)
        reqs = Requisitos(request.POST, queryset=req_qs, prefix='req')
        pregs = Preguntas(request.POST, queryset=preg_qs, prefix='preg')
        if form.is_valid() and reqs.is_valid() and pregs.is_valid():
            with transaction.atomic():
                form.save()
                for fs in (reqs, pregs):
                    for obj in fs.save(commit=False):
                        obj.colegio = colegio
                        obj.save()
                    for obj in fs.deleted_objects:
                        if isinstance(obj, Requisito) and obj.documento_set.exists():
                            obj.activo = False          # tiene documentos subidos: se desactiva
                            obj.save(update_fields=['activo'])
                        else:
                            obj.delete()
            messages.success(request, 'Prematrícula configurada.')
            return redirect('prematricula:configuracion')
        messages.error(request, 'Revise los campos marcados.')
    else:
        form = ConfiguracionForm(instance=conf)
        reqs = Requisitos(queryset=req_qs, prefix='req')
        pregs = Preguntas(queryset=preg_qs, prefix='preg')
    return render(request, 'prematricula/configuracion.html', {
        'colegio': colegio, 'form': form, 'reqs': reqs, 'pregs': pregs,
        'url_portal': request.build_absolute_uri(reverse('prematricula:inicio')),
    })


@admin_requerido
@require_POST
def activar(request):
    """Activa a los estudiantes nuevos aprobados de un año (cuando el año empieza)."""
    colegio = _colegio_o_404(request)
    ano = int(request.POST.get('ano'))
    ids = Solicitud.objects.filter(colegio=colegio, ano_lectivo=ano, estado='APROBADA',
                                   tipo='NUEVO').values('estudiante_creado')
    n = Estudiante.objects.filter(colegio=colegio, is_active=False, id__in=ids).update(is_active=True)
    messages.success(request, f'Se activaron {n} estudiante(s) matriculados para {ano}.')
    return redirect(f"{reverse('prematricula:panel')}?ano={ano}")


@admin_requerido
@require_POST
def alternar(request):
    """Habilita o deshabilita la prematrícula en el portal con un clic."""
    colegio = _colegio_o_404(request)
    conf = logica.configuracion_de(colegio)
    conf.abierta = request.POST.get('abrir') == '1'
    conf.save(update_fields=['abierta'])
    if conf.abierta and not conf.esta_abierta():
        messages.warning(request, f'Quedó habilitada, pero la fecha de cierre ({conf.fecha_cierre:%d/%m/%Y}) ya pasó: '
                                  f'cámbiela en «Configurar» para que salga en el portal.')
    elif conf.abierta:
        messages.success(request, f'Prematrícula {conf.ano_lectivo} habilitada: ya sale en el portal del colegio.')
    else:
        messages.info(request, 'Prematrícula deshabilitada: ya no sale en el portal. Las solicitudes recibidas se conservan.')
    return redirect('prematricula:panel')
