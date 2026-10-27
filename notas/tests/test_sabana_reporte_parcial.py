"""Sábana del reporte parcial: estudiantes × asignaturas con la marca y los totales."""
from io import BytesIO

import openpyxl
from django.contrib.auth.models import User

from notas.models import AsignacionDocente, Curso, Docente, Materia, ReporteParcial

from .base import CLAVE, ColegioDePrueba


class SabanaReporteParcial(ColegioDePrueba):

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.ana, cls.beto, cls.caro = cls.estudiantes
        u2 = User.objects.create_user('profe2', password=CLAVE, first_name='Sol', last_name='Sociales')
        cls.docente2 = Docente.objects.create(colegio=cls.a, user=u2)
        cls.sociales = Materia.objects.create(colegio=cls.a, nombre='Sociales')
        cls.asig_soc = AsignacionDocente.objects.create(colegio=cls.a, docente=cls.docente2,
                                                        materia=cls.sociales, curso=cls.curso)
        ingles = Materia.objects.create(colegio=cls.a, nombre='Inglés')
        AsignacionDocente.objects.create(colegio=cls.a, docente=cls.docente2, materia=ingles, curso=cls.curso)
        for est, asig, si in [(cls.ana, cls.asig, True), (cls.ana, cls.asig_soc, True),
                              (cls.beto, cls.asig, False), (cls.beto, cls.asig_soc, False),
                              (cls.caro, cls.asig, True), (cls.caro, cls.asig_soc, False)]:
            ReporteParcial.objects.create(colegio=cls.a, estudiante=est, asignacion=asig, periodo=cls.p1,
                                          presenta_dificultades=si,
                                          sugerencia='Entregar el taller' if si else '')
        # Otro colegio: no debe salir
        otro = Curso.objects.create(colegio=cls.b, nombre='701', grado=7, subgrupo='1')
        cls.otro_curso = otro

    def url(self, **extra):
        q = '&'.join(f'{k}={v}' for k, v in {'periodo': self.p1.id, 'curso': self.curso.id, **extra}.items())
        return f'/reporte-parcial/sabana/?{q}'

    def test_el_director_ve_la_sabana_con_conteos(self):
        r = self.cliente(self.u_docente).get(self.url())
        self.assertEqual(r.status_code, 200)
        self.assertEqual([m['nombre'] for m in r.context['materias']], ['Inglés', 'Matemáticas', 'Sociales'])
        totales = {f['est'].id: f['total'] for f in r.context['filas']}
        self.assertEqual(totales, {self.ana.id: 2, self.beto.id: 0, self.caro.id: 1})
        self.assertEqual(r.context['total_reportes'], 3)
        self.assertEqual(r.context['estudiantes_reportados'], 2)
        self.assertEqual(r.context['estudiantes_sin_reporte'], 1)
        self.assertEqual([m['total'] for m in r.context['materias']], [0, 2, 1])
        self.assertEqual([m['nombre'] for m in r.context['pendientes']], ['Inglés'])
        self.assertContains(r, 'Entregar el taller')
        self.assertContains(r, 'Asignaturas sin reporte guardado')

    def test_alerta_con_tres_o_mas(self):
        from notas import reporte_sabana
        ingles = AsignacionDocente.objects.get(curso=self.curso, materia__nombre='Inglés')
        ReporteParcial.objects.create(colegio=self.a, estudiante=self.ana, asignacion=ingles, periodo=self.p1,
                                      presenta_dificultades=True)
        datos = reporte_sabana.calcular(self.curso, self.p1)
        self.assertEqual(datos['estudiantes_alerta'], 1)
        self.assertTrue(next(f for f in datos['filas'] if f['est'] == self.ana)['alerta'])

    def test_excel(self):
        r = self.cliente(self.rectora).get(self.url(formato='excel'))
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r['Content-Type'])
        ws = openpyxl.load_workbook(BytesIO(r.content)).active
        filas = {ws.cell(row=i, column=2).value: [ws.cell(row=i, column=c).value for c in range(4, 8)]
                 for i in range(6, 9)}
        self.assertEqual(filas['Prueba Ana'], [None, 'X', 'X', 2])
        self.assertEqual(filas['Prueba Beto'], [None, None, None, 0])

    def test_pdf(self):
        r = self.cliente(self.u_docente).get(self.url(formato='pdf'))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], 'application/pdf')

    def test_docente_que_no_es_director_no_entra(self):
        self.assertEqual(self.cliente(self.docente2.user).get(self.url()).status_code, 403)

    def test_solo_cursos_del_colegio(self):
        r = self.cliente(self.rectora).get(self.url(curso=self.otro_curso.id))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context['curso'], self.curso)
        self.assertNotIn(self.otro_curso, r.context['cursos'])
