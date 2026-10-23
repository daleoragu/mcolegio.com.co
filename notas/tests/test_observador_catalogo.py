"""Catálogo de faltas (Tipo I, II, III) y observaciones subidas desde Excel."""
import datetime
import io

from django.utils import timezone
from openpyxl import Workbook, load_workbook

from notas import observador_catalogo as oc
from notas.models import FaltaCatalogo, FichaEstudiante, Notificacion, RegistroObservador

from .base import ColegioDePrueba


def _excel(filas, hoja='Observaciones'):
    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    ws.append(oc.COLUMNAS)
    for f in filas:
        ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = 'obs.xlsx'
    return buf


class CatalogoDeFaltas(ColegioDePrueba):

    def test_ejemplos_sin_repetir(self):
        n = oc.cargar_ejemplos(self.a)
        self.assertEqual(n, len(oc.EJEMPLOS))
        self.assertEqual(oc.cargar_ejemplos(self.a), 0)
        self.assertEqual(set(FaltaCatalogo.objects.values_list('tipo_falta', flat=True)), {'I', 'II', 'III'})

    def test_importar_catalogo_desde_excel(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['Tipo', 'Numeral', 'Descripción'])
        ws.append(['Tipo 2', 'Art. 9', 'Ciberacoso a un compañero'])
        ws.append(['1', '', 'Llegar tarde'])
        ws.append(['IV', '', 'No existe'])
        ws.append(['I', '', ''])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        n, errores = oc.importar_catalogo(self.a, buf)
        self.assertEqual(n, 2)
        self.assertEqual(len(errores), 2)
        self.assertEqual(FaltaCatalogo.objects.get(numeral='Art. 9').tipo_falta, 'II')

    def test_pantalla_del_catalogo(self):
        c = self.cliente(self.rectora)
        c.post('/observador/catalogo-faltas/', {'accion': 'crear', 'tipo_falta': 'III', 'descripcion': 'Porte de armas'})
        f = FaltaCatalogo.objects.get()
        self.assertContains(c.get('/observador/catalogo-faltas/?tipo=III'), 'Porte de armas')
        self.assertNotContains(c.get('/observador/catalogo-faltas/?tipo=I'), 'Porte de armas')
        c.post('/observador/catalogo-faltas/', {'accion': 'alternar', 'id': f.id})
        f.refresh_from_db()
        self.assertFalse(f.activa)
        self.assertEqual(oc.catalogo_para_buscar(self.a), [])
        r = c.get('/observador/catalogo-faltas/?descargar=1')
        self.assertEqual(load_workbook(io.BytesIO(r.content)).active['C2'].value, 'Porte de armas')
        # Un docente no administra el catálogo
        self.assertEqual(self.cliente(self.u_docente).get('/observador/catalogo-faltas/').status_code, 302)

    def test_el_docente_escoge_la_falta_al_anotar(self):
        f = FaltaCatalogo.objects.create(colegio=self.a, tipo_falta='II', descripcion='Ciberacoso')
        otra = FaltaCatalogo.objects.create(colegio=self.b, tipo_falta='I', descripcion='De otro colegio')
        est = self.estudiantes[0]
        c = self.cliente(self.u_docente)
        self.assertContains(c.get(f'/observador/crear/{est.id}/'), 'Buscar en el catálogo de faltas')
        c.post(f'/observador/crear/{est.id}/', {'fecha_suceso': '2026-03-10', 'tipo': 'COMPORTAMENTAL',
                                                'subtipo': 'NEGATIVA', 'descripcion': 'Ciberacoso', 'falta': f.id})
        r = RegistroObservador.objects.get()
        self.assertEqual((r.falta, r.tipo_falta), (f, 'II'))
        # Una falta de otro colegio no se acepta
        c.post(f'/observador/crear/{est.id}/', {'fecha_suceso': '2026-03-11', 'tipo': 'COMPORTAMENTAL',
                                                'subtipo': 'NEGATIVA', 'descripcion': 'x', 'falta': otra.id})
        self.assertIsNone(RegistroObservador.objects.get(fecha_suceso='2026-03-11').falta)


class ObservadorEnExcel(ColegioDePrueba):

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        FichaEstudiante.objects.create(estudiante=cls.estudiantes[0], numero_documento='1110555')
        cls.falta = FaltaCatalogo.objects.create(colegio=cls.a, tipo_falta='I', numeral='Art. 3',
                                                 descripcion='Usar el celular en clase')

    def test_plantilla_trae_estudiantes_y_catalogo(self):
        r = self.cliente(self.u_docente).get(f'/observador/excel/?curso_id={self.curso.id}&descargar=1')
        wb = load_workbook(io.BytesIO(r.content))
        docs = [c.value for c in wb['Observaciones']['A'][1:]]
        self.assertIn('1110555', docs)
        self.assertIn('beto', docs)                      # sin documento: el usuario
        self.assertEqual(wb['Catálogo de faltas']['D2'].value, 'Usar el celular en clase')

    def test_subir_observaciones(self):
        filas = [
            ['1110555', 'PRUEBA ANA', '10/03/2026', '', '', '', self.falta.id, 'En clase de inglés', 'Dice que era una emergencia'],
            [1110555.0, 'PRUEBA ANA', datetime.datetime(2026, 3, 12), 'Académica', 'Positiva', '', '', 'Excelente exposición', ''],
            ['beto', 'PRUEBA BETO', '', '', '', '', '', '', ''],          # sin nada: se salta
        ]
        c = self.cliente(self.u_docente)
        r = c.post('/observador/excel/', {'curso_id': self.curso.id, 'archivo': _excel(filas)})
        self.assertContains(r, '2 observaciones guardadas')
        a, b = RegistroObservador.objects.order_by('fecha_suceso')
        self.assertEqual((a.tipo, a.subtipo, a.tipo_falta, a.falta), ('COMPORTAMENTAL', 'NEGATIVA', 'I', self.falta))
        self.assertTrue(a.descripcion.startswith('Usar el celular en clase\nEn clase de inglés'))
        self.assertEqual(a.descargo_estudiante, 'Dice que era una emergencia')
        self.assertIsNotNone(a.fecha_descargo)
        self.assertEqual(a.docente_reporta, self.docente)
        self.assertEqual((b.tipo, b.subtipo, b.tipo_falta), ('ACADEMICA', 'POSITIVA', ''))
        self.assertEqual(Notificacion.objects.filter(tipo='OBSERVADOR').count(), 2)

    def test_con_un_error_no_se_guarda_nada(self):
        manana = (timezone.localdate() + datetime.timedelta(days=2)).strftime('%d/%m/%Y')
        filas = [
            ['1110555', '', '10/03/2026', 'Comportamental', 'A mejorar', 'II', '', 'Bien', ''],
            ['999', '', '10/03/2026', 'Comportamental', '', '', '', 'No es del curso', ''],
            ['1110555', '', manana, 'Comportamental', '', '', '', 'Futuro', ''],
            ['1110555', '', 'ayer', 'Comportamental', '', '', '', 'Sin fecha', ''],
            ['1110555', '', '10/03/2026', 'Comportamental', '', '', 4321, '', ''],
            ['1110555', '', '10/03/2026', '', '', '', '', 'Sin clase', ''],
        ]
        registros, errores = oc.importar_observaciones(self.a, self.docente, _excel(filas), self.estudiantes)
        self.assertEqual(registros, [])
        self.assertEqual(len(errores), 5, errores)
        self.assertEqual(RegistroObservador.objects.count(), 0)

    def test_solo_cursos_de_su_sede(self):
        otro = self.crear_estudiante('zoe', colegio=self.b)
        registros, errores = oc.importar_observaciones(
            self.a, self.docente, _excel([['zoe', '', '10/03/2026', 'Académica', '', '', '', 'x', '']]), self.estudiantes)
        self.assertEqual(len(errores), 1)
        self.assertFalse(RegistroObservador.objects.filter(estudiante=otro).exists())
