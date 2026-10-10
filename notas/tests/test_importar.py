"""Importaciones masivas: plantillas, vista previa, actualizar sin duplicar y errores claros."""
import io

import openpyxl
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile

from notas import importar as imp
from notas.models import AsignacionDocente, Curso, Docente, Estudiante, FichaDocente, FichaEstudiante, Materia, Sede

from .base import ColegioDePrueba


def xlsx(titulos, filas, hoja='Datos'):
    libro = openpyxl.Workbook()
    ws = libro.active
    ws.title = hoja
    ws.append(titulos)
    for f in filas:
        ws.append(f)
    salida = io.BytesIO()
    libro.save(salida)
    return SimpleUploadedFile('datos.xlsx', salida.getvalue())


def csv(texto, nombre='datos.csv', codificacion='utf-8'):
    return SimpleUploadedFile(nombre, texto.encode(codificacion))


class Importar(ColegioDePrueba):

    def setUp(self):
        self.c = self.cliente(self.rectora)

    def subir(self, tipo, archivo):
        return self.c.post(f'/admin/importacion/{tipo}/revisar/', {'archivo': archivo})

    def confirmar(self, tipo):
        return self.c.post(f'/admin/importacion/{tipo}/aplicar/')

    # --- página y plantillas ---------------------------------------------------
    def test_pagina_y_plantillas(self):
        r = self.c.get('/admin/importacion/')
        self.assertContains(r, 'Asignación académica')
        for tipo in imp.ORDEN:
            for extra in ('', '?datos=1'):
                r = self.c.get(f'/admin/importacion/{tipo}/plantilla/{extra}')
                libro = openpyxl.load_workbook(io.BytesIO(r.content))
                self.assertEqual(libro.sheetnames, ['Instrucciones', 'Datos', 'Listas'])
        # La de estudiantes con datos trae a los tres del 601 y vuelve a entrar sin cambios.
        r = self.c.get('/admin/importacion/estudiantes/plantilla/?datos=1')
        archivo = SimpleUploadedFile('actual.xlsx', r.content)
        r = self.subir('estudiantes', archivo)
        self.assertEqual(r.context['resumen']['igual'], 3)

    def test_solo_administradores(self):
        self.assertEqual(self.cliente(self.u_docente).get('/admin/importacion/').status_code, 403)

    # --- cursos ---------------------------------------------------------------
    def test_cursos(self):
        Sede.objects.create(colegio=self.a, nombre='Sede Rural')
        FichaDocente.objects.create(docente=self.docente, numero_documento='1105123456')
        r = self.subir('cursos', xlsx(['Grado *', 'Subgrupo', 'Sede', 'Director de grado'], [
            ['Séptimo', '01', 'sede rural', '1.105.123.456'], ['7', '02', '', ''],
            ['6', '1', '', 'Pedro Profe'], ['doce', '', '', ''], ['8', '', 'Sede Luna', '']]))
        res = {x['resumen']: x for x in r.context['resultados']}
        self.assertEqual(res['701']['accion'], 'crear')
        self.assertEqual(res['601']['accion'], 'igual')                      # ya existía con ese grado y director
        self.assertEqual(r.context['resumen']['error'], 2)
        self.confirmar('cursos')
        c701 = Curso.objects.get(colegio=self.a, nombre='701')
        self.assertEqual((c701.grado, c701.sede.nombre, c701.director_grado), (7, 'Sede Rural', self.docente))
        self.assertTrue(Curso.objects.filter(colegio=self.a, nombre='702').exists())

    # --- docentes -------------------------------------------------------------
    def test_docentes_sin_duplicar(self):
        archivo = lambda: csv('Nombres;Apellidos;Cédula;Correo\nLaura;Martínez Gómez;1.105.000.111;laura@x.co\n'
                              'Juan;Pérez;22333444;correo-malo\n', codificacion='cp1252')
        r = self.subir('docentes', archivo())
        self.assertEqual((r.context['resumen']['crear'], r.context['resumen']['error']), (1, 1))
        r = self.confirmar('docentes')
        self.assertEqual(r.context['credenciales'][0]['usuario'], 'laura.martinez')
        u = User.objects.get(username='laura.martinez')
        self.assertTrue(u.check_password('laura.martinez'))
        self.assertEqual(u.docente.ficha.numero_documento, '1105000111')
        # Otra vez el mismo archivo: queda igual, no se duplica.
        r = self.subir('docentes', archivo())
        self.assertEqual(r.context['resumen']['igual'], 1)
        self.assertEqual(Docente.objects.filter(colegio=self.a).count(), 2)
        # Descargar la lista de usuarios nuevos
        self.assertEqual(self.c.get('/admin/importacion-credenciales/').status_code, 200)

    def test_documento_de_otro_colegio(self):
        otro = User.objects.create_user('otro', password='x')
        d = Docente.objects.create(colegio=self.b, user=otro)
        FichaDocente.objects.create(docente=d, numero_documento='999')
        r = self.subir('docentes', csv('Nombres,Apellidos,Documento\nAna,Ruiz,999\n'))
        self.assertIn('otro colegio', r.context['resultados'][0]['errores'][0])

    # --- materias y asignación -----------------------------------------------
    def test_materias_y_asignacion(self):
        self.subir('materias', xlsx(['Asignatura', 'Área', 'Sigla'], [['Física', 'Ciencias Naturales', 'FIS'],
                                                                     ['Matemáticas', 'Matemáticas', 'MAT']]))
        self.confirmar('materias')
        fisica = Materia.objects.get(colegio=self.a, nombre='FÍSICA')
        Curso.objects.create(colegio=self.a, nombre='602', grado=6, subgrupo='2')
        FichaDocente.objects.create(docente=self.docente, numero_documento='1105123456')
        otra = User.objects.create_user('marta.ruiz', password='x', first_name='Marta', last_name='Ruiz')
        marta = Docente.objects.create(colegio=self.a, user=otra)
        r = self.subir('asignacion', xlsx(['Docente', 'Materia', 'Cursos', 'IH'], [
            ['1105123456', 'física', '601, 602', 3],
            ['marta.ruiz', 'Matemáticas', '601', 5],            # Matemáticas 601 era de Pedro: cambia de docente
            ['Pedro Profe', 'Química', '601', 2]]))
        acciones = [(x['resumen'].upper(), x['accion']) for x in r.context['resultados']]
        self.assertIn(('FÍSICA · 601 · PEDRO PROFE', 'crear'), acciones)
        self.assertIn(('FÍSICA · 602 · PEDRO PROFE', 'crear'), acciones)
        self.assertIn(('MATEMÁTICAS · 601 · MARTA RUIZ', 'actualizar'), acciones)
        self.assertEqual(r.context['resumen']['error'], 1)          # Química no existe
        self.confirmar('asignacion')
        self.assertEqual(AsignacionDocente.objects.filter(colegio=self.a, materia=fisica).count(), 2)
        self.asig.refresh_from_db()
        self.assertEqual((self.asig.docente, self.asig.intensidad_horaria_semanal), (marta, 5))

    # --- estudiantes ----------------------------------------------------------
    def test_estudiantes_crear_actualizar_y_ficha(self):
        ana = self.estudiantes[0]
        FichaEstudiante.objects.create(estudiante=ana, numero_documento='1001')
        Curso.objects.create(colegio=self.a, nombre='701', grado=7)
        r = self.subir('estudiantes', xlsx(
            ['Apellidos', 'Nombres', 'Curso', 'Tipo doc', 'Documento', 'Fecha de nacimiento', 'RH', 'Acudiente', 'Celular acudiente'],
            [['Prueba', 'Ana', '701', 'TI', '1001', '12/04/2014', 'O positivo', 'Rosa Prueba', '3105556677'],
             ['Zapata', 'Luis', '601', 'Registro civil', '2002', '2015-01-30', 'B+', '', ''],
             ['Ríos', 'Eva', '901', '', '', '', '', '', ''],
             ['Zapata', 'Luis', '601', 'TI', '2002', '', '', '', '']]))
        res = r.context['resultados']
        self.assertEqual([x['accion'] for x in res], ['actualizar', 'crear', 'error', 'error'])
        self.assertIn('curso: 601 → 701', res[0]['cambios'])
        self.assertIn('fila 3', res[3]['errores'][0])
        self.confirmar('estudiantes')
        ana.refresh_from_db()
        self.assertEqual(ana.curso.nombre, '701')
        self.assertEqual((ana.ficha.grupo_sanguineo, ana.ficha.celular_acudiente), ('O+', '3105556677'))
        luis = Estudiante.objects.get(ficha__numero_documento='2002')
        self.assertEqual((luis.user.username, luis.ficha.tipo_documento, str(luis.ficha.fecha_nacimiento)),
                         ('luis.zapata', 'RC', '2015-01-30'))

    def test_archivo_sin_columnas_obligatorias(self):
        r = self.subir('estudiantes', xlsx(['Nombre completo', 'Edad'], [['X', 3]]))
        self.assertRedirects(r, '/admin/importacion/', fetch_redirect_response=False)
        r = self.c.get('/admin/importacion/')
        self.assertContains(r, 'No se encontraron las columnas obligatorias')

    def test_no_se_aplica_sin_revisar(self):
        r = self.confirmar('cursos')
        self.assertRedirects(r, '/admin/importacion/', fetch_redirect_response=False)

    def test_grados(self):
        for v, g in [('6', 6), ('6°', 6), ('Sexto', 6), ('transición', 0), ('Undécimo', 11), ('once', 11),
                     ('Grado 3', 3), ('12', None), ('', None), ('Jardín', -1)]:
            self.assertEqual(imp.parse_grado(v), g, v)


class ImportacionGrande(ColegioDePrueba):
    """500 estudiantes de una vez no pueden pasar el límite de tiempo del servidor."""

    def test_quinientos_estudiantes_rapido_y_la_clave_se_actualiza_al_entrar(self):
        import time
        from django.contrib.auth.hashers import identify_hasher
        from django.test import override_settings
        produccion = ['django.contrib.auth.hashers.PBKDF2PasswordHasher',          # como config/settings.py
                      'django.contrib.auth.hashers.PBKDF2SHA1PasswordHasher',
                      'notas.hashers.PBKDF2InicialHasher']
        filas = [[f'Nombre{i}', f'Apellido{i} Prueba', '601', 'TI', str(1100000000 + i)] for i in range(500)]
        c = self.cliente(self.rectora)
        with override_settings(PASSWORD_HASHERS=produccion):
            c.post('/admin/importacion/estudiantes/revisar/', {'archivo': xlsx(
                ['Nombres *', 'Apellidos *', 'Curso *', 'Tipo de documento', 'Documento'], filas)})
            inicio = time.time()
            r = c.post('/admin/importacion/estudiantes/aplicar/')
            duracion = time.time() - inicio
            self.assertEqual(r.context['resumen']['crear'], 500)
            self.assertLess(duracion, 60, f'tardó {duracion:.1f} s')
            u = User.objects.get(username='nombre0.apellido0')
            self.assertEqual(identify_hasher(u.password).algorithm, 'pbkdf2_inicial')
            # Entra con su usuario como clave y la contraseña queda con el método normal.
            entra = self.cliente(host='a.localhost')
            self.assertTrue(entra.login(username='nombre0.apellido0', password='nombre0.apellido0'))
            u.refresh_from_db()
            self.assertEqual(identify_hasher(u.password).algorithm, 'pbkdf2_sha256')
