"""Pendientes de octubre: línea 5 del encabezado, Tahoma y planillas atadas a la escala del colegio."""
import io
import json
from decimal import Decimal

import openpyxl
from django.template.loader import render_to_string

from notas.models import Calificacion, EscalaValoracion
from notas.planillas import pdf as pdf_mod
from notas.planillas.excel import a_sesion, aplicar, generar_libro, leer_libro
from notas.planillas.guardar import guardar_estudiante, rango

from .base import ColegioDePrueba


class Encabezado(ColegioDePrueba):

    def test_linea_5_y_tahoma(self):
        self.a.linea_encabezado_5 = 'Resolución de aprobación 0123 de 2026'
        self.a.linea_encabezado_5_fuente = 'Tahoma'
        self.a.save()
        html = render_to_string('notas/fragmentos/encabezado_pdf.html', {'colegio': self.a})
        self.assertIn('Resolución de aprobación 0123 de 2026', html)
        self.assertIn("font-family: 'Tahoma'", html)
        from notas.models import Colegio
        self.assertIn('Tahoma', dict(Colegio.FONT_CHOICES))


class EscalaDeUnoADiez(ColegioDePrueba):
    """Un colegio que califica de 1 a 10: la planilla no puede botar un 7."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        EscalaValoracion.objects.filter(colegio=cls.a).delete()
        for nombre, mi, ma in [('Bajo', '1.0', '5.9'), ('Básico', '6.0', '7.9'),
                               ('Alto', '8.0', '8.9'), ('Superior', '9.0', '10.0')]:
            EscalaValoracion.objects.create(colegio=cls.a, nombre_desempeno=nombre,
                                            valor_minimo=Decimal(mi), valor_maximo=Decimal(ma))

    def test_rango(self):
        self.assertEqual(rango(self.a), (Decimal('1.0'), Decimal('10.0')))

    def test_guardar_respeta_la_escala(self):
        ana = self.estudiantes[0]
        final = guardar_estudiante(self.a, self.asig, self.p1, ana,
                                   {'SABER': [('N1', '7'), ('N2', '9'), ('N3', '11')], 'HACER': [('N1', '8')]})
        # SABER = (7+9)/2 = 8 (el 11 no cuenta) → 8*0.6 + 8*0.4 = 8.0
        self.assertEqual(final, Decimal('8.00'))
        self.assertEqual(Calificacion.objects.get(estudiante=ana, tipo_nota='SABER').notas_detalladas.count(), 2)

    def test_pdf_desempeno_y_rojo_segun_la_escala(self):
        ana, beto = self.estudiantes[:2]
        guardar_estudiante(self.a, self.asig, self.p1, ana, {'SABER': [('N1', '9.5')], 'HACER': [('N1', '9')]})
        guardar_estudiante(self.a, self.asig, self.p1, beto, {'SABER': [('N1', '5')], 'HACER': [('N1', '5')]})
        filas = {f['nombre']: f for f in pdf_mod.datos_planilla(self.a, self.p1, self.asig)['filas']}
        self.assertEqual(filas['PRUEBA ANA']['desempeno'], 'SUPERIOR')
        self.assertFalse(filas['PRUEBA ANA']['pierde'])
        self.assertEqual(filas['PRUEBA BETO']['desempeno'], 'BAJO')
        self.assertTrue(filas['PRUEBA BETO']['pierde'])     # lo decide la nota de aprobación (6,0), no el nombre

    def test_excel_acepta_7_y_rechaza_11(self):
        libro = openpyxl.load_workbook(io.BytesIO(generar_libro(self.a, self.p1, [self.asig])))
        ws = libro[libro.sheetnames[1]]
        meta = next(json.loads(c.value) for c in ws[4] if isinstance(c.value, str) and 'mcolegio' in c.value)
        # La validación del Excel ya viene con la escala del colegio.
        rangos = [(dv.formula1, dv.formula2) for dv in ws.data_validations.dataValidation if dv.type == 'decimal']
        self.assertIn(('1.0', '10.0'), rangos)
        ws.cell(meta['fila'], meta['comp']['SABER'][0], 7)
        ws.cell(meta['fila'] + 1, meta['comp']['SABER'][0], 11)
        salida = io.BytesIO()
        libro.save(salida)
        salida.seek(0)
        resultado = leer_libro(salida, self.a, self.u_docente)
        hoja = resultado['hojas'][0]
        self.assertTrue(any('11' in e and '10' in e for e in (str(x) for x in hoja['errores'])), hoja['errores'])
        aplicar(a_sesion(resultado), self.a, self.u_docente)
        ana = self.estudiantes[0]
        self.assertEqual(Calificacion.objects.get(estudiante=ana, tipo_nota='SABER').valor_nota, Decimal('7.00'))

    def test_la_pagina_en_linea_recibe_la_escala(self):
        r = self.cliente(self.u_docente).get(
            f'/docente/ingresar-notas/?asignacion_id={self.asig.id}&periodo_id={self.p1.id}')
        self.assertContains(r, 'SUPERIOR')
        self.assertContains(r, '10.0')


class AgregarVariasPreguntas(ColegioDePrueba):

    def test_pregunta_cuantas_y_las_agrega(self):
        from puntoexacto.models import Examen
        c = self.cliente(self.u_docente)
        c.post('/puntoexacto/nuevo/', {
            'titulo': 'Parcial', 'fecha': '2026-10-06', 'componente': 'SABER', 'numero_preguntas': '5',
            'numero_opciones': '4', 'hojas_por_pagina': '1', 'metodo': 'con_piso', 'nota_todo_mal': '1',
            'nota_nada_marcado': '1', 'opciones_por_pregunta': '{}'})
        ex = Examen.objects.get()
        self.assertContains(c.get(f'/puntoexacto/{ex.id}/clave/'), '¿Cuántas preguntas agregar?')
        c.post(f'/puntoexacto/{ex.id}/clave/', {'metodo': 'con_piso', 'accion': 'agregar', 'cantidad_agregar': '10'})
        ex.refresh_from_db()
        self.assertEqual(ex.numero_preguntas, 15)
        self.assertEqual(ex.preguntas.count(), 15)
        # Nunca más de 150 ni menos de 1.
        c.post(f'/puntoexacto/{ex.id}/clave/', {'metodo': 'con_piso', 'accion': 'agregar', 'cantidad_agregar': '900'})
        ex.refresh_from_db()
        self.assertEqual(ex.numero_preguntas, 150)
