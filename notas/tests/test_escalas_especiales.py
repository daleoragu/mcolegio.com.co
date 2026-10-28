"""Escalas especiales: una materia (o una materia en ciertos grados) con otra escala de valoración."""
from decimal import Decimal

from notas import escalas
from notas.models import (AreaConocimiento, AsignacionDocente, Calificacion, Curso, EscalaEspecial, Materia,
                          NivelEscalaEspecial, PonderacionAreaMateria)

from .base import ColegioDePrueba

# Inglés: el BAJO va hasta 3,4.
NIVELES_INGLES = [('BAJO', '1.0', '3.4'), ('BÁSICO', '3.5', '3.9'), ('ALTO', '4.0', '4.5'), ('SUPERIOR', '4.6', '5.0')]
# Matemáticas solo en 10.°: el BAJO va hasta 3,2.
NIVELES_MAT10 = [('BAJO', '1.0', '3.2'), ('BÁSICO', '3.3', '3.9'), ('ALTO', '4.0', '4.5'), ('SUPERIOR', '4.6', '5.0')]


def crear_escala(colegio, nombre, materias, niveles, grados=()):
    e = EscalaEspecial.objects.create(colegio=colegio, nombre=nombre, grados=list(grados))
    e.materias.set(materias)
    for n, a, b in niveles:
        NivelEscalaEspecial.objects.create(escala=e, nombre_desempeno=n, valor_minimo=Decimal(a),
                                           valor_maximo=Decimal(b))
    return e


class EscalasEspeciales(ColegioDePrueba):

    def setUp(self):
        self.ingles = Materia.objects.create(colegio=self.a, nombre='Inglés')
        self.area_hum = AreaConocimiento.objects.create(colegio=self.a, nombre='Humanidades')
        self.area_mat = AreaConocimiento.objects.create(colegio=self.a, nombre='Matemáticas')
        PonderacionAreaMateria.objects.create(colegio=self.a, area=self.area_hum, materia=self.ingles, peso_porcentual=100)
        PonderacionAreaMateria.objects.create(colegio=self.a, area=self.area_mat, materia=self.materia, peso_porcentual=100)
        self.asig_ing = AsignacionDocente.objects.create(colegio=self.a, docente=self.docente, materia=self.ingles,
                                                         curso=self.curso)
        self.c1001 = Curso.objects.create(colegio=self.a, nombre='1001', grado=10, subgrupo='01')
        self.asig_mat10 = AsignacionDocente.objects.create(colegio=self.a, docente=self.docente, materia=self.materia,
                                                           curso=self.c1001)
        self.esc_ing = crear_escala(self.a, 'Escala de Inglés', [self.ingles], NIVELES_INGLES)
        self.esc_mat = crear_escala(self.a, 'Matemáticas media', [self.materia], NIVELES_MAT10, grados=[10])

    def nota(self, est, materia, valor, periodo=None):
        Calificacion.objects.create(colegio=self.a, estudiante=est, materia=materia, periodo=periodo or self.p1,
                                    tipo_nota='PROM_PERIODO', valor_nota=Decimal(valor))

    def test_que_escala_rige(self):
        esc = escalas.Escalas(self.a)
        self.assertTrue(esc.es_especial(self.ingles.id, 6))
        self.assertTrue(esc.es_especial(self.ingles.id, 10))
        self.assertFalse(esc.es_especial(self.materia.id, 6))      # Matemáticas en 6.°: la del colegio
        self.assertTrue(esc.es_especial(self.materia.id, 10))
        self.assertEqual(esc.desempeno(Decimal('3.2'), self.ingles.id, 6), 'BAJO')
        self.assertEqual(esc.desempeno(Decimal('3.2'), self.materia.id, 6), 'BÁSICO')
        self.assertEqual(esc.desempeno(Decimal('3.2'), self.materia.id, 10), 'BAJO')
        self.assertEqual(esc.aprobacion(self.ingles.id, 6), Decimal('3.5'))
        self.assertEqual(esc.aprobacion(self.materia.id, 6), Decimal('3.0'))
        self.assertTrue(esc.pierde(Decimal('3.4'), self.ingles.id, 6))
        self.assertFalse(esc.pierde(Decimal('3.4'), self.materia.id, 6))

    def test_el_area_usa_la_especial_solo_si_todas_sus_materias_la_comparten(self):
        esc = escalas.Escalas(self.a)
        self.assertIs(esc.niveles_area([self.ingles.id], 6), esc.niveles(self.ingles.id, 6))
        self.assertIs(esc.niveles_area([self.ingles.id, self.materia.id], 6), esc.general)

    def test_boletin_del_periodo(self):
        from notas.boletin.logic import get_datos_boletin_curso
        ana = self.estudiantes[0]
        self.nota(ana, self.ingles, '3.2')
        self.nota(ana, self.materia, '3.2')
        datos = get_datos_boletin_curso(self.a, self.curso, self.p1, ana)[0]
        materias = {m['nombre']: m for area in datos['areas'] for m in area['materias']}
        self.assertEqual(materias['Inglés']['v_n'], 'BAJO')
        self.assertEqual(materias['Matemáticas']['v_n'], 'BÁSICO')
        areas = {a['nombre']: a for a in datos['areas']}
        self.assertEqual(areas['HUMANIDADES']['desempeno_area'], 'BAJO')
        self.assertEqual(areas['HUMANIDADES']['estado_color_css_class'], 'desempeno-reprobado')
        self.assertEqual(areas['MATEMÁTICAS']['estado_color_css_class'], 'desempeno-aprobado')

    def test_boletin_final_cuenta_areas_perdidas_con_la_escala_de_la_materia(self):
        from notas.boletin.logic import get_datos_boletin_final
        ana = self.estudiantes[0]
        for p in (self.p1, self.p2):
            self.nota(ana, self.ingles, '3.3', p)
            self.nota(ana, self.materia, '3.3', p)
        datos, _ = get_datos_boletin_final(self.a, self.curso, 2026, ana)
        self.assertEqual(datos[0]['areas_reprobadas'], 1)          # solo Inglés
        self.assertEqual(list(datos[0]['detalle_areas_reprobadas']), ['HUMANIDADES'])

    def test_alertas_tempranas(self):
        from notas.alertas import alertas_de_curso
        ana = self.estudiantes[0]
        self.nota(ana, self.ingles, '3.4')
        self.nota(ana, self.materia, '3.4')
        fila = next(f for f in alertas_de_curso(self.a, self.curso, 2026, self.p1) if f['estudiante'] == ana)
        self.assertEqual([p['materia'] for p in fila['perdidas']], ['Inglés'])
        self.assertEqual(fila['perdidas'][0]['falta'], Decimal('0.1'))

    def test_reporte_parcial_y_planillas(self):
        from notas import reporte_aviso
        from notas.planillas.excel import aprobacion_de, escala_de
        self.assertTrue(reporte_aviso.esta_en_bajo(self.a, Decimal('3.3'), self.asig_ing))
        self.assertFalse(reporte_aviso.esta_en_bajo(self.a, Decimal('3.3'), self.asig))
        self.assertEqual(reporte_aviso.desempeno(self.a, Decimal('3.3'), self.asig_ing), 'BAJO')
        self.assertEqual(aprobacion_de(self.a, self.asig_ing), Decimal('3.5'))
        self.assertEqual(aprobacion_de(self.a, self.asig_mat10), Decimal('3.3'))
        self.assertEqual(escala_de(self.a, self.asig_ing)[0][1], Decimal('3.4'))
        self.assertEqual(escala_de(self.a, self.asig)[0][1], Decimal('2.9'))

    def test_planilla_excel_con_dos_escalas(self):
        import io

        import openpyxl

        from notas.planillas.excel import generar_libro
        datos = generar_libro(self.a, self.p1, [self.asig, self.asig_ing])
        libro = openpyxl.load_workbook(io.BytesIO(datos))
        self.assertIn('Escala', libro.sheetnames)
        self.assertIn('Escala2', libro.sheetnames)
        self.assertEqual(libro['Escala2']['B2'].value, 3.4)

    def test_sabana(self):
        from notas.views.sabana_views import _get_sabana_acumulada_data
        ana = self.estudiantes[0]
        self.nota(ana, self.ingles, '3.2')
        self.nota(ana, self.materia, '3.2')
        sabana, areas, *_ = _get_sabana_acumulada_data(self.a, self.curso, self.p1)
        fila = next(s for s in sabana if s['info'] == ana)
        celdas = [c for c in fila['filas_notas'][0]['celdas'] if not c['is_area']]
        nombres = [m.nombre for a in areas for m in a.materias_del_curso_ordenadas]
        pierde = {n: c['pierde'] for n, c in zip(nombres, celdas)}
        self.assertEqual(pierde, {'Inglés': True, 'Matemáticas': False})

    def test_ingreso_de_notas_manda_la_escala_de_la_materia(self):
        import json
        r = self.cliente(self.u_docente).get(
            f'/docente/ingresar-notas/?asignacion_id={self.asig_ing.id}&periodo_id={self.p1.id}')
        self.assertEqual(r.status_code, 200)
        escala = json.loads(r.context['escala_valoracion_json'])
        self.assertEqual(Decimal(str(escala[0]['valor_maximo'])), Decimal('3.4'))


class PantallaEscalasEspeciales(ColegioDePrueba):

    def test_crear_editar_y_validar(self):
        c = self.cliente(self.rectora)
        r = c.get('/admin/escalas-especiales/nueva/')
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'BÁSICO')                           # arranca con la escala del colegio
        datos = {'nombre': 'Escala de Matemáticas', 'materias': [self.materia.id], 'grados': ['10', '11'],
                 'nivel_nombre': ['bajo', 'básico', 'alto', 'superior'],
                 'nivel_min': ['1,0', '3.5', '4.0', '4.6'], 'nivel_max': ['3,4', '3.9', '4.5', '5.0']}
        r = c.post('/admin/escalas-especiales/nueva/', datos)
        self.assertEqual(r.status_code, 302)
        e = EscalaEspecial.objects.get()
        self.assertEqual(e.grados, [10, 11])
        self.assertEqual([n.nombre_desempeno for n in e.niveles.all()], ['BAJO', 'BÁSICO', 'ALTO', 'SUPERIOR'])
        self.assertEqual(e.niveles.first().valor_maximo, Decimal('3.4'))
        self.assertContains(c.get('/admin/escalas-especiales/'), 'Escala de Matemáticas')

        # Otra escala para la misma materia en 11.°: no se puede.
        r = c.post('/admin/escalas-especiales/nueva/', {**datos, 'nombre': 'Otra', 'grados': ['11']})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'ya tiene la escala')
        # Rangos cruzados.
        r = c.post(f'/admin/escalas-especiales/{e.pk}/', {**datos, 'nivel_max': ['3.6', '3.9', '4.5', '5.0']})
        self.assertContains(r, 'se cruzan')
        # Eliminar.
        c.post(f'/admin/escalas-especiales/{e.pk}/eliminar/')
        self.assertFalse(EscalaEspecial.objects.exists())

    def test_solo_administracion(self):
        self.assertEqual(self.cliente(self.u_docente).get('/admin/escalas-especiales/').status_code, 403)
