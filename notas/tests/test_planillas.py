# notas/tests/test_planillas.py
"""Plan de notas, cálculo de la definitiva y el Excel de ida y vuelta."""
import io
from decimal import Decimal

import openpyxl

from notas.models import Calificacion, PlanNotas
from notas.models.comunicaciones import Notificacion
from notas.planillas.columnas import columnas_del_plan, ubicar
from notas.planillas.excel import generar_libro, leer_libro, aplicar, a_sesion
from notas.planillas.guardar import guardar_estudiante

from .base import ColegioDePrueba


def definitiva(est, periodo):
    return Calificacion.objects.get(estudiante=est, periodo=periodo, tipo_nota='PROM_PERIODO').valor_nota


class Calculo(ColegioDePrueba):

    def test_definitiva_igual_que_la_plataforma(self):
        ana = self.estudiantes[0]
        guardar_estudiante(self.a, self.asig, self.p1, ana,
                           {'SABER': [('Nota 1', '4'), ('Nota 3', '2')], 'HACER': [('Nota 1', '5')]})
        self.assertEqual(definitiva(ana, self.p1), Decimal('3.80'))     # 0,6·3 + 0,4·5

    def test_notas_fuera_de_rango_no_cuentan(self):
        beto = self.estudiantes[1]
        guardar_estudiante(self.a, self.asig, self.p1, beto,
                           {'SABER': [('N1', '7'), ('N2', '4'), ('N3', 'aus')], 'HACER': []})
        self.assertEqual(definitiva(beto, self.p1), Decimal('2.40'))    # solo el 4 cuenta

    def test_cada_nota_va_a_la_columna_de_su_nombre(self):
        self.assertEqual(ubicar([('Nota 1', 4), ('Nota 3', 2)], ['Nota 1', 'Nota 2', 'Nota 3']), [4, None, 2])

    def test_sin_plan_salen_las_del_colegio(self):
        self.assertEqual(len(columnas_del_plan(self.asig, self.p1, 'SABER')), 5)


class ExcelIdaYVuelta(ColegioDePrueba):

    def bajar(self):
        return openpyxl.load_workbook(io.BytesIO(generar_libro(self.a, self.p1, [self.asig])))

    def subir(self, libro, usuario=None):
        salida = io.BytesIO()
        libro.save(salida)
        salida.seek(0)
        return leer_libro(salida, self.a, usuario or self.u_docente)

    def test_lo_que_se_sube_es_lo_que_queda(self):
        import json
        libro = self.bajar()
        ws = libro[libro.sheetnames[1]]
        meta = next(json.loads(c.value) for c in ws[4] if isinstance(c.value, str) and 'mcolegio' in c.value)
        fila_ana = meta['fila']                               # ana es la primera por apellido y nombre
        ws.cell(fila_ana, meta['comp']['SABER'][0], 4.5)      # SABER, nota 1
        ws.cell(fila_ana, meta['comp']['HACER'][0], 3)        # HACER, nota 1
        resultado = self.subir(libro)
        hoja = resultado['hojas'][0]
        self.assertEqual(hoja['estado'], 'ok')
        aplicar(a_sesion(resultado), self.a, self.u_docente)
        ana = self.estudiantes[0]
        esperado = next(e['definitiva'] for e in hoja['estudiantes'] if e['id'] == ana.id)
        self.assertEqual(definitiva(ana, self.p1), esperado)

    def test_otro_docente_no_puede_subir_la_planilla(self):
        from django.contrib.auth.models import User
        from notas.models import Docente
        otra = User.objects.create_user('otra', password='x')
        Docente.objects.create(colegio=self.a, user=otra)
        hoja = self.subir(self.bajar(), usuario=otra)['hojas'][0]
        self.assertEqual(hoja['estado'], 'error')

    def test_avisa_estudiantes_que_llegaron_despues(self):
        libro = self.bajar()
        nuevo = self.crear_estudiante('dani')
        hoja = self.subir(libro)['hojas'][0]
        self.assertIn('PRUEBA DANI', hoja['faltan'])

    def test_estudiante_nuevo_avisa_al_docente(self):
        Notificacion.objects.all().delete()       # los de la preparación ya avisaron
        self.crear_estudiante('eva')
        self.crear_estudiante('fer')
        avisos = Notificacion.objects.filter(destinatario=self.u_docente, tipo='PLANILLA')
        self.assertEqual(avisos.count(), 1)
        self.assertTrue(avisos.first().mensaje.startswith('2 estudiantes nuevos'))


class ObservacionDeLaAsignatura(ColegioDePrueba):

    def obs(self, est):
        return Calificacion.objects.get(estudiante=est, periodo=self.p1, tipo_nota='PROM_PERIODO').observacion

    def test_se_guarda_y_es_opcional(self):
        ana, beto, _ = self.estudiantes
        guardar_estudiante(self.a, self.asig, self.p1, ana, {'SABER': [('Nota 1', '4')]},
                           observacion='  Muy   participativa ')
        guardar_estudiante(self.a, self.asig, self.p1, beto, {'SABER': [('Nota 1', '4')]})
        self.assertEqual(self.obs(ana), 'Muy participativa')
        self.assertEqual(self.obs(beto), '')

    def test_subir_el_excel_no_la_borra(self):
        ana = self.estudiantes[0]
        guardar_estudiante(self.a, self.asig, self.p1, ana, {}, observacion='Refuerzo en casa')
        libro = openpyxl.load_workbook(io.BytesIO(generar_libro(self.a, self.p1, [self.asig])))
        salida = io.BytesIO(); libro.save(salida); salida.seek(0)
        aplicar(a_sesion(leer_libro(salida, self.a, self.u_docente)), self.a, self.u_docente)
        self.assertEqual(self.obs(ana), 'Refuerzo en casa')

    def test_la_planilla_en_linea_la_guarda(self):
        import json
        from django.urls import reverse
        ana = self.estudiantes[0]
        c = self.cliente(self.u_docente)
        r = c.post(reverse('notas:ingresar_notas_periodo'), json.dumps({'asignacion_id': self.asig.id, 'periodo_id': self.p1.id,
                                    'estudiantes': [{'id': str(ana.id), 'notas': {'saber': [{'descripcion': 'Nota 1', 'valor': '4'}]},
                                                     'inasistencias': '0', 'observacion': 'Excelente trabajo'}]}),
                   content_type='application/json')
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(self.obs(ana), 'Excelente trabajo')


class BoletinConNotasDelExcel(ColegioDePrueba):

    def test_el_boletin_no_se_cae_con_inclusion_vacia(self):
        """Una definitiva guardada sin indicador de inclusión (NULL) no debe tumbar el boletín."""
        from notas.boletin.logic import get_datos_boletin_curso
        ana = self.estudiantes[0]
        guardar_estudiante(self.a, self.asig, self.p1, ana, {'SABER': [('Nota 1', '4')]})
        Calificacion.objects.filter(estudiante=ana, tipo_nota='PROM_PERIODO').update(observacion_inclusion=None)
        get_datos_boletin_curso(self.a, self.curso, self.p1)     # antes: AttributeError

    def test_guardar_no_deja_inclusion_en_null(self):
        beto = self.estudiantes[1]
        guardar_estudiante(self.a, self.asig, self.p1, beto, {'SABER': [('Nota 1', '4')]})
        self.assertEqual(Calificacion.objects.get(estudiante=beto, tipo_nota='PROM_PERIODO').observacion_inclusion, '')


class PlanillasPDFyAsistencia(ColegioDePrueba):
    """Mis planillas: notas y asistencia, cada una en línea, Excel y PDF."""

    def test_mis_planillas_trae_notas_y_asistencia(self):
        h = self.cliente(self.u_docente).get(f'/docente/planillas/?periodo={self.p1.id}').content.decode()
        a = self.asig.id
        self.assertIn(f'/docente/planillas/{self.p1.id}/pdf/{a}/', h)                       # notas PDF
        self.assertIn(f'asistencia/?asignacion_id={a}', h)                                    # asistencia en línea
        self.assertIn(f'reportes/asistencia/excel/?asignacion_id={a}', h)                     # asistencia Excel
        self.assertIn(f'reportes/asistencia/pdf/?asignacion_id={a}', h)                       # asistencia PDF
        self.assertIn(f'/docente/planillas/{self.p1.id}/pdf/', h)                             # todas en PDF

    def test_datos_del_pdf_de_notas(self):
        from decimal import Decimal
        from notas.planillas import pdf as pdf_mod
        est = self.estudiantes[0]
        guardar_estudiante(self.a, self.asig, self.p1, est, {'SABER': [('Taller', '4'), ('Quiz', '3')],
                                                             'HACER': [('Proyecto', '5')]})
        d = pdf_mod.datos_planilla(self.a, self.p1, self.asig)
        fila = next(f for f in d['filas'] if f['nombre'].startswith('PRUEBA ANA'))
        # SABER 60 % (prom. 3.5) + HACER 40 % (5.0) = 4.1
        self.assertEqual(fila['final'], Decimal('4.10'))
        self.assertEqual(fila['desempeno'], 'ALTO')
        self.assertEqual([c['nombre'] for c in d['componentes']][:2], ['SABER', 'HACER'][:2])

    def test_pdf_de_otro_docente_no(self):
        from notas.models import Docente
        from django.contrib.auth.models import User
        otro = User.objects.create_user('otro', password='x')
        Docente.objects.create(colegio=self.a, user=otro)
        r = self.cliente(otro).get(f'/docente/planillas/{self.p1.id}/pdf/{self.asig.id}/')
        self.assertEqual(r.status_code, 403)
