# notas/tests/test_sabana_excel.py
"""La sábana se descarga en Excel (antes fallaba con error 500)."""
import io

import openpyxl

from notas.tests.base import ColegioDePrueba


class SabanaEnExcel(ColegioDePrueba):
    def test_descarga_con_el_encabezado_del_colegio(self):
        self.a.linea_encabezado_2 = 'Aprobada por Resolución 0456 de 2019'
        self.a.save()
        c = self.cliente(self.rectora)
        r = c.get('/docente/exportar-sabana-excel/', {'curso_id': self.curso.id, 'periodo_id': self.p2.id})
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r['Content-Type'])
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        textos = [str(c.value) for fila in ws.iter_rows(max_row=8) for c in fila if c.value]
        self.assertIn('COLEGIO A', textos)
        self.assertTrue(any('Resolución 0456' in t for t in textos), textos)
        self.assertTrue(any('601' in t for t in textos))

    def test_se_imprime_horizontal_y_a_lo_ancho(self):
        c = self.cliente(self.rectora)
        r = c.get('/docente/exportar-sabana-excel/', {'curso_id': self.curso.id, 'periodo_id': self.p2.id})
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws.page_setup.orientation, 'landscape')
        self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(ws.page_setup.fitToWidth, 1)
