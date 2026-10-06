# notas/tests/test_asistencia_excel.py
"""Asistencia en Excel: plantilla con lo registrado, subida con revisión."""
import datetime
import io

import openpyxl
from django.contrib.auth.models import User
from django.urls import reverse

from notas.models import Asistencia, Docente
from notas.planillas import asistencia as planilla
from notas.reportes.excel_generator import AsistenciaExcelGenerator

from .base import CLAVE, ColegioDePrueba

FEB_3 = datetime.date(2026, 2, 3)      # martes del primer periodo
MAR_4 = datetime.date(2026, 3, 4)      # miércoles, en la segunda hoja


class AsistenciaExcel(ColegioDePrueba):

    def libro(self):
        wb = AsistenciaExcelGenerator(colegio=self.a).generate_report(self.asig, self.p1, 'todos')
        salida = io.BytesIO(); wb.save(salida); salida.seek(0)
        return openpyxl.load_workbook(salida)

    def celda(self, wb, fecha, estudiante):
        for ws in wb.worksheets:
            for col in range(3, ws.max_column + 1):
                if str(ws.cell(9, col).value or '').endswith(fecha.isoformat()):
                    for fila in range(12, ws.max_row + 1):
                        if ws.cell(fila, 2).value == estudiante.id:
                            return ws.cell(fila, col)
        raise AssertionError('no está la casilla')

    def subir(self, wb, usuario=None):
        salida = io.BytesIO(); wb.save(salida); salida.seek(0)
        return planilla.leer(salida, self.a, usuario or self.u_docente)

    def test_la_plantilla_trae_lo_registrado_y_un_mes_por_hoja(self):
        ana = self.estudiantes[0]
        Asistencia.objects.create(colegio=self.a, estudiante=ana, asignacion=self.asig, fecha=FEB_3, estado='A')
        wb = self.libro()
        self.assertGreaterEqual(len(wb.worksheets), 2)
        self.assertEqual(self.celda(wb, FEB_3, ana).value, 'X')
        self.assertEqual(wb.worksheets[0].cell(9, 1).value, planilla.MARCA)

    def test_lee_todos_los_meses_y_quita_fallas_borradas(self):
        ana, beto, _ = self.estudiantes
        Asistencia.objects.create(colegio=self.a, estudiante=ana, asignacion=self.asig, fecha=FEB_3, estado='A')
        wb = self.libro()
        self.celda(wb, FEB_3, ana).value = None          # la falla de febrero era un error
        self.celda(wb, MAR_4, beto).value = 'aj'         # marzo: segunda hoja
        r = self.subir(wb)
        self.assertEqual(r['errores'], [])
        self.assertEqual(r['resumen'], {'marcadas': 1, 'quitadas': 1})
        planilla.aplicar(r['cambios'], self.a, self.u_docente)
        self.assertEqual(Asistencia.objects.get(estudiante=ana, fecha=FEB_3).estado, 'P')
        b = Asistencia.objects.get(estudiante=beto, fecha=MAR_4)
        self.assertEqual((b.estado, b.justificada), ('A', True))

    def test_plantilla_vieja_no_borra_lo_registrado_en_linea(self):
        ana = self.estudiantes[0]
        Asistencia.objects.create(colegio=self.a, estudiante=ana, asignacion=self.asig, fecha=FEB_3, estado='A')
        wb = self.libro()
        for ws in wb.worksheets:
            ws.cell(9, 1).value = None                   # como las descargadas antes del cambio
        self.celda(wb, FEB_3, ana).value = None
        r = self.subir(wb)
        self.assertEqual(r['cambios'], [])
        self.assertTrue(any('plantilla anterior' in a for a in r['avisos']))

    def test_otro_docente_no_puede_subirla(self):
        otra = User.objects.create_user('otra', password=CLAVE)
        Docente.objects.create(colegio=self.a, user=otra)
        r = self.subir(self.libro(), usuario=otra)
        self.assertTrue(r['errores'])
        self.assertEqual(r['cambios'], [])

    def test_codigo_invalido_se_avisa(self):
        wb = self.libro()
        self.celda(wb, FEB_3, self.estudiantes[0]).value = 'Z'
        r = self.subir(wb)
        self.assertTrue(any('«Z» no es válido' in e for e in r['errores']))

    def test_flujo_en_la_pagina_revisar_y_confirmar(self):
        beto = self.estudiantes[1]
        wb = self.libro()
        self.celda(wb, FEB_3, beto).value = 'T'
        archivo = io.BytesIO(); wb.save(archivo); archivo.seek(0); archivo.name = 'asistencia.xlsx'
        c = self.cliente(self.u_docente)
        url = reverse('notas:importar_asistencia_excel')
        r = c.post(url, {'accion': 'revisar', 'archivo_excel': archivo})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'Guardar en la plataforma')
        self.assertFalse(Asistencia.objects.exists())   # todavía nada
        r = c.post(url, {'accion': 'confirmar'})
        self.assertRedirects(r, reverse('notas:consulta_asistencia'), fetch_redirect_response=False)
        self.assertEqual(Asistencia.objects.get(estudiante=beto, fecha=FEB_3).estado, 'T')

    def test_otro_docente_no_descarga_la_plantilla(self):
        otra = User.objects.create_user('otra2', password=CLAVE)
        Docente.objects.create(colegio=self.a, user=otra)
        r = self.cliente(otra).get(reverse('notas:generar_reporte_individual_excel'),
                                   {'asignacion_id': self.asig.id, 'periodo_id': self.p1.id})
        self.assertEqual(r.status_code, 403)
