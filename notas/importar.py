# notas/importar.py
"""Importaciones masivas: cursos, docentes, materias, asignación académica y estudiantes.

El camino es siempre el mismo:
  1. Se descarga la plantilla (vacía o con lo que ya hay en la plataforma).
  2. Se sube el Excel (o CSV). Los títulos de las columnas se reconocen aunque
     vengan con tildes, mayúsculas o con otro nombre conocido («Cédula» =
     «Documento»), y en cualquier orden.
  3. Vista previa: fila por fila, qué se crea, qué se actualiza, qué queda igual
     y qué tiene un error y por qué. Todavía no se guarda nada.
  4. Al confirmar se vuelve a revisar contra la base (por si algo cambió) y se
     guarda todo en una sola transacción.

Volver a subir el mismo archivo no duplica nada: cada fila se busca primero por
su llave (documento del docente o del estudiante, nombre del curso, nombre de la
materia, materia + curso de la asignación) y, si ya existe, se actualiza.

Los usuarios nuevos se crean igual que a mano: usuario «nombre.apellido» y
contraseña inicial igual al usuario. Al final se pueden descargar.
"""
import csv
import datetime
import io
import re
import unicodedata
from collections import namedtuple
from decimal import Decimal, InvalidOperation

from django.contrib.auth.models import Group, User
from django.core.validators import validate_email
from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils.text import slugify

Col = namedtuple('Col', 'clave titulo obligatoria sinonimos ayuda ejemplo')

MAX_FILAS = 3000


def norm(texto):
    """'Número de Documento ' -> 'numero de documento'."""
    t = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode().lower()
    return ' '.join(re.sub(r'[^a-z0-9]+', ' ', t).split())


# ---------------------------------------------------------------------------
# Las columnas de cada importación
# ---------------------------------------------------------------------------

TIPOS = {
    'cursos': {
        'titulo': 'Cursos', 'icono': 'fa-door-open',
        'descripcion': 'Grados y grupos del colegio, con su sede y director de grado.',
        'columnas': [
            Col('grado', 'Grado', True, ['nivel grado', 'grado escolar'],
                'Número o nombre del grado: 0 o Transición, 1 a 11, Sexto, Undécimo…', '6'),
            Col('subgrupo', 'Subgrupo', False, ['grupo', 'seccion', 'paralelo'],
                'Solo si el grado tiene varios cursos: 01, 02… o A, B…', '01'),
            Col('nombre', 'Nombre del curso', False, ['curso', 'nombre'],
                'Vacío = la plataforma lo arma con el grado y el subgrupo, como los demás cursos.', '601'),
            Col('sede', 'Sede', False, ['nombre sede'], 'Nombre de una sede ya creada.', 'Sede Principal'),
            Col('director', 'Director de grado', False, ['director', 'director de grupo', 'documento director'],
                'Documento, usuario, correo o nombre completo de un docente ya creado.', '1105123456'),
        ],
    },
    'docentes': {
        'titulo': 'Docentes', 'icono': 'fa-chalkboard-teacher',
        'descripcion': 'Cuentas de los docentes. Se reconocen por el documento.',
        'columnas': [
            Col('nombres', 'Nombres', True, ['nombre', 'primer nombre'], '', 'Laura Marcela'),
            Col('apellidos', 'Apellidos', True, ['apellido'], '', 'Martínez Gómez'),
            Col('documento', 'Documento', True, ['cedula', 'numero de documento', 'no documento', 'identificacion', 'cc'],
                'Con él se reconoce al docente si se vuelve a subir el archivo.', '1105123456'),
            Col('correo', 'Correo', False, ['email', 'correo electronico', 'e mail'],
                'Sirve para recuperar la contraseña.', 'laura@colegio.edu.co'),
            Col('telefono', 'Teléfono', False, ['celular', 'telefono celular', 'movil'], '', '3105556677'),
            Col('titulo', 'Título profesional', False, ['titulo', 'profesion'], '', 'Licenciada en Matemáticas'),
        ],
    },
    'materias': {
        'titulo': 'Materias', 'icono': 'fa-book-open',
        'descripcion': 'Asignaturas y el área a la que pertenecen.',
        'columnas': [
            Col('nombre', 'Materia', True, ['asignatura', 'nombre materia', 'nombre asignatura', 'nombre'], '', 'Matemáticas'),
            Col('area', 'Área', True, ['area de conocimiento', 'nombre area'], '', 'Matemáticas'),
            Col('abreviatura', 'Abreviatura', False, ['sigla', 'abrev'], 'Hasta 10 letras.', 'MAT'),
        ],
    },
    'ponderacion': {
        'titulo': 'Ponderación de áreas', 'icono': 'fa-balance-scale',
        'descripcion': 'Cuánto pesa cada materia en su área: general o distinto por grado (opcional).',
        'columnas': [
            Col('area', 'Área', True, ['area de conocimiento', 'nombre area'], 'Un área ya creada.',
                'Ciencias Naturales y Educación Ambiental'),
            Col('materia', 'Materia', True, ['asignatura'], 'Una materia de esa área.', 'Física'),
            Col('grado', 'Grado', False, ['grados', 'nivel'],
                'Vacío = general (todos los grados). Un grado (6, Décimo…), varios (10, 11) o un nivel: '
                'Preescolar, Primaria, Secundaria, Media.', '10, 11'),
            Col('peso', 'Peso (%)', True, ['porcentaje', 'ponderacion', 'peso porcentual', 'peso'],
                'De 0 a 100. Los de una misma área y grado deben sumar 100.', '40'),
        ],
    },
    'asignacion': {
        'titulo': 'Asignación académica', 'icono': 'fa-sitemap',
        'descripcion': 'Qué docente dicta qué materia en qué curso, con su intensidad horaria.',
        'columnas': [
            Col('docente', 'Docente', True, ['documento docente', 'profesor', 'documento'],
                'Documento, usuario, correo o nombre completo de un docente ya creado.', '1105123456'),
            Col('materia', 'Materia', True, ['asignatura'], 'Nombre de una materia ya creada.', 'Matemáticas'),
            Col('curso', 'Curso(s)', True, ['cursos', 'grupo', 'grupos'],
                'Uno o varios separados por coma: 601, 602, 603.', '601, 602'),
            Col('ih', 'Intensidad horaria', False, ['ih', 'horas', 'intensidad', 'horas semanales', 'intensidad horaria semanal'],
                'Horas por semana.', '5'),
        ],
    },
    'estudiantes': {
        'titulo': 'Estudiantes', 'icono': 'fa-users',
        'descripcion': 'Matrícula: cuentas de los estudiantes, su curso y su ficha. Se reconocen por el documento.',
        'columnas': [
            Col('nombres', 'Nombres', True, ['nombre', 'primer nombre'], '', 'Ana Sofía'),
            Col('apellidos', 'Apellidos', True, ['apellido'], '', 'Acosta Bernal'),
            Col('curso', 'Curso', True, ['grupo', 'grado'], 'Nombre de un curso ya creado.', '601'),
            Col('tipo_documento', 'Tipo de documento', False, ['tipo doc', 'tipo id', 'tipo identificacion', 'td'],
                'TI, RC, CC, CE u OT.', 'TI'),
            Col('documento', 'Documento', False, ['numero de documento', 'no documento', 'identificacion', 'numero documento'],
                'Muy recomendado: con él se reconoce al estudiante si se vuelve a subir el archivo.', '1105222333'),
            Col('fecha_nacimiento', 'Fecha de nacimiento', False, ['fecha nacimiento', 'nacimiento', 'f nacimiento'],
                'dd/mm/aaaa', '12/04/2014'),
            Col('lugar_nacimiento', 'Lugar de nacimiento', False, ['lugar nacimiento'], '', 'Honda'),
            Col('eps', 'EPS', False, [], '', 'Nueva EPS'),
            Col('rh', 'Grupo sanguíneo', False, ['rh', 'grupo sanguineo y rh', 'tipo de sangre'], 'O+, A-, …', 'O+'),
            Col('enfermedades', 'Enfermedades o alergias', False, ['alergias', 'enfermedades alergias'], '', ''),
            Col('nombre_madre', 'Nombre de la madre', False, ['madre', 'mama'], '', 'Carolina Bernal'),
            Col('celular_madre', 'Celular de la madre', False, ['telefono madre', 'cel madre'], '', '3105556677'),
            Col('nombre_padre', 'Nombre del padre', False, ['padre', 'papa'], '', ''),
            Col('celular_padre', 'Celular del padre', False, ['telefono padre', 'cel padre'], '', ''),
            Col('nombre_acudiente', 'Nombre del acudiente', False, ['acudiente'], '', 'Carolina Bernal'),
            Col('celular_acudiente', 'Celular del acudiente', False, ['telefono acudiente', 'cel acudiente'], '', '3105556677'),
            Col('correo_acudiente', 'Correo del acudiente', False, ['email acudiente', 'correo'], '', 'carolina@correo.com'),
            Col('colegio_anterior', 'Colegio anterior', False, ['institucion anterior', 'procedencia'], '', ''),
            Col('grado_anterior', 'Grado anterior', False, [], '', ''),
        ],
    },
}
ORDEN = ['cursos', 'docentes', 'materias', 'ponderacion', 'asignacion', 'estudiantes']


def columnas(tipo):
    return TIPOS[tipo]['columnas']


# ---------------------------------------------------------------------------
# Leer el archivo
# ---------------------------------------------------------------------------

class ArchivoInvalido(Exception):
    pass


def _indice_de_titulos(tipo):
    indice = {}
    for c in columnas(tipo):
        for nombre in [c.clave, c.titulo] + list(c.sinonimos):
            indice.setdefault(norm(nombre.replace('_', ' ')), c.clave)
    return indice


def _mapear(fila_titulos, tipo):
    """{posición: clave} de los títulos reconocidos (quita el asterisco de obligatorio)."""
    indice = _indice_de_titulos(tipo)
    mapa, usadas = {}, set()
    for i, t in enumerate(fila_titulos):
        clave = indice.get(norm(str(t or '').replace('*', '')))
        if clave and clave not in usadas:
            mapa[i] = clave
            usadas.add(clave)
    return mapa


def _filas_de_xlsx(archivo):
    from openpyxl import load_workbook
    try:
        libro = load_workbook(archivo, data_only=True, read_only=True)
    except Exception:
        raise ArchivoInvalido('No se pudo abrir el Excel. Guárdelo como .xlsx y vuelva a subirlo.')
    hoja = libro['Datos'] if 'Datos' in libro.sheetnames else next(
        (h for h in libro.worksheets if h.sheet_state == 'visible'), libro.worksheets[0])
    return [list(r) for r in hoja.iter_rows(values_only=True)]


def _filas_de_csv(archivo):
    crudo = archivo.read()
    for codificacion in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            texto = crudo.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    muestra = texto[:4000]
    separador = max([';', ',', '\t'], key=lambda s: muestra.count(s))
    return [r for r in csv.reader(io.StringIO(texto), delimiter=separador)]


def leer(archivo, tipo):
    """[(número de fila, {clave: valor})] y la lista de columnas reconocidas."""
    nombre = (getattr(archivo, 'name', '') or '').lower()
    if nombre.endswith(('.xlsx', '.xlsm')):
        filas = _filas_de_xlsx(archivo)
    elif nombre.endswith(('.csv', '.txt')):
        filas = _filas_de_csv(archivo)
    elif nombre.endswith('.xls'):
        raise ArchivoInvalido('Ese es el formato viejo de Excel (.xls). Ábralo y guárdelo como .xlsx.')
    else:
        raise ArchivoInvalido('Suba un archivo de Excel (.xlsx) o CSV.')
    # El título puede no estar en la primera fila: se busca en las primeras 10 la que más columnas reconoce.
    mejor, mapa = None, {}
    for k, f in enumerate(filas[:10]):
        m = _mapear(f, tipo)
        if len(m) > len(mapa):
            mejor, mapa = k, m
    faltan = [c.titulo for c in columnas(tipo) if c.obligatoria and c.clave not in mapa.values()]
    if mejor is None or faltan:
        raise ArchivoInvalido('No se encontraron las columnas obligatorias: ' + ', '.join(faltan or
                              [c.titulo for c in columnas(tipo) if c.obligatoria]) +
                              '. Use la plantilla o revise los títulos de la primera fila.')
    salida = []
    for k, f in enumerate(filas[mejor + 1:], start=mejor + 2):
        datos = {}
        for i, clave in mapa.items():
            v = f[i] if i < len(f) else None
            if isinstance(v, str):
                v = ' '.join(v.split())
            if isinstance(v, float) and v.is_integer():
                v = int(v)                      # 1105123456.0 -> 1105123456
            if isinstance(v, (datetime.date, datetime.datetime)):
                v = v.strftime('%d/%m/%Y')      # se guarda en la sesión hasta confirmar
            elif v is not None and not isinstance(v, (str, int, float)):
                v = str(v)
            datos[clave] = v
        if any(v not in (None, '') for v in datos.values()):
            salida.append((k, datos))
        if len(salida) > MAX_FILAS:
            raise ArchivoInvalido(f'El archivo tiene más de {MAX_FILAS} filas. Divídalo en partes.')
    reconocidas = [c.titulo for c in columnas(tipo) if c.clave in mapa.values()]
    return salida, reconocidas


# ---------------------------------------------------------------------------
# Ayudas para reconocer cosas que ya existen
# ---------------------------------------------------------------------------

def _texto(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _documento(v):
    """'1.105.123.456' o 1105123456.0 -> '1105123456'."""
    t = _texto(v)
    return re.sub(r'[\s.\-]', '', t)


_GRADOS_TEXTO = None


def parse_grado(v):
    global _GRADOS_TEXTO
    from .models.perfiles import GRADOS
    if _GRADOS_TEXTO is None:
        _GRADOS_TEXTO = {norm(n): g for g, n, _ in GRADOS}
        _GRADOS_TEXTO.update({'once': 11, 'decimo primero': 11, 'pre jardin': -2, 'prekinder': -2, 'kinder': -1,
                              'transicion': 0, 'grado 0': 0, 'cero': 0})
    t = norm(v)
    if not t:
        return None
    m = re.fullmatch(r'(?:grado )?(-?\d{1,2})(?: ?o)?', t)        # «6», «6°», «grado 6», «6o»
    if m:
        n = int(m.group(1))
        return n if -2 <= n <= 11 else None
    return _GRADOS_TEXTO.get(t)


class Docentes:
    """Busca un docente del colegio por documento, usuario, correo o nombre completo."""

    def __init__(self, colegio):
        from .models import Docente
        self.por = {}
        self.ambiguos = set()
        for d in Docente.objects.filter(colegio=colegio).select_related('user', 'ficha'):
            u = d.user
            claves = {u.username.lower(), norm(f'{u.first_name} {u.last_name}'), norm(f'{u.last_name} {u.first_name}')}
            if u.email:
                claves.add(u.email.lower())
            ficha = getattr(d, 'ficha', None)
            if ficha is not None and ficha.numero_documento:
                claves.add(_documento(ficha.numero_documento))
            for c in claves:
                if c in self.por and self.por[c].id != d.id:
                    self.ambiguos.add(c)
                self.por[c] = d

    def buscar(self, v):
        t = _texto(v)
        for c in (_documento(t), t.lower(), norm(t)):
            if c and c in self.por:
                if c in self.ambiguos:
                    return None, f'«{t}» coincide con varios docentes: use el documento'
                return self.por[c], ''
        return None, f'no hay un docente «{t}» en el colegio'


def _clave_curso(nombre):
    return re.sub(r'\s+', '', norm(nombre)).upper()


class Usuarios:
    """Genera usuarios «nombre.apellido» sin repetir, también dentro del mismo archivo."""

    def __init__(self):
        self.usados = None

    def nuevo(self, nombres, apellidos):
        from unidecode import unidecode
        if self.usados is None:
            self.usados = set(User.objects.values_list('username', flat=True))
        primer_nombre = unidecode(str(nombres).split(' ')[0].lower())
        primer_apellido = unidecode(str(apellidos).split(' ')[0].lower())
        base = f'{slugify(primer_nombre)}.{slugify(primer_apellido)}'.strip('.') or 'usuario'
        final, k = base, 1
        while final in self.usados:
            final = f'{base}{k}'
            k += 1
        self.usados.add(final)
        return final


def _fecha(v):
    if v in (None, ''):
        return None, ''
    if isinstance(v, datetime.datetime):
        return v.date(), ''
    if isinstance(v, datetime.date):
        return v, ''
    t = _texto(v)
    for formato in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d/%m/%y', '%Y/%m/%d'):
        try:
            return datetime.datetime.strptime(t, formato).date(), ''
        except ValueError:
            continue
    return None, f'la fecha «{t}» no se entiende (use dd/mm/aaaa)'


def _resultado(fila, accion, resumen, datos=None, cambios=None, errores=None, avisos=None):
    return {'fila': fila, 'accion': accion, 'resumen': resumen, 'datos': datos or {},
            'cambios': cambios or [], 'errores': errores or [], 'avisos': avisos or []}


def _duplicado(vistos, clave, fila, que):
    if clave in vistos:
        return f'{que} ya viene en la fila {vistos[clave]} de este mismo archivo'
    vistos[clave] = fila
    return ''


# ---------------------------------------------------------------------------
# Revisar (sin guardar)
# ---------------------------------------------------------------------------

def revisar(tipo, filas, colegio):
    """Lista de resultados, uno por fila (o por curso, en la asignación)."""
    return {'cursos': _revisar_cursos, 'docentes': _revisar_docentes, 'materias': _revisar_materias,
            'ponderacion': _revisar_ponderacion, 'asignacion': _revisar_asignacion,
            'estudiantes': _revisar_estudiantes}[tipo](filas, colegio)


def resumen(resultados):
    r = {'crear': 0, 'actualizar': 0, 'igual': 0, 'error': 0}
    for x in resultados:
        r[x['accion']] += 1
    r['total'] = len(resultados)
    r['aplicables'] = r['crear'] + r['actualizar']
    return r


def _revisar_cursos(filas, colegio):
    from .models import Curso, Sede
    from .models.perfiles import formato_del_colegio, nombre_curso_sugerido, normalizar_subgrupo
    existentes = {_clave_curso(c.nombre): c for c in Curso.objects.filter(colegio=colegio).select_related('sede', 'director_grado')}
    sedes = {norm(s.nombre): s for s in Sede.objects.filter(colegio=colegio)}
    docentes = Docentes(colegio)
    formato = formato_del_colegio(colegio)
    vistos, salida = {}, []
    for n, d in filas:
        errores = []
        grado = parse_grado(d.get('grado'))
        if grado is None:
            errores.append(f'el grado «{_texto(d.get("grado"))}» no se reconoce (use 0 a 11 o el nombre del grado)')
        subgrupo = normalizar_subgrupo(_texto(d.get('subgrupo')))
        nombre = _texto(d.get('nombre')).upper()
        if not nombre and grado is not None:
            nombre = nombre_curso_sugerido(grado, subgrupo, formato).upper()
        sede = None
        if _texto(d.get('sede')):
            sede = sedes.get(norm(d.get('sede')))
            if sede is None:
                errores.append(f'no hay una sede «{_texto(d.get("sede"))}»')
        director = None
        if _texto(d.get('director')):
            director, error = docentes.buscar(d.get('director'))
            if error:
                errores.append(error)
        if nombre:
            dup = _duplicado(vistos, _clave_curso(nombre), n, f'el curso {nombre}')
            if dup:
                errores.append(dup)
        datos = {'nombre': nombre, 'grado': grado, 'subgrupo': subgrupo,
                 'sede': sede.id if sede else None, 'director': director.id if director else None}
        if errores:
            salida.append(_resultado(n, 'error', nombre or '(sin nombre)', datos, errores=errores))
            continue
        actual = existentes.get(_clave_curso(nombre))
        if actual is None:
            extra = [f'sede {sede.nombre}'] if sede else []
            extra += [f'director {director.user.get_full_name()}'] if director else []
            salida.append(_resultado(n, 'crear', nombre, datos, cambios=extra))
            continue
        cambios = []
        if actual.grado != grado:
            cambios.append(f'grado: {actual.get_grado_display() if actual.grado is not None else "—"} → {dict(_grados()).get(grado)}')
        if subgrupo and actual.subgrupo != subgrupo:
            cambios.append(f'subgrupo: {actual.subgrupo or "—"} → {subgrupo}')
        if sede and actual.sede_id != sede.id:
            cambios.append(f'sede: {actual.sede.nombre if actual.sede else "—"} → {sede.nombre}')
        if director and actual.director_grado_id != director.id:
            antes = actual.director_grado.user.get_full_name() if actual.director_grado else '—'
            cambios.append(f'director: {antes} → {director.user.get_full_name()}')
        datos['id'] = actual.id
        salida.append(_resultado(n, 'actualizar' if cambios else 'igual', nombre, datos, cambios=cambios))
    return salida


def _grados():
    from .models.perfiles import GRADOS
    return [(g, nombre) for g, nombre, _ in GRADOS]


def _revisar_docentes(filas, colegio):
    from .models import FichaDocente
    propios = {_documento(f.numero_documento): f.docente for f in
               FichaDocente.objects.filter(docente__colegio=colegio).exclude(numero_documento__isnull=True)
               .select_related('docente__user')}
    ajenos = set(_documento(x) for x in FichaDocente.objects.exclude(docente__colegio=colegio)
                 .exclude(numero_documento__isnull=True).values_list('numero_documento', flat=True))
    vistos, salida = {}, []
    for n, d in filas:
        nombres, apellidos = _texto(d.get('nombres')), _texto(d.get('apellidos'))
        doc, correo = _documento(d.get('documento')), _texto(d.get('correo')).lower()
        errores = []
        if not nombres:
            errores.append('faltan los nombres')
        if not apellidos:
            errores.append('faltan los apellidos')
        if not doc:
            errores.append('falta el documento')
        elif len(doc) > 20:
            errores.append('el documento es muy largo')
        if correo:
            try:
                validate_email(correo)
            except ValidationError:
                errores.append(f'el correo «{correo}» no es válido')
        if doc:
            dup = _duplicado(vistos, doc, n, f'el documento {doc}')
            if dup:
                errores.append(dup)
            if doc in ajenos and doc not in propios:
                errores.append(f'el documento {doc} ya es de un docente de otro colegio')
        nombre_ver = f'{apellidos} {nombres}'.strip().upper()
        datos = {'nombres': nombres, 'apellidos': apellidos, 'documento': doc, 'correo': correo,
                 'telefono': _texto(d.get('telefono'))[:20], 'titulo': _texto(d.get('titulo'))[:200]}
        if errores:
            salida.append(_resultado(n, 'error', nombre_ver or '(sin nombre)', datos, errores=errores))
            continue
        actual = propios.get(doc)
        if actual is None:
            salida.append(_resultado(n, 'crear', nombre_ver, datos, cambios=[f'documento {doc}']))
            continue
        u, ficha = actual.user, actual.ficha
        cambios = []
        if norm(u.first_name) != norm(nombres) or norm(u.last_name) != norm(apellidos):
            cambios.append(f'nombre: {u.get_full_name()} → {nombres} {apellidos}')
        if correo and (u.email or '').lower() != correo:
            cambios.append(f'correo: {u.email or "—"} → {correo}')
        if datos['telefono'] and (ficha.telefono or '') != datos['telefono']:
            cambios.append(f'teléfono: {ficha.telefono or "—"} → {datos["telefono"]}')
        if datos['titulo'] and (ficha.titulo_profesional or '') != datos['titulo']:
            cambios.append('título profesional')
        datos['id'] = actual.id
        salida.append(_resultado(n, 'actualizar' if cambios else 'igual', nombre_ver, datos, cambios=cambios,
                                 avisos=[f'usuario {u.username}']))
    return salida


def _revisar_materias(filas, colegio):
    from .models import Materia
    from .models.academicos import AreaConocimiento, PonderacionAreaMateria
    materias = {norm(m.nombre): m for m in Materia.objects.filter(colegio=colegio)}
    areas = {norm(a.nombre): a for a in AreaConocimiento.objects.filter(colegio=colegio)}
    enlaces = set(PonderacionAreaMateria.objects.filter(colegio=colegio).values_list('area_id', 'materia_id'))
    vistos, salida = {}, []
    for n, d in filas:
        nombre, area = _texto(d.get('nombre')).upper(), _texto(d.get('area')).upper()
        abrev = _texto(d.get('abreviatura')).upper()
        errores = []
        if not nombre:
            errores.append('falta el nombre de la materia')
        if not area:
            errores.append('falta el área')
        if len(abrev) > 10:
            errores.append('la abreviatura tiene más de 10 letras')
        if nombre:
            dup = _duplicado(vistos, norm(nombre), n, f'la materia {nombre}')
            if dup:
                errores.append(dup)
        datos = {'nombre': nombre, 'area': area, 'abreviatura': abrev}
        if errores:
            salida.append(_resultado(n, 'error', nombre or '(sin nombre)', datos, errores=errores))
            continue
        m, a = materias.get(norm(nombre)), areas.get(norm(area))
        if m is None:
            salida.append(_resultado(n, 'crear', nombre, datos,
                                     cambios=[f'área {area}' + ('' if a else ' (nueva)')]))
            continue
        cambios = []
        if a is None or (a.id, m.id) not in enlaces:
            cambios.append(f'se une al área {area}' + ('' if a else ' (nueva)'))
        if abrev and (m.abreviatura or '') != abrev:
            cambios.append(f'abreviatura: {m.abreviatura or "—"} → {abrev}')
        datos['id'] = m.id
        salida.append(_resultado(n, 'actualizar' if cambios else 'igual', nombre, datos, cambios=cambios))
    return salida


NIVELES_GRADO = {'preescolar': [-2, -1, 0], 'primaria': [1, 2, 3, 4, 5], 'basica primaria': [1, 2, 3, 4, 5],
                 'secundaria': [6, 7, 8, 9], 'basica secundaria': [6, 7, 8, 9], 'basica': [6, 7, 8, 9],
                 'media': [10, 11], 'media tecnica': [10, 11], 'media academica': [10, 11]}


def grados_de(v):
    """'' -> [None] (general); '10, 11' -> [10, 11]; 'Media' -> [10, 11]. None si algo no se entiende."""
    t = norm(v)
    if t in ('', 'general', 'todos', 'todos los grados', 'todo'):
        return [None]
    salida = []
    for trozo in re.split(r'[,;/]|\s+y\s+', _texto(v)):
        k = norm(trozo)
        if not k:
            continue
        if k in NIVELES_GRADO:
            salida += NIVELES_GRADO[k]
            continue
        g = parse_grado(trozo)
        if g is None:
            return None
        salida.append(g)
    return sorted(set(salida)) or [None]


def _nombre_grados(grados):
    if grados == [None]:
        return 'general'
    nombres = dict(_grados())
    return ', '.join(nombres.get(g, str(g)) for g in grados)


def _revisar_ponderacion(filas, colegio):
    from .models.academicos import AreaConocimiento, PonderacionAreaMateria, PonderacionGrado
    areas = {norm(a.nombre): a for a in AreaConocimiento.objects.filter(colegio=colegio)}
    generales = {(p.area_id, p.materia_id): p for p in
                 PonderacionAreaMateria.objects.filter(colegio=colegio).select_related('materia')}
    materias_de_area = {}
    for (area_id, _), p in generales.items():
        materias_de_area.setdefault(area_id, {})[norm(p.materia.nombre)] = p.materia
    propios = {(p.area_id, p.materia_id, p.grado): p.peso_porcentual
               for p in PonderacionGrado.objects.filter(colegio=colegio)}
    vistos, salida, sumas = {}, [], {}
    for n, d in filas:
        errores = []
        area = areas.get(norm(d.get('area'))) if _texto(d.get('area')) else None
        if area is None:
            errores.append(f'no hay un área «{_texto(d.get("area"))}»' if _texto(d.get('area')) else 'falta el área')
        materia = None
        if area is not None:
            materia = materias_de_area.get(area.id, {}).get(norm(d.get('materia')))
            if materia is None:
                errores.append(f'«{_texto(d.get("materia"))}» no es una materia del área {area.nombre} '
                               '(únala al área en Materias)' if _texto(d.get('materia')) else 'falta la materia')
        grados = grados_de(d.get('grado'))
        if grados is None:
            errores.append(f'el grado «{_texto(d.get("grado"))}» no se reconoce (deje vacío para el general)')
        peso = None
        try:
            peso = Decimal(_texto(d.get('peso')).replace('%', '').replace(',', '.').strip())
            if not Decimal('0') <= peso <= Decimal('100'):
                raise InvalidOperation
        except (InvalidOperation, ValueError):
            errores.append(f'el peso «{_texto(d.get("peso"))}» debe ser un número de 0 a 100')
        if area and materia and grados:
            for g in grados:
                dup = _duplicado(vistos, (area.id, materia.id, g), n,
                                 f'{materia.nombre} en {area.nombre} ({_nombre_grados([g])})')
                if dup:
                    errores.append(dup)
                    break
        resumen_fila = (f'{materia.nombre if materia else _texto(d.get("materia"))} en '
                        f'{area.nombre if area else _texto(d.get("area"))} · {_nombre_grados(grados or [None])}')
        datos = {'area': area.id if area else None, 'materia': materia.id if materia else None,
                 'grados': grados, 'peso': str(peso) if peso is not None else None}
        if errores:
            salida.append(_resultado(n, 'error', resumen_fila, datos, errores=errores))
            continue
        cambios = []
        for g in grados:
            antes = generales[(area.id, materia.id)].peso_porcentual if g is None else propios.get((area.id, materia.id, g))
            if antes is None or Decimal(antes) != peso:
                cambios.append(f'{_nombre_grados([g])}: {"—" if antes is None else f"{Decimal(antes):g} %"} → {peso:g} %')
            sumas.setdefault((area.id, g), [Decimal('0'), area.nombre, n])
            sumas[(area.id, g)][0] += peso
            sumas[(area.id, g)][2] = n
        salida.append(_resultado(n, 'actualizar' if cambios else 'igual', resumen_fila, datos,
                                 cambios=cambios))
    # Aviso cuando las materias de un área (en un grado) no suman 100 en el archivo.
    por_fila = {r['fila']: r for r in salida}
    for (area_id, g), (total, nombre_area, fila) in sumas.items():
        if abs(total - 100) > Decimal('0.01'):
            por_fila[fila]['avisos'].append(f'{nombre_area} ({_nombre_grados([g])}) suma {total:g} % en este archivo, no 100 %')
    return salida


def _revisar_asignacion(filas, colegio):
    from .models import AsignacionDocente, Curso, Materia
    docentes = Docentes(colegio)
    materias = {norm(m.nombre): m for m in Materia.objects.filter(colegio=colegio)}
    cursos = {_clave_curso(c.nombre): c for c in Curso.objects.filter(colegio=colegio)}
    actuales = {}
    for a in AsignacionDocente.objects.filter(colegio=colegio).select_related('docente__user'):
        actuales.setdefault((a.materia_id, a.curso_id), []).append(a)
    vistos, salida = {}, []
    for n, d in filas:
        errores = []
        docente, error = docentes.buscar(d.get('docente')) if _texto(d.get('docente')) else (None, 'falta el docente')
        if error:
            errores.append(error)
        materia = materias.get(norm(d.get('materia'))) if _texto(d.get('materia')) else None
        if materia is None:
            errores.append(f'no hay una materia «{_texto(d.get("materia"))}» (impórtela primero en Materias)'
                           if _texto(d.get('materia')) else 'falta la materia')
        ih_txt = _texto(d.get('ih'))
        ih = None
        if ih_txt:
            if re.fullmatch(r'\d{1,2}', ih_txt):
                ih = int(ih_txt)
            else:
                errores.append(f'la intensidad «{ih_txt}» debe ser un número de horas')
        nombres_cursos = [c for c in re.split(r'[,;/]|\s+y\s+', _texto(d.get('curso'))) if c.strip()]
        if not nombres_cursos:
            errores.append('falta el curso')
        quien = docente.user.get_full_name().upper() if docente else _texto(d.get('docente'))
        que = materia.nombre if materia else _texto(d.get('materia')).upper()
        if errores:
            salida.append(_resultado(n, 'error', f'{que} · {", ".join(nombres_cursos)} · {quien}', {}, errores=errores))
            continue
        for nombre_curso in nombres_cursos:
            curso = cursos.get(_clave_curso(nombre_curso))
            resumen_fila = f'{que} · {nombre_curso.strip().upper()} · {quien}'
            if curso is None:
                salida.append(_resultado(n, 'error', resumen_fila, {}, errores=[f'no hay un curso «{nombre_curso.strip()}»']))
                continue
            dup = _duplicado(vistos, (materia.id, curso.id), n, f'{que} en {curso.nombre}')
            if dup:
                salida.append(_resultado(n, 'error', resumen_fila, {}, errores=[dup]))
                continue
            datos = {'docente': docente.id, 'materia': materia.id, 'curso': curso.id, 'ih': ih}
            hay = actuales.get((materia.id, curso.id), [])
            propia = next((a for a in hay if a.docente_id == docente.id), None)
            if propia:
                cambios = [f'intensidad: {propia.intensidad_horaria_semanal} → {ih}'] if ih is not None and propia.intensidad_horaria_semanal != ih else []
                datos['id'] = propia.id
                salida.append(_resultado(n, 'actualizar' if cambios else 'igual', resumen_fila, datos, cambios=cambios))
            elif len(hay) == 1:
                datos['id'] = hay[0].id
                salida.append(_resultado(n, 'actualizar', resumen_fila, datos,
                                         cambios=[f'cambia de docente: {hay[0].docente.user.get_full_name()} → {docente.user.get_full_name()}'],
                                         avisos=['las notas ya registradas se conservan']))
            elif len(hay) > 1:
                salida.append(_resultado(n, 'error', resumen_fila, datos, errores=[
                    f'{que} en {curso.nombre} ya la tienen {len(hay)} docentes: corríjalo en Asignación académica']))
            else:
                salida.append(_resultado(n, 'crear', resumen_fila, datos,
                                         cambios=[f'{ih} h/semana'] if ih is not None else []))
    return salida


_TIPOS_DOC = None


def _tipo_documento(v):
    global _TIPOS_DOC
    from .models import FichaEstudiante
    if _TIPOS_DOC is None:
        _TIPOS_DOC = {}
        for codigo, nombre in FichaEstudiante.TIPO_DOCUMENTO_CHOICES:
            _TIPOS_DOC[norm(codigo)] = codigo
            _TIPOS_DOC[norm(nombre)] = codigo
        _TIPOS_DOC.update({'t i': 'TI', 'tarjeta': 'TI', 'registro': 'RC', 'nuip': 'RC', 'cedula': 'CC', 'c c': 'CC',
                           'r c': 'RC', 'extranjeria': 'CE', 'pep': 'OT', 'ppt': 'OT', 'pasaporte': 'OT'})
    t = norm(v)
    return _TIPOS_DOC.get(t) if t else ''


def _rh(v):
    t = _texto(v).upper().replace(' ', '').replace('POSITIVO', '+').replace('NEGATIVO', '-').replace('POS', '+').replace('NEG', '-')
    return t if t in ('O+', 'O-', 'A+', 'A-', 'B+', 'B-', 'AB+', 'AB-') else None


FICHA_TEXTO = ['lugar_nacimiento', 'eps', 'enfermedades', 'nombre_madre', 'celular_madre', 'nombre_padre',
               'celular_padre', 'nombre_acudiente', 'celular_acudiente', 'colegio_anterior', 'grado_anterior']
CAMPO_FICHA = {'enfermedades': 'enfermedades_alergias'}
LARGO_FICHA = {'celular_madre': 20, 'celular_padre': 20, 'celular_acudiente': 20, 'grado_anterior': 20,
               'lugar_nacimiento': 100, 'eps': 100}


def _revisar_estudiantes(filas, colegio):
    from .models import Curso, Estudiante, FichaEstudiante
    cursos = {_clave_curso(c.nombre): c for c in Curso.objects.filter(colegio=colegio)}
    propios_doc, propios_nombre = {}, {}
    for e in Estudiante.objects.filter(colegio=colegio).select_related('user', 'curso', 'ficha'):
        ficha = getattr(e, 'ficha', None)
        if ficha is not None and ficha.numero_documento:
            propios_doc[_documento(ficha.numero_documento)] = e
        propios_nombre.setdefault(norm(f'{e.user.last_name} {e.user.first_name}'), []).append(e)
    ajenos = set(_documento(x) for x in FichaEstudiante.objects.exclude(estudiante__colegio=colegio)
                 .exclude(numero_documento__isnull=True).values_list('numero_documento', flat=True))
    vistos, salida = {}, []
    for n, d in filas:
        nombres, apellidos = _texto(d.get('nombres')).upper(), _texto(d.get('apellidos')).upper()
        doc = _documento(d.get('documento'))
        errores, avisos = [], []
        if not nombres:
            errores.append('faltan los nombres')
        if not apellidos:
            errores.append('faltan los apellidos')
        curso = None
        if not _texto(d.get('curso')):
            errores.append('falta el curso')
        else:
            curso = cursos.get(_clave_curso(d.get('curso')))
            if curso is None:
                errores.append(f'no hay un curso «{_texto(d.get("curso"))}» (créelo o impórtelo primero)')
        tipo_doc = _tipo_documento(d.get('tipo_documento'))
        if tipo_doc is None:
            errores.append(f'el tipo de documento «{_texto(d.get("tipo_documento"))}» no se reconoce (TI, RC, CC, CE, OT)')
        if doc and len(doc) > 20:
            errores.append('el documento es muy largo')
        fecha, error = _fecha(d.get('fecha_nacimiento'))
        if error:
            errores.append(error)
        rh = _rh(d.get('rh')) if _texto(d.get('rh')) else None
        if _texto(d.get('rh')) and rh is None:
            avisos.append(f'el grupo sanguíneo «{_texto(d.get("rh"))}» no se reconoce y se deja vacío')
        correo = _texto(d.get('correo_acudiente')).lower()
        if correo:
            try:
                validate_email(correo)
            except ValidationError:
                errores.append(f'el correo «{correo}» no es válido')
        if doc:
            dup = _duplicado(vistos, doc, n, f'el documento {doc}')
            if dup:
                errores.append(dup)
            if doc in ajenos and doc not in propios_doc:
                errores.append(f'el documento {doc} ya es de un estudiante de otro colegio')
        else:
            dup = _duplicado(vistos, norm(f'{apellidos} {nombres}'), n, f'{apellidos} {nombres}')
            if dup:
                errores.append(dup)
        nombre_ver = f'{apellidos} {nombres}'.strip()
        ficha = {k: _texto(d.get(k))[:LARGO_FICHA.get(k, 200)] for k in FICHA_TEXTO if _texto(d.get(k))}
        datos = {'nombres': nombres, 'apellidos': apellidos, 'documento': doc, 'tipo_documento': tipo_doc or '',
                 'curso': curso.id if curso else None, 'fecha_nacimiento': fecha.isoformat() if fecha else '',
                 'rh': rh or '', 'correo_acudiente': correo, 'ficha': ficha}
        if errores:
            salida.append(_resultado(n, 'error', nombre_ver or '(sin nombre)', datos, errores=errores, avisos=avisos))
            continue
        actual = propios_doc.get(doc) if doc else None
        if actual is None and not doc:
            mismos = propios_nombre.get(norm(f'{apellidos} {nombres}'), [])
            if len(mismos) == 1:
                actual = mismos[0]
                avisos.append('reconocido por el nombre (sin documento)')
            elif len(mismos) > 1:
                salida.append(_resultado(n, 'error', nombre_ver, datos, errores=[
                    'hay varios estudiantes con ese nombre: agregue el documento'], avisos=avisos))
                continue
        if actual is None:
            if not doc:
                avisos.append('sin documento: si vuelve a subir el archivo, se reconoce por el nombre')
            salida.append(_resultado(n, 'crear', f'{nombre_ver} → {curso.nombre}', datos, avisos=avisos))
            continue
        cambios = []
        if actual.curso_id != curso.id:
            cambios.append(f'curso: {actual.curso.nombre if actual.curso else "—"} → {curso.nombre}')
        if norm(actual.user.first_name) != norm(nombres) or norm(actual.user.last_name) != norm(apellidos):
            cambios.append(f'nombre: {actual.user.last_name} {actual.user.first_name} → {nombre_ver}')
        if not actual.is_active:
            cambios.append('vuelve a quedar activo')
        f_actual = getattr(actual, 'ficha', None)
        nuevos = [k for k, v in ficha.items()
                  if f_actual is None or (getattr(f_actual, CAMPO_FICHA.get(k, k)) or '') != v]
        if fecha and (f_actual is None or f_actual.fecha_nacimiento != fecha):
            nuevos.append('fecha_nacimiento')
        if rh and (f_actual is None or f_actual.grupo_sanguineo != rh):
            nuevos.append('rh')
        if correo and (f_actual is None or (f_actual.email_acudiente or '').lower() != correo):
            nuevos.append('correo del acudiente')
        if doc and f_actual is not None and tipo_doc and f_actual.tipo_documento != tipo_doc:
            nuevos.append('tipo de documento')
        if nuevos:
            cambios.append(f'ficha: {len(nuevos)} dato(s)')
        datos['id'] = actual.id
        salida.append(_resultado(n, 'actualizar' if cambios else 'igual', f'{nombre_ver} → {curso.nombre}', datos,
                                 cambios=cambios, avisos=avisos + [f'usuario {actual.user.username}']))
    return salida


# ---------------------------------------------------------------------------
# Aplicar
# ---------------------------------------------------------------------------

@transaction.atomic
def aplicar(tipo, filas, colegio):
    """Revisa otra vez contra la base y guarda. Devuelve (resumen, credenciales nuevas)."""
    resultados = revisar(tipo, filas, colegio)
    credenciales = []
    hacer = {'cursos': _aplicar_curso, 'docentes': _aplicar_docente, 'materias': _aplicar_materia,
             'ponderacion': _aplicar_ponderacion, 'asignacion': _aplicar_asignacion,
             'estudiantes': _aplicar_estudiante}[tipo]
    usuarios = Usuarios()
    for r in resultados:
        if r['accion'] in ('crear', 'actualizar'):
            nueva = hacer(r, colegio, usuarios)
            if nueva:
                credenciales.append(nueva)
    return resumen(resultados), credenciales, resultados


def _aplicar_curso(r, colegio, usuarios):
    from .models import Curso
    d = r['datos']
    curso = Curso.objects.get(id=d['id'], colegio=colegio) if d.get('id') else Curso(colegio=colegio, nombre=d['nombre'])
    curso.grado = d['grado']
    if d['subgrupo'] or not d.get('id'):
        curso.subgrupo = d['subgrupo']
    if d['sede']:
        curso.sede_id = d['sede']
    if d['director']:
        curso.director_grado_id = d['director']
    curso.save()


def _aplicar_docente(r, colegio, usuarios):
    from .models import Docente, FichaDocente
    d = r['datos']
    if d.get('id'):
        docente = Docente.objects.select_related('user').get(id=d['id'], colegio=colegio)
        u = docente.user
        u.first_name, u.last_name = d['nombres'], d['apellidos']
        if d['correo']:
            u.email = d['correo']
        u.save()
        ficha, _ = FichaDocente.objects.get_or_create(docente=docente)
        if d['telefono']:
            ficha.telefono = d['telefono']
        if d['titulo']:
            ficha.titulo_profesional = d['titulo']
        ficha.save()
        return None
    usuario = usuarios.nuevo(d['nombres'], d['apellidos'])
    u = User.objects.create_user(username=usuario, password=usuario, first_name=d['nombres'],
                                 last_name=d['apellidos'], email=d['correo'] or '')
    grupo, _ = Group.objects.get_or_create(name='Docentes')
    u.groups.add(grupo)
    docente = Docente.objects.create(user=u, colegio=colegio)
    FichaDocente.objects.create(docente=docente, numero_documento=d['documento'] or None,
                                telefono=d['telefono'] or None, titulo_profesional=d['titulo'] or None)
    return {'rol': 'Docente', 'nombre': f'{d["apellidos"]} {d["nombres"]}'.upper(), 'usuario': usuario,
            'clave': usuario, 'detalle': d['documento']}


def _aplicar_materia(r, colegio, usuarios):
    from .models import Materia
    from .models.academicos import AreaConocimiento, PonderacionAreaMateria
    d = r['datos']
    if d.get('id'):
        materia = Materia.objects.get(id=d['id'], colegio=colegio)
        if d['abreviatura']:
            materia.abreviatura = d['abreviatura']
            materia.save()
    else:
        materia = Materia.objects.create(colegio=colegio, nombre=d['nombre'], abreviatura=d['abreviatura'] or None)
    area, _ = AreaConocimiento.objects.get_or_create(colegio=colegio, nombre=d['area'])
    PonderacionAreaMateria.objects.get_or_create(colegio=colegio, area=area, materia=materia,
                                                 defaults={'peso_porcentual': 0})


def _aplicar_ponderacion(r, colegio, usuarios):
    from .models.academicos import PonderacionAreaMateria, PonderacionGrado
    d = r['datos']
    peso = Decimal(d['peso'])
    for g in d['grados']:
        if g is None:
            PonderacionAreaMateria.objects.filter(colegio=colegio, area_id=d['area'], materia_id=d['materia']).update(
                peso_porcentual=peso)
        else:
            PonderacionGrado.objects.update_or_create(colegio=colegio, area_id=d['area'], materia_id=d['materia'],
                                                      grado=g, defaults={'peso_porcentual': peso})


def _aplicar_asignacion(r, colegio, usuarios):
    from .models import AsignacionDocente
    d = r['datos']
    if d.get('id'):
        a = AsignacionDocente.objects.get(id=d['id'], colegio=colegio)
        a.docente_id = d['docente']
        if d['ih'] is not None:
            a.intensidad_horaria_semanal = d['ih']
        a.save()
    else:
        AsignacionDocente.objects.create(colegio=colegio, docente_id=d['docente'], materia_id=d['materia'],
                                         curso_id=d['curso'], intensidad_horaria_semanal=d['ih'] or 0)


def _aplicar_estudiante(r, colegio, usuarios):
    from .models import Estudiante, FichaEstudiante
    d = r['datos']
    nueva = None
    if d.get('id'):
        est = Estudiante.objects.select_related('user').get(id=d['id'], colegio=colegio)
        u = est.user
        u.first_name, u.last_name = d['nombres'], d['apellidos']
        u.save()
        est.curso_id = d['curso']
        est.is_active = True
        est.save()
    else:
        usuario = usuarios.nuevo(d['nombres'], d['apellidos'])
        u = User.objects.create_user(username=usuario, password=usuario, first_name=d['nombres'], last_name=d['apellidos'])
        grupo, _ = Group.objects.get_or_create(name='Estudiantes')
        u.groups.add(grupo)
        est = Estudiante.objects.create(user=u, curso_id=d['curso'], colegio=colegio)
        nueva = {'rol': 'Estudiante', 'nombre': f'{d["apellidos"]} {d["nombres"]}', 'usuario': usuario,
                 'clave': usuario, 'detalle': ''}
    ficha, _ = FichaEstudiante.objects.get_or_create(estudiante=est)
    if d['documento']:
        ficha.numero_documento = d['documento']
    if d['tipo_documento']:
        ficha.tipo_documento = d['tipo_documento']
    for k, v in d['ficha'].items():
        setattr(ficha, CAMPO_FICHA.get(k, k), v)
    if d['fecha_nacimiento']:
        ficha.fecha_nacimiento = datetime.date.fromisoformat(d['fecha_nacimiento'])
    if d['rh']:
        ficha.grupo_sanguineo = d['rh']
    if d['correo_acudiente']:
        ficha.email_acudiente = d['correo_acudiente']
    ficha.save()
    if nueva:
        nueva['detalle'] = est.curso.nombre if est.curso_id else ''
    return nueva


# ---------------------------------------------------------------------------
# Plantillas (vacías o con lo que ya hay)
# ---------------------------------------------------------------------------

def _filas_actuales(tipo, colegio):
    from .models import AsignacionDocente, Curso, Docente, Estudiante, Materia
    from .models.academicos import PonderacionAreaMateria
    if tipo == 'cursos':
        for c in Curso.objects.filter(colegio=colegio).select_related('sede', 'director_grado__ficha', 'director_grado__user').order_by('orden', 'nombre'):
            director = ''
            if c.director_grado_id:
                ficha = getattr(c.director_grado, 'ficha', None)
                director = (ficha.numero_documento if ficha and ficha.numero_documento else c.director_grado.user.username)
            yield [c.grado if c.grado is not None else '', c.subgrupo, c.nombre, c.sede.nombre if c.sede else '', director]
    elif tipo == 'docentes':
        for dct in Docente.objects.filter(colegio=colegio).select_related('user', 'ficha').order_by('user__last_name', 'user__first_name'):
            f = getattr(dct, 'ficha', None)
            yield [dct.user.first_name, dct.user.last_name, (f.numero_documento if f else '') or '', dct.user.email,
                   (f.telefono if f else '') or '', (f.titulo_profesional if f else '') or '']
    elif tipo == 'materias':
        areas = {}
        for p in PonderacionAreaMateria.objects.filter(colegio=colegio).select_related('area'):
            areas.setdefault(p.materia_id, p.area.nombre)
        for m in Materia.objects.filter(colegio=colegio).order_by('nombre'):
            yield [m.nombre, areas.get(m.id, ''), m.abreviatura or '']
    elif tipo == 'ponderacion':
        from .models.academicos import PonderacionGrado
        nombres = dict(_grados())
        for p in (PonderacionAreaMateria.objects.filter(colegio=colegio).select_related('area', 'materia')
                  .order_by('area__nombre', 'materia__nombre')):
            yield [p.area.nombre, p.materia.nombre, '', f'{p.peso_porcentual:g}']
        for p in (PonderacionGrado.objects.filter(colegio=colegio).select_related('area', 'materia')
                  .order_by('grado', 'area__nombre', 'materia__nombre')):
            yield [p.area.nombre, p.materia.nombre, nombres.get(p.grado, p.grado), f'{p.peso_porcentual:g}']
    elif tipo == 'asignacion':
        for a in (AsignacionDocente.objects.filter(colegio=colegio)
                  .select_related('docente__user', 'docente__ficha', 'materia', 'curso')
                  .order_by('curso__orden', 'curso__nombre', 'materia__nombre')):
            f = getattr(a.docente, 'ficha', None)
            yield [(f.numero_documento if f and f.numero_documento else a.docente.user.username),
                   a.materia.nombre, a.curso.nombre, a.intensidad_horaria_semanal]
    elif tipo == 'estudiantes':
        for e in (Estudiante.objects.filter(colegio=colegio, is_active=True).select_related('user', 'curso', 'ficha')
                  .order_by('curso__orden', 'curso__nombre', 'user__last_name', 'user__first_name')):
            f = getattr(e, 'ficha', None)
            g = (lambda campo: (getattr(f, campo) or '') if f else '')
            yield [e.user.first_name, e.user.last_name, e.curso.nombre if e.curso else '', g('tipo_documento'),
                   g('numero_documento'), f.fecha_nacimiento.strftime('%d/%m/%Y') if f and f.fecha_nacimiento else '',
                   g('lugar_nacimiento'), g('eps'), g('grupo_sanguineo'), g('enfermedades_alergias'),
                   g('nombre_madre'), g('celular_madre'), g('nombre_padre'), g('celular_padre'),
                   g('nombre_acudiente'), g('celular_acudiente'), g('email_acudiente'),
                   g('colegio_anterior'), g('grado_anterior')]


def plantilla(tipo, colegio, con_datos=False):
    """Bytes del .xlsx: hoja «Datos» para llenar, «Instrucciones» y «Listas» de lo que ya existe."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from .models import Curso, Docente, Materia, Sede
    from .models.academicos import AreaConocimiento

    color = (getattr(colegio, 'color_primario', '') or '#193661').lstrip('#')[:6].upper()
    if not re.fullmatch(r'[0-9A-F]{6}', color):
        color = '193661'
    libro = Workbook()
    ws = libro.active
    ws.title = 'Datos'
    cols = columnas(tipo)
    for i, c in enumerate(cols, start=1):
        celda = ws.cell(1, i, c.titulo + (' *' if c.obligatoria else ''))
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor=color if c.obligatoria else '6C757D')
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(34, len(c.titulo) + 6))
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'A2'
    if con_datos:
        for fila in _filas_actuales(tipo, colegio):
            ws.append([str(v) if v is not None else '' for v in fila])
    # Todo como texto: que Excel no le quite los ceros a «01» ni vuelva fecha un «6-1».
    for i in range(1, len(cols) + 1):
        for fila in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 600), min_col=i, max_col=i):
            for celda in fila:
                celda.number_format = '@'

    # Listas de lo que ya existe, para elegir sin escribir mal.
    listas = libro.create_sheet('Listas')
    valores = {
        'Cursos': [c.nombre for c in Curso.objects.filter(colegio=colegio).order_by('orden', 'nombre')],
        'Materias': [m.nombre for m in Materia.objects.filter(colegio=colegio).order_by('nombre')],
        'Sedes': [s.nombre for s in Sede.objects.filter(colegio=colegio).order_by('nombre')],
        'Docentes (documento)': [],
        'Grados': [f'{g} - {n}' for g, n in _grados()],
        'Tipos de documento': ['TI', 'RC', 'CC', 'CE', 'OT'],
        'Áreas': [a.nombre for a in AreaConocimiento.objects.filter(colegio=colegio).order_by('nombre')],
        'Grado o nivel': ['', 'Preescolar', 'Primaria', 'Secundaria', 'Media'] + [n for _, n in _grados()],
    }
    for dct in Docente.objects.filter(colegio=colegio).select_related('user', 'ficha').order_by('user__last_name'):
        f = getattr(dct, 'ficha', None)
        valores['Docentes (documento)'].append(
            f'{(f.numero_documento if f and f.numero_documento else dct.user.username)}')
    for j, (titulo, lista) in enumerate(valores.items(), start=1):
        listas.cell(1, j, titulo).font = Font(bold=True)
        listas.column_dimensions[get_column_letter(j)].width = 26
        for k, v in enumerate(lista, start=2):
            listas.cell(k, j, v)

    def lista_de(nombre):
        j = list(valores).index(nombre) + 1
        n = len(valores[nombre])
        letra = get_column_letter(j)
        return f'=Listas!${letra}$2:${letra}${max(n, 1) + 1}' if n else None

    desplegables = {'curso': 'Cursos', 'materia': 'Materias', 'sede': 'Sedes', 'tipo_documento': 'Tipos de documento',
                    'area': 'Áreas'}
    if tipo == 'ponderacion':
        desplegables['grado'] = 'Grado o nivel'
    if tipo == 'cursos':
        desplegables.pop('curso')
    for i, c in enumerate(cols, start=1):
        fuente = desplegables.get(c.clave)
        if fuente and lista_de(fuente) and not (tipo == 'asignacion' and c.clave == 'curso'):
            dv = DataValidation(type='list', formula1=lista_de(fuente), allow_blank=True, showErrorMessage=False)
            dv.add(f'{get_column_letter(i)}2:{get_column_letter(i)}1000')
            ws.add_data_validation(dv)

    ins = libro.create_sheet('Instrucciones')
    ins.column_dimensions['A'].width = 28
    ins.column_dimensions['B'].width = 13
    ins.column_dimensions['C'].width = 70
    ins.column_dimensions['D'].width = 26
    ins.append([f'Importar {TIPOS[tipo]["titulo"].lower()} · {colegio.nombre}'])
    ins['A1'].font = Font(bold=True, size=13)
    ins.append(['Llene la hoja «Datos», una fila por registro. Las columnas con * son obligatorias. '
                'Si una fila ya existe en la plataforma, se actualiza; no se duplica.'])
    ins.append([])
    ins.append(['Columna', '¿Obligatoria?', 'Qué escribir', 'Ejemplo'])
    for celda in ins[4]:
        celda.font = Font(bold=True)
    for c in cols:
        ins.append([c.titulo, 'Sí' if c.obligatoria else 'No', c.ayuda, c.ejemplo])
    libro.move_sheet('Instrucciones', offset=-libro.sheetnames.index('Instrucciones'))
    libro.active = libro.sheetnames.index('Datos')
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def credenciales_xlsx(credenciales, colegio):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    libro = Workbook()
    ws = libro.active
    ws.title = 'Usuarios nuevos'
    ws.append([f'{colegio.nombre} · usuarios creados en la importación'])
    ws['A1'].font = Font(bold=True, size=12)
    ws.append(['La contraseña inicial es igual al usuario. Pídales cambiarla al entrar.'])
    ws.append([])
    ws.append(['Rol', 'Nombre', 'Usuario', 'Contraseña inicial', 'Curso / documento'])
    for celda in ws[4]:
        celda.font = Font(bold=True)
    for c in credenciales:
        ws.append([c['rol'], c['nombre'], c['usuario'], c['clave'], c.get('detalle', '')])
    for letra, ancho in zip('ABCDE', (12, 38, 24, 20, 20)):
        ws.column_dimensions[letra].width = ancho
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
