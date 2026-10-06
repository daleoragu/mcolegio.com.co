# notas/planillas/excel.py
"""La planilla en Excel: se descarga con el diseño del colegio y se vuelve a subir.

Diseño (basado en la planilla que usaba el docente):
  fila 1   escudo + nombre del colegio, en el color del colegio
  fila 2   docente · asignatura · curso · periodo
  fila 3   N. | Estudiante | COMPONENTE (peso %) … | Final | Desempeño | Fallas
  fila 4   nombre de cada nota (se puede cambiar) | Def
  fila 5+  un estudiante por fila

La definitiva de cada componente, la final y el desempeño son fórmulas que
siguen la misma regla de la plataforma (notas/planillas/guardar.py), así el
docente ve en Excel el mismo número que verá en línea.

Cada hoja lleva, en una columna oculta, de qué asignación y periodo es y
dónde está cada cosa. Con eso se sube de vuelta sin depender del nombre de la
hoja ni de que el docente no mueva nada más.
"""
import io
import json
import re
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..models.academicos import (AsignacionDocente, EscalaValoracion, IndicadorLogroPeriodo,
                                  InasistenciasManualesPeriodo, PeriodoAcademico)
from ..models.perfiles import Estudiante
from .columnas import (columnas_del_plan, componentes_activos, configuracion, limpiar_columnas,
                       notas_guardadas, plan_completo)
from .guardar import CENTESIMA, MAXIMA, MINIMA, a_decimal

MARCA = 'mcolegio-planilla'
VERSION = 1
FILA_DATOS = 5
FUENTE = 'Arial'

# Colores por puesto en la escala, de la más baja a la más alta: los mismos
# tonos que usa la planilla en línea (rojo, amarillo, verde, azul).
TONOS_ESCALA = [('F8D7DA', '842029'), ('FFF3CD', '664D03'), ('D1E7DD', '0F5132'), ('CFE2FF', '084298')]
ESCALA_POR_DEFECTO = [(Decimal('1.0'), Decimal('2.9'), 'BAJO'), (Decimal('3.0'), Decimal('3.9'), 'BÁSICO'),
                      (Decimal('4.0'), Decimal('4.5'), 'ALTO'), (Decimal('4.6'), Decimal('5.0'), 'SUPERIOR')]


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------

def _hex(color, por_defecto='0D6EFD'):
    c = (color or '').strip().lstrip('#')
    if len(c) == 3:
        c = ''.join(ch * 2 for ch in c)
    return c.upper() if re.fullmatch(r'[0-9A-Fa-f]{6}', c or '') else por_defecto


def _mezclar(color, con_blanco):
    """El color aclarado: con_blanco=0.85 deja un 15 % del color."""
    r, g, b = (int(color[i:i + 2], 16) for i in (0, 2, 4))
    r, g, b = (round(v + (255 - v) * con_blanco) for v in (r, g, b))
    return f'{r:02X}{g:02X}{b:02X}'


def _tono(k, total):
    """Color del nivel k de la escala: el más bajo rojo y el más alto azul."""
    if total <= 1:
        return TONOS_ESCALA[-1]
    return TONOS_ESCALA[round(k * (len(TONOS_ESCALA) - 1) / (total - 1))]


def _relleno(color):
    return PatternFill('solid', start_color=color, end_color=color)


def _nombre_hoja(texto, usados):
    base = re.sub(r'[\[\]\:\*\?\/\\]', '', texto).strip()[:31] or 'Planilla'
    nombre, k = base, 2
    while nombre.lower() in usados:
        sufijo = f' ({k})'
        nombre = base[:31 - len(sufijo)] + sufijo
        k += 1
    usados.add(nombre.lower())
    return nombre


def _num(peso):
    """Peso en % a fracción para la fórmula: 33.33 -> '0.3333'."""
    return format((Decimal(peso) / 100).quantize(Decimal('0.0001')).normalize(), 'f')


def escala_del_colegio(colegio):
    filas = list(EscalaValoracion.objects.filter(colegio=colegio).order_by('valor_minimo')
                 .values_list('valor_minimo', 'valor_maximo', 'nombre_desempeno'))
    return [(Decimal(a), Decimal(b), n) for a, b, n in filas] or ESCALA_POR_DEFECTO


def _recortar_bordes(img):
    """Quita el margen blanco o transparente que traen muchos escudos."""
    from PIL import ImageChops
    try:
        if img.mode == 'RGBA':
            caja = img.getchannel('A').point(lambda a: 255 if a > 10 else 0).getbbox()
        else:
            from PIL import Image
            blanco = Image.new('RGB', img.size, (255, 255, 255))
            dif = ImageChops.difference(img.convert('RGB'), blanco).convert('L')
            caja = dif.point(lambda v: 255 if v > 18 else 0).getbbox()
        if caja and (caja[2] - caja[0]) > 8 and (caja[3] - caja[1]) > 8:
            return img.crop(caja)
    except Exception:
        pass
    return img


def _imagen_escudo(colegio, alto=62):
    """El escudo listo para el Excel, o None. Funciona en local y con Spaces."""
    from openpyxl.drawing.image import Image as ImagenExcel
    for campo in ('escudo', 'logo_izquierdo'):
        archivo = getattr(colegio, campo, None)
        if not archivo:
            continue
        try:
            from PIL import Image
            archivo.open('rb')
            try:
                datos = archivo.read()
            finally:
                archivo.close()
            img = Image.open(io.BytesIO(datos))
            img.load()
            if img.mode not in ('RGB', 'RGBA'):
                img = img.convert('RGBA')
            img = _recortar_bordes(img)
            ancho = max(1, round(img.width * alto / img.height))
            img = img.resize((ancho, alto))
            salida = io.BytesIO()
            img.save(salida, format='PNG')
            salida.seek(0)
            return ImagenExcel(salida)
        except Exception:
            continue
    return None


def estudiantes_de(asignacion, periodo):
    """Los estudiantes de esa asignación en ese año (respeta el historial de promoción)."""
    from ..boletin.logic import estudiantes_del_curso_en
    return list(estudiantes_del_curso_en(asignacion.colegio, asignacion.curso, periodo.ano_lectivo))


# ---------------------------------------------------------------------------
# Descargar
# ---------------------------------------------------------------------------

def generar_libro(colegio, periodo, asignaciones):
    """Bytes del .xlsx con una hoja por asignación."""
    config = configuracion(colegio)
    primario = _hex(getattr(colegio, 'color_primario', ''))
    texto_primario = _hex(getattr(colegio, 'color_texto_primario', ''), 'FFFFFF')
    secundario = _hex(getattr(colegio, 'color_secundario', ''), '6C757D')
    escala = escala_del_colegio(colegio)

    libro = Workbook()
    libro.remove(libro.active)

    # Hoja oculta con la escala de valoración: la usa la fórmula de desempeño.
    hoja_escala = libro.create_sheet('Escala')
    hoja_escala.append(['Desde', 'Hasta', 'Desempeño'])
    for a, b, n in escala:
        hoja_escala.append([float(a), float(b), n])
    hoja_escala.sheet_state = 'hidden'
    fin_escala = len(escala) + 1

    usados = {'escala'}
    for asignacion in asignaciones:
        nombre = _nombre_hoja(f'{asignacion.curso.nombre} {asignacion.materia.nombre}', usados)
        hoja = libro.create_sheet(nombre)
        _dibujar_hoja(hoja, colegio, periodo, asignacion, config, escala, fin_escala,
                      primario, texto_primario, secundario)
    if not asignaciones:
        libro.create_sheet('Sin asignaciones')['A1'] = 'No hay asignaturas para este docente.'
    # La primera hoja visible queda activa al abrir.
    for i, h in enumerate(libro.worksheets):
        if h.sheet_state == 'visible':
            libro.active = i
            break

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _dibujar_hoja(ws, colegio, periodo, asignacion, config, escala, fin_escala,
                  primario, texto_primario, secundario):
    estudiantes = estudiantes_de(asignacion, periodo)
    plan = plan_completo(asignacion, periodo, estudiantes, config)
    fallas = dict(InasistenciasManualesPeriodo.objects.filter(
        asignacion=asignacion, periodo=periodo, estudiante__in=estudiantes)
        .values_list('estudiante_id', 'cantidad'))

    fino = Side(style='thin', color=_mezclar(primario, 0.55))
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    centro = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cabecera = Font(name=FUENTE, bold=True, color=texto_primario, size=10)
    relleno_cab = _relleno(primario)
    relleno_def = _relleno(_mezclar(primario, 0.85))
    relleno_final = _relleno(_mezclar(primario, 0.72))
    desbloqueada = Protection(locked=False)

    # --- Columnas -----------------------------------------------------------
    col = 3
    mapa = {}
    for codigo, datos in plan.items():
        ini = col
        fin = col + len(datos['columnas']) - 1
        mapa[codigo] = [ini, fin, fin + 1]
        col = fin + 2
    c_final, c_desemp, c_fallas, c_id = col, col + 1, col + 2, col + 3
    ultima = c_fallas
    L = get_column_letter

    n_est = len(estudiantes)
    fila_fin = FILA_DATOS + max(n_est, 1) - 1

    # --- Filas 1 y 2: encabezado institucional -------------------------------
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ultima)
    t = ws.cell(1, 1, (colegio.nombre or '').upper())
    t.font = Font(name=FUENTE, bold=True, italic=True, size=16, color=primario)
    t.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 46

    docente = asignacion.docente.user.get_full_name() if asignacion.docente_id else ''
    linea = (f'Docente: {docente}   ·   Asignatura: {asignacion.materia.nombre}   ·   '
             f'Curso: {asignacion.curso.nombre}   ·   {periodo.get_nombre_display()} {periodo.ano_lectivo}')
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ultima)
    s = ws.cell(2, 1, linea)
    s.font = Font(name=FUENTE, size=10, color=secundario)
    s.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 20

    escudo = _imagen_escudo(colegio)
    if escudo is not None:
        ws.add_image(escudo, 'A1')

    # --- Filas 3 y 4: títulos -----------------------------------------------
    def titulo(fila, columna, texto):
        celda = ws.cell(fila, columna, texto)
        celda.font, celda.fill, celda.alignment, celda.border = cabecera, relleno_cab, centro, borde
        return celda

    for columna, texto in ((1, 'N.'), (2, 'Estudiante'), (c_final, 'Final'),
                           (c_desemp, 'Desempeño'), (c_fallas, 'Fallas')):
        titulo(3, columna, texto)
        titulo(4, columna, None)
        ws.merge_cells(start_row=3, start_column=columna, end_row=4, end_column=columna)

    largo = 6
    for codigo, datos in plan.items():
        ini, fin, cdef = mapa[codigo]
        peso = format(Decimal(datos['peso']).normalize(), 'f')
        titulo(3, ini, f"{datos['nombre']} ({peso}%)")
        for c in range(ini + 1, cdef + 1):
            titulo(3, c, None)
        ws.merge_cells(start_row=3, start_column=ini, end_row=3, end_column=cdef)
        for k, nombre_nota in enumerate(datos['columnas']):
            celda = titulo(4, ini + k, nombre_nota)
            celda.alignment = Alignment(horizontal='center', vertical='bottom', text_rotation=90, wrap_text=True)
            celda.protection = desbloqueada      # el docente puede renombrar la nota
            largo = max(largo, len(nombre_nota))
        titulo(4, cdef, 'Def')
    ws.row_dimensions[3].height = 22
    ws.row_dimensions[4].height = max(36, min(120, largo * 6.2))

    # --- Datos ---------------------------------------------------------------
    letras_final = []
    for i, est in enumerate(estudiantes):
        f = FILA_DATOS + i
        ws.cell(f, 1, i + 1).alignment = Alignment(horizontal='center')
        nombre_est = f'{est.user.last_name} {est.user.first_name}'.strip().upper() or est.user.username
        ws.cell(f, 2, nombre_est)
        for codigo, datos in plan.items():
            ini, fin, cdef = mapa[codigo]
            valores = datos['valores'].get(est.id, [None] * len(datos['columnas']))
            for k in range(len(datos['columnas'])):
                v = valores[k] if k < len(valores) else None
                celda = ws.cell(f, ini + k, float(v) if v is not None else None)
                celda.number_format = '0.0#'
                celda.alignment = Alignment(horizontal='center')
                celda.protection = desbloqueada
            rango = f'{L(ini)}{f}:{L(fin)}{f}'
            # Solo cuentan las notas entre 1 y 5, como en la plataforma: un 7
            # pegado (la validación no frena lo pegado) no cambia la definitiva.
            validas = f'{rango},">={MINIMA}",{rango},"<={MAXIMA}"'
            dfx = ws.cell(f, cdef, f'=IF(COUNTIFS({validas})=0,"",'
                                   f'ROUND(AVERAGEIFS({rango},{validas}),2))')
            dfx.number_format, dfx.fill = '0.00', relleno_def
            dfx.font = Font(name=FUENTE, bold=True, size=10)
            dfx.alignment = Alignment(horizontal='center')
        defs = [f'{L(mapa[c][2])}{f}' for c in plan]
        if defs:
            suma = '+'.join(f'N({d})*{_num(plan[c]["peso"])}' for d, c in zip(defs, plan))
            formula = f'=IF(COUNT({",".join(defs)})=0,"",ROUND({suma},2))'
        else:
            formula = ''
        fin_c = ws.cell(f, c_final, formula or None)
        fin_c.number_format, fin_c.fill = '0.00', relleno_final
        fin_c.font = Font(name=FUENTE, bold=True, size=11)
        fin_c.alignment = Alignment(horizontal='center')
        letra_final = f'{L(c_final)}{f}'
        desempeno = ws.cell(f, c_desemp,
                            f'=IF({letra_final}="","",IFERROR(LOOKUP({letra_final},Escala!$A$2:$A${fin_escala},'
                            f'Escala!$C$2:$C${fin_escala}),""))')
        desempeno.alignment = Alignment(horizontal='center')
        desempeno.font = Font(name=FUENTE, bold=True, size=9)
        fa = ws.cell(f, c_fallas, fallas.get(est.id) or None)
        fa.alignment = Alignment(horizontal='center')
        fa.protection = desbloqueada
        ws.cell(f, c_id, est.id)

    # Bordes y fuente en toda la tabla
    for fila in ws.iter_rows(min_row=FILA_DATOS, max_row=fila_fin, min_col=1, max_col=ultima):
        for celda in fila:
            celda.border = borde
            if celda.font is None or celda.font.name != FUENTE:
                celda.font = Font(name=FUENTE, size=10, bold=celda.font.b if celda.font else False)

    # --- Validación: solo notas de 1 a 5 y fallas enteras ---------------------
    if n_est:
        dv = DataValidation(type='decimal', operator='between', formula1=str(MINIMA), formula2=str(MAXIMA),
                            allow_blank=True, showErrorMessage=True, errorTitle='Nota fuera de rango',
                            error=f'Las notas van de {MINIMA} a {MAXIMA}. Use coma o punto para los decimales.')
        for codigo in plan:
            ini, fin, _ = mapa[codigo]
            dv.add(f'{L(ini)}{FILA_DATOS}:{L(fin)}{fila_fin}')
        ws.add_data_validation(dv)
        dvf = DataValidation(type='whole', operator='between', formula1='0', formula2='999', allow_blank=True,
                             showErrorMessage=True, errorTitle='Fallas', error='Escriba un número entero de fallas.')
        dvf.add(f'{L(c_fallas)}{FILA_DATOS}:{L(c_fallas)}{fila_fin}')
        ws.add_data_validation(dvf)

        # --- Colores del desempeño, con la escala del colegio -----------------
        rango_final = f'{L(c_final)}{FILA_DATOS}:{L(c_desemp)}{fila_fin}'
        ref = f'${L(c_final)}{FILA_DATOS}'
        for k, (a, b, _) in enumerate(escala):
            fondo, letra = _tono(k, len(escala))
            ws.conditional_formatting.add(rango_final, FormulaRule(
                formula=[f'AND(ISNUMBER({ref}),{ref}>={a},{ref}<={b})'],
                fill=_relleno(fondo), font=Font(color=letra, bold=True)))
        # Casillas de nota que la plataforma no va a aceptar: en rojo.
        for codigo in plan:
            ini, fin, _ = mapa[codigo]
            primera = f'{L(ini)}{FILA_DATOS}'
            ws.conditional_formatting.add(f'{L(ini)}{FILA_DATOS}:{L(fin)}{fila_fin}', FormulaRule(
                formula=[f'AND({primera}<>"",OR(NOT(ISNUMBER({primera})),{primera}<{MINIMA},{primera}>{MAXIMA}))'],
                fill=_relleno('F8D7DA'), font=Font(color='842029', bold=True)))
        aprobar = escala[1][0] if len(escala) >= 2 else Decimal('3.0')
        for codigo in plan:
            cdef = mapa[codigo][2]
            ws.conditional_formatting.add(f'{L(cdef)}{FILA_DATOS}:{L(cdef)}{fila_fin}', CellIsRule(
                operator='lessThan', formula=[str(aprobar)], font=Font(color='C00000', bold=True)))

    # --- Leyenda --------------------------------------------------------------
    fila_ley = fila_fin + 2
    ws.merge_cells(start_row=fila_ley, start_column=1, end_row=fila_ley, end_column=ultima)
    ley = ws.cell(fila_ley, 1,
                  f'Escriba notas de {MINIMA} a {MAXIMA} en las casillas blancas. «Def», «Final» y «Desempeño» '
                  f'se calculan solos, igual que en la plataforma. Puede cambiar el nombre de cada nota en la fila '
                  f'de títulos. Para guardarlo en la plataforma: Mis planillas › Subir Excel.')
    ley.font = Font(name=FUENTE, italic=True, size=9, color=secundario)
    ley.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[fila_ley].height = 30

    # --- Datos ocultos para subirla de vuelta --------------------------------
    meta = {'marca': MARCA, 'v': VERSION, 'asignacion': asignacion.id, 'periodo': periodo.id,
            'fila': FILA_DATOS, 'comp': mapa, 'fallas': c_fallas, 'id': c_id}
    ws.cell(3, c_id, 'ID')
    ws.cell(4, c_id, json.dumps(meta))
    ws.column_dimensions[L(c_id)].hidden = True

    # --- Medidas, impresión y protección --------------------------------------
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 40
    for codigo in plan:
        ini, fin, cdef = mapa[codigo]
        for c in range(ini, fin + 1):
            ws.column_dimensions[L(c)].width = 5.6
        ws.column_dimensions[L(cdef)].width = 6.6
    ws.column_dimensions[L(c_final)].width = 7.4
    ws.column_dimensions[L(c_desemp)].width = 12
    ws.column_dimensions[L(c_fallas)].width = 7
    ws.freeze_panes = f'C{FILA_DATOS}'
    ws.sheet_properties.tabColor = primario
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '3:4'
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    # Protegida sin clave: evita borrar una fórmula por accidente, pero el
    # docente puede quitarla (Revisar › Desproteger hoja) si lo necesita.
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.formatCells = False


# ---------------------------------------------------------------------------
# Subir
# ---------------------------------------------------------------------------

class ErrorDeArchivo(Exception):
    pass


def _meta_de(ws):
    """Busca la marca de mColegio en la fila 4 (la columna oculta)."""
    for celda in ws[4]:
        v = celda.value
        if isinstance(v, str) and MARCA in v:
            try:
                meta = json.loads(v)
            except ValueError:
                return None
            if meta.get('marca') == MARCA:
                return meta
    return None


def _a_nota(valor):
    """(Decimal|None, error|None) para lo que haya en una casilla.

    Una nota escrita como texto («4,5» pegado desde otro lado) no se acepta:
    el Excel tampoco la cuenta en su definitiva, y lo que el docente ve en el
    Excel tiene que ser lo que se guarda.
    """
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None, None
    if isinstance(valor, str):
        if a_decimal(valor) is not None:
            return None, f'«{valor}» está guardado como texto y el Excel no lo cuenta; escríbalo de nuevo'
        return None, f'«{valor}» no es un número'
    v = a_decimal(valor)
    if v is None:
        return None, f'«{valor}» no es un número'
    if not (MINIMA <= v <= MAXIMA):
        return None, f'«{format(v.normalize(), "f")}» no está entre {MINIMA} y {MAXIMA}'
    return v.quantize(CENTESIMA, rounding=ROUND_HALF_UP), None


def _definitiva(notas_por_comp, pesos):
    """La misma cuenta de guardar.py, sin tocar la base: para la vista previa."""
    total = Decimal('0')
    for codigo, peso in pesos.items():
        validas = [v for _, v in notas_por_comp.get(codigo, [])]
        prom = (sum(validas) / len(validas)).quantize(CENTESIMA, rounding=ROUND_HALF_UP) if validas else Decimal('0')
        total += prom * Decimal(peso) / 100
    return total.quantize(CENTESIMA, rounding=ROUND_HALF_UP)


def leer_libro(archivo, colegio, usuario):
    """Lee el Excel subido y devuelve lo que pasaría al guardarlo. No guarda nada.

    Devuelve {'hojas': [...], 'ignoradas': [...]}. Cada hoja trae sus columnas,
    los estudiantes con sus notas válidas, los errores por casilla y si algo
    cambia frente a lo que hay en la plataforma.
    """
    try:
        libro = load_workbook(archivo, data_only=True, read_only=False)
    except Exception:
        raise ErrorDeArchivo('El archivo no se pudo abrir como Excel (.xlsx).')

    hojas, ignoradas = [], []
    for ws in libro.worksheets:
        if ws.sheet_state != 'visible':
            continue
        meta = _meta_de(ws)
        if meta is None:
            ignoradas.append(ws.title)
            continue
        hojas.append(_leer_hoja(ws, meta, colegio, usuario))
    if not hojas:
        raise ErrorDeArchivo('El archivo no tiene planillas de mColegio. Descárguelo desde «Mis planillas» '
                             'y súbalo con la misma estructura.')
    return {'hojas': hojas, 'ignoradas': ignoradas}


def _leer_hoja(ws, meta, colegio, usuario):
    hoja = {'titulo': ws.title, 'asignacion': meta.get('asignacion'), 'periodo': meta.get('periodo'),
            'estado': 'ok', 'error': '', 'estudiantes': [], 'columnas': {}, 'errores': [],
            'cambian': 0, 'notas': 0}
    asignacion = AsignacionDocente.objects.filter(id=meta.get('asignacion'), colegio=colegio) \
        .select_related('materia', 'curso', 'docente__user').first()
    periodo = PeriodoAcademico.objects.filter(id=meta.get('periodo'), colegio=colegio).first()
    if asignacion is None or periodo is None:
        hoja.update(estado='error', error='Esta hoja es de otro colegio o de una asignatura que ya no existe.')
        return hoja
    hoja['nombre'] = f'{asignacion.curso.nombre} · {asignacion.materia.nombre}'
    hoja['periodo_nombre'] = f'{periodo.get_nombre_display()} {periodo.ano_lectivo}'
    if not (usuario.is_superuser or (asignacion.docente_id and asignacion.docente.user_id == usuario.id)):
        hoja.update(estado='error', error='Esta asignatura no es suya.')
        return hoja
    if not periodo.esta_activo:
        hoja.update(estado='error', error=f'El {hoja["periodo_nombre"]} está cerrado para notas.')
        return hoja
    if not IndicadorLogroPeriodo.objects.filter(asignacion=asignacion, periodo=periodo).exists():
        hoja.update(estado='error', error='Faltan los indicadores de logro de este periodo. Defínalos en '
                                          '«Ingresar notas» y vuelva a subir el archivo.')
        return hoja

    activos = {c: p for c, _, p in componentes_activos(asignacion)}
    mapa = {c: v for c, v in (meta.get('comp') or {}).items() if c in activos}
    for codigo, (ini, fin, _) in mapa.items():
        nombres = [ws.cell(4, c).value for c in range(ini, fin + 1)]
        hoja['columnas'][codigo] = limpiar_columnas(nombres)
    faltan = [c for c in activos if c not in mapa]
    if faltan:
        hoja.update(estado='error', error='Los componentes de la asignatura cambiaron desde que se descargó '
                                          'este Excel. Descárguelo de nuevo.')
        return hoja

    validos = {e.id: e for e in estudiantes_de(asignacion, periodo)}
    actuales = {c: notas_guardadas(asignacion, periodo, c, list(validos.values())) for c in mapa}
    fallas_actuales = dict(InasistenciasManualesPeriodo.objects.filter(
        asignacion=asignacion, periodo=periodo).values_list('estudiante_id', 'cantidad'))
    from ..models.academicos import Calificacion
    def_actual = dict(Calificacion.objects.filter(
        materia=asignacion.materia, periodo=periodo, tipo_nota='PROM_PERIODO', colegio=colegio,
        estudiante_id__in=list(validos)).values_list('estudiante_id', 'valor_nota'))

    c_id, c_fallas = meta.get('id'), meta.get('fallas')
    fila = meta.get('fila', FILA_DATOS)
    vistos = set()
    while True:
        id_celda = ws.cell(fila, c_id).value
        nombre = ws.cell(fila, 2).value
        if id_celda in (None, '') and not nombre:
            break
        try:
            est_id = int(id_celda)
        except (TypeError, ValueError):
            fila += 1
            continue
        if est_id not in validos or est_id in vistos:
            hoja['errores'].append(f'Fila {fila}: {nombre or "estudiante"} ya no está en este curso; se omite.')
            fila += 1
            continue
        vistos.add(est_id)
        est = {'id': est_id, 'nombre': nombre or str(validos[est_id]), 'notas': {}, 'fallas': None}
        for codigo, (ini, fin, _) in mapa.items():
            lista = []
            for k, c in enumerate(range(ini, fin + 1)):
                v, error = _a_nota(ws.cell(fila, c).value)
                if error:
                    hoja['errores'].append(f'{ws.cell(fila, c).coordinate} ({est["nombre"]}): {error}; '
                                           f'no se guarda.')
                elif v is not None:
                    lista.append((hoja['columnas'][codigo][k], v))
            est['notas'][codigo] = lista
            hoja['notas'] += len(lista)
        if c_fallas:
            crudo = ws.cell(fila, c_fallas).value
            v = a_decimal(crudo)
            est['fallas'] = int(v) if v is not None and v >= 0 else 0
        # ¿Cambia algo frente a la plataforma?
        cambia = est['fallas'] != (fallas_actuales.get(est_id) or 0)
        for codigo in mapa:
            antes = sorted((d, Decimal(v)) for d, v in actuales[codigo].get(est_id, []))
            despues = sorted(est['notas'][codigo])
            cambia = cambia or antes != despues
        est['cambia'] = cambia
        est['definitiva'] = _definitiva(est['notas'], activos)
        est['definitiva_antes'] = def_actual.get(est_id)
        hoja['cambian'] += 1 if cambia else 0
        hoja['estudiantes'].append(est)
        fila += 1
    # Columnas renombradas o agregadas también son un cambio.
    for codigo in mapa:
        if hoja['columnas'][codigo] != columnas_del_plan(asignacion, periodo, codigo):
            hoja['plan_cambia'] = True
    return hoja


def a_sesion(resultado):
    """Lo leído, en algo que cabe en la sesión (sin Decimal)."""
    def conv(x):
        if isinstance(x, Decimal):
            return str(x)
        if isinstance(x, dict):
            return {k: conv(v) for k, v in x.items()}
        if isinstance(x, (list, tuple)):
            return [conv(v) for v in x]
        return x
    return conv(resultado)


def aplicar(resultado, colegio, usuario):
    """Guarda lo leído. Vuelve a verificar permisos: la sesión no es fuente de verdad."""
    from .columnas import guardar_plan
    from .guardar import guardar_estudiante
    from django.db import transaction

    resumen = []
    for hoja in resultado.get('hojas', []):
        if hoja.get('estado') != 'ok':
            continue
        asignacion = AsignacionDocente.objects.filter(id=hoja['asignacion'], colegio=colegio).first()
        periodo = PeriodoAcademico.objects.filter(id=hoja['periodo'], colegio=colegio).first()
        if asignacion is None or periodo is None or not periodo.esta_activo:
            continue
        if not (usuario.is_superuser or (asignacion.docente_id and asignacion.docente.user_id == usuario.id)):
            continue
        validos = {e.id: e for e in estudiantes_de(asignacion, periodo)}
        guardados = 0
        with transaction.atomic():
            for codigo, columnas in hoja.get('columnas', {}).items():
                guardar_plan(asignacion, periodo, codigo, columnas)
            for est in hoja.get('estudiantes', []):
                e = validos.get(est['id'])
                if e is None or not est.get('cambia'):
                    continue
                notas = {c: [{'descripcion': d, 'valor': v} for d, v in lista]
                         for c, lista in est.get('notas', {}).items()}
                guardar_estudiante(colegio, asignacion, periodo, e, notas, inasistencias=est.get('fallas'))
                guardados += 1
        resumen.append((hoja.get('nombre') or hoja.get('titulo'), guardados))
    return resumen
