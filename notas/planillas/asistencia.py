# notas/planillas/asistencia.py
"""Asistencia en Excel: la plantilla que se descarga y la subida con revisión.

La plantilla trae una hoja por mes, una columna por día hábil del periodo y
una fila por estudiante. Desde esta versión trae escrito lo que ya está
registrado en la plataforma (X, T o AJ), así que al subirla:

  * X o A  -> ausente
  * AJ     -> ausente con excusa (justificada)
  * T      -> llegó tarde
  * vacío  -> asistió. Si en la plataforma tenía una falla ese día, se quita.

Las plantillas descargadas antes de este cambio venían en blanco: con ellas
un vacío no dice nada, así que solo se agregan o cambian fallas y nunca se
quitan (se avisa en la revisión).

Nada se guarda al subir: primero se muestra qué cambia y el docente confirma.
"""
import datetime
from collections import defaultdict

from openpyxl import load_workbook

from ..models.academicos import Asistencia, AsignacionDocente
from ..models.perfiles import Docente, Estudiante
from ..permisos import es_admin_colegio

MARCA = 'mcolegio-asistencia v2'      # celda A9 (fila oculta) de cada hoja nueva
FILA_META, FILA_DATOS = 9, 12
COL_ID, COL_PRIMER_DIA = 2, 3

CODIGOS = {'X': ('A', False), 'A': ('A', False), 'AJ': ('A', True), 'T': ('T', False), 'P': ('P', False)}
NOMBRE_ESTADO = {'A': 'Ausente', 'AJ': 'Ausente con excusa', 'T': 'Tarde', 'P': 'Asistió'}


def codigo_de(registro):
    """Lo que se escribe en la plantilla para un registro guardado ('' si asistió)."""
    if registro.estado == 'A':
        return 'AJ' if registro.justificada else 'X'
    if registro.estado == 'T':
        return 'T'
    return ''


def registros_de(asignacion, fechas):
    """{(estudiante_id, fecha): Asistencia} de esa asignación en esas fechas."""
    if not fechas:
        return {}
    qs = Asistencia.objects.filter(asignacion=asignacion, fecha__in=list(fechas))
    return {(a.estudiante_id, a.fecha): a for a in qs}


def puede_usar(usuario, asignacion):
    """El docente de la asignación, o un administrador de ese colegio."""
    if es_admin_colegio(usuario, asignacion.colegio):
        return True
    return Docente.objects.filter(user=usuario, colegio=asignacion.colegio, pk=asignacion.docente_id).exists()


def _estado_clave(estado, justificada):
    return 'AJ' if (estado == 'A' and justificada) else estado


def leer(archivo, colegio, usuario):
    """Lee el Excel y devuelve qué cambiaría, sin guardar nada.

    {'hojas': [...], 'cambios': [...], 'errores': [...], 'avisos': [...],
     'resumen': {...}, 'asignaciones': [...]}
    """
    errores, avisos = [], []
    try:
        wb = load_workbook(archivo, data_only=True)
    except Exception:
        return {'errores': ['El archivo no se pudo abrir. Súbalo en formato .xlsx (Excel).'], 'cambios': [], 'hojas': [],
                'avisos': [], 'resumen': {}, 'asignaciones': []}

    cambios = []                       # (est_id, asig_id, fecha, nuevo_estado, justificada, antes)
    hojas = []
    asignaciones = {}
    vistos = set()
    for ws in wb.worksheets:
        # 1. Columnas de fechas: fila 9 trae "asignacion_id|AAAA-MM-DD"
        columnas = {}
        for col in range(COL_PRIMER_DIA, ws.max_column + 1):
            meta = ws.cell(FILA_META, col).value
            if not meta or '|' not in str(meta):
                continue
            try:
                a_id, f = str(meta).split('|')
                columnas[col] = (int(a_id), datetime.date.fromisoformat(f.strip()))
            except (ValueError, TypeError):
                continue
        if not columnas:
            continue                   # hoja sin fechas (portada, "Sin datos", etc.)

        nueva = str(ws.cell(FILA_META, 1).value or '').strip() == MARCA
        ids = {a for a, _ in columnas.values()}
        if len(ids) != 1:
            errores.append(f'Hoja «{ws.title}»: mezcla varias asignaturas; descargue la plantilla de nuevo.')
            continue
        a_id = ids.pop()
        asignacion = asignaciones.get(a_id) or AsignacionDocente.objects.filter(
            pk=a_id, colegio=colegio).select_related('curso', 'materia', 'docente__user').first()
        if asignacion is None:
            errores.append(f'Hoja «{ws.title}»: la asignatura de esta plantilla no es de este colegio.')
            continue
        if not puede_usar(usuario, asignacion):
            errores.append(f'Hoja «{ws.title}»: {asignacion.materia} {asignacion.curso} no es una de sus asignaturas.')
            continue
        asignaciones[a_id] = asignacion

        del_curso = {e.id: e for e in Estudiante.objects.filter(
            colegio=colegio, curso=asignacion.curso, is_active=True).select_related('user')}
        guardados = registros_de(asignacion, {f for _, f in columnas.values()})
        hoja = {'titulo': ws.title, 'asignacion': f'{asignacion.curso} · {asignacion.materia}', 'nueva': nueva,
                'dias': len(columnas), 'ausencias': 0, 'tardes': 0}

        for fila in range(FILA_DATOS, ws.max_row + 1):
            est_id = ws.cell(fila, COL_ID).value
            nombre = ws.cell(fila, 1).value
            if est_id in (None, ''):
                continue
            try:
                est_id = int(est_id)
            except (TypeError, ValueError):
                continue
            if est_id not in del_curso:
                avisos.append(f'{nombre or est_id}: ya no está en {asignacion.curso}; su fila no se tuvo en cuenta.')
                continue
            for col, (_, fecha) in columnas.items():
                crudo = ws.cell(fila, col).value
                texto = str(crudo).strip().upper() if crudo is not None else ''
                if texto in ('', '.', '-', '✓', 'P'):
                    texto = ''
                if texto and texto not in CODIGOS:
                    errores.append(f'Hoja «{ws.title}», {nombre}, día {fecha.day}: «{crudo}» no es válido. Use X, T, AJ o deje vacío.')
                    continue
                clave = (est_id, a_id, fecha)
                if clave in vistos:
                    continue
                vistos.add(clave)
                antes = guardados.get((est_id, fecha))
                antes_clave = _estado_clave(antes.estado, antes.justificada) if antes else None
                if texto:
                    estado, just = CODIGOS[texto]
                    if estado == 'A':
                        hoja['ausencias'] += 1
                    elif estado == 'T':
                        hoja['tardes'] += 1
                    nuevo_clave = _estado_clave(estado, just)
                    if nuevo_clave != (antes_clave or 'P') and not (nuevo_clave == 'P' and antes is None):
                        cambios.append((est_id, a_id, fecha.isoformat(), estado, just, antes_clave))
                elif nueva and antes is not None and antes.estado in ('A', 'T'):
                    # vacío en plantilla nueva = asistió: se quita la falla
                    cambios.append((est_id, a_id, fecha.isoformat(), 'P', False, antes_clave))
        if not nueva:
            avisos.append(f'Hoja «{ws.title}»: es una plantilla anterior (venía en blanco). Se agregan las fallas marcadas, '
                          'pero las casillas vacías no quitan fallas que ya estén en la plataforma.')
        hojas.append(hoja)

    if not hojas and not errores:
        errores.append('El archivo no es una plantilla de asistencia de la plataforma (no tiene las fechas ocultas). '
                       'Descárguela desde «Consultar asistencia».')

    # Para mostrar: nombres y resumen
    nombres = {e.id: f'{e.user.last_name} {e.user.first_name}'.strip() for e in Estudiante.objects.filter(
        id__in={c[0] for c in cambios}).select_related('user')}
    filas = []
    for est_id, a_id, f, estado, just, antes in sorted(cambios, key=lambda c: (nombres.get(c[0], ''), c[2])):
        filas.append({'estudiante': nombres.get(est_id, est_id), 'asignacion': str(asignaciones[a_id].materia),
                      'curso': str(asignaciones[a_id].curso), 'fecha': datetime.date.fromisoformat(f).strftime('%d/%m/%Y'),
                      'antes': NOMBRE_ESTADO.get(antes, 'Sin registro') if antes else 'Sin registro',
                      'ahora': NOMBRE_ESTADO[_estado_clave(estado, just)]})
    por_tipo = defaultdict(int)
    for c in cambios:
        por_tipo['quitadas' if c[3] == 'P' else 'marcadas'] += 1
    return {'hojas': hojas, 'cambios': [list(c) for c in cambios], 'filas': filas, 'errores': errores,
            'avisos': avisos, 'resumen': dict(por_tipo), 'asignaciones': sorted(asignaciones)}


def aplicar(cambios, colegio, usuario):
    """Guarda los cambios ya revisados. Devuelve cuántos registros tocó."""
    permitidas = {a.id: a for a in AsignacionDocente.objects.filter(colegio=colegio, id__in={c[1] for c in cambios})
                  if puede_usar(usuario, a)}
    n = 0
    for est_id, a_id, fecha, estado, just, _antes in cambios:
        asignacion = permitidas.get(a_id)
        if asignacion is None:
            continue
        if not Estudiante.objects.filter(pk=est_id, colegio=colegio, curso=asignacion.curso).exists():
            continue
        Asistencia.objects.update_or_create(
            colegio=colegio, estudiante_id=est_id, asignacion=asignacion,
            fecha=datetime.date.fromisoformat(fecha),
            defaults={'estado': estado, 'justificada': bool(just) and estado == 'A'})
        n += 1
    return n
