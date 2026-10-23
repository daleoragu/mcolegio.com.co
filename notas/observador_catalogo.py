"""Catálogo de faltas (Tipo I, II y III) y observaciones subidas desde Excel.

- El colegio arma su catálogo a mano, desde un Excel o partiendo de unos ejemplos
  generales del Decreto 1965 de 2013 (que debe ajustar a su manual de convivencia).
- El docente busca en el catálogo al escribir una observación.
- Sin internet, el docente llena un Excel con las observaciones del curso (fecha,
  falta, descripción y descargo del estudiante) y lo sube después.
"""
import datetime
import io
import re

from django.db import transaction
from django.utils import timezone

from .models import FaltaCatalogo, RegistroObservador

TIPOS = {'I': 'I', '1': 'I', 'TIPO I': 'I', 'TIPO 1': 'I', 'UNO': 'I',
         'II': 'II', '2': 'II', 'TIPO II': 'II', 'TIPO 2': 'II', 'DOS': 'II',
         'III': 'III', '3': 'III', 'TIPO III': 'III', 'TIPO 3': 'III', 'TRES': 'III'}

# Ejemplos generales (Decreto 1965 de 2013, art. 40). Cada colegio los ajusta a su manual.
EJEMPLOS = [
    ('I', 'Llegar tarde al colegio o a clase sin justificación.'),
    ('I', 'Interrumpir la clase con charlas, ruidos o juegos que afectan el trabajo de los compañeros.'),
    ('I', 'Usar el celular u otros aparatos en clase sin autorización del docente.'),
    ('I', 'No portar el uniforme según lo establecido en el manual de convivencia.'),
    ('I', 'Usar vocabulario soez o irrespetuoso con compañeros o docentes, sin que llegue a agresión.'),
    ('I', 'Conflicto manejado de forma inadecuada (discusión, empujón esporádico) sin daño al cuerpo ni a la salud.'),
    ('I', 'Arrojar basuras o dañar por descuido el aseo y la presentación de los espacios del colegio.'),
    ('I', 'Salir del aula o de la institución sin autorización.'),
    ('II', 'Agresión física, verbal o gestual repetida contra un compañero (acoso escolar), sin características de delito.'),
    ('II', 'Ciberacoso: publicar o difundir por redes sociales mensajes, fotos o videos que ofenden a otra persona.'),
    ('II', 'Agresión que causa daño al cuerpo o a la salud sin generar incapacidad.'),
    ('II', 'Reincidencia en situaciones Tipo I después del proceso pedagógico.'),
    ('II', 'Discriminar o excluir a un compañero por su condición, origen, creencias u orientación.'),
    ('II', 'Dañar de forma intencional bienes de la institución o de los compañeros.'),
    ('III', 'Agresión que constituye presunto delito contra la libertad, integridad y formación sexual.'),
    ('III', 'Agresión física que causa incapacidad médica o lesiones personales.'),
    ('III', 'Porte, consumo o distribución de sustancias psicoactivas dentro de la institución.'),
    ('III', 'Porte de armas o elementos que pongan en peligro la integridad de las personas.'),
    ('III', 'Hurto, extorsión, amenazas u otra conducta que constituya presunto delito.'),
]


def tipo_de(texto):
    t = re.sub(r'\s+', ' ', str(texto or '').strip().upper())
    return TIPOS.get(t, '')


def cargar_ejemplos(colegio):
    """Agrega los ejemplos que el colegio todavía no tenga. Devuelve cuántos agregó."""
    ya = {d.strip().lower() for d in FaltaCatalogo.objects.filter(colegio=colegio).values_list('descripcion', flat=True)}
    nuevas = [FaltaCatalogo(colegio=colegio, tipo_falta=t, descripcion=d) for t, d in EJEMPLOS if d.lower() not in ya]
    FaltaCatalogo.objects.bulk_create(nuevas)
    return len(nuevas)


def catalogo_para_buscar(colegio):
    """[{id, tipo, numeral, descripcion}] de las faltas activas, para el buscador del formulario."""
    return [{'id': f.id, 'tipo': f.tipo_falta, 'numeral': f.numeral, 'descripcion': f.descripcion}
            for f in FaltaCatalogo.objects.filter(colegio=colegio, activa=True)]


# ----------------------------------------------------------------------------- Excel

def _libro():
    from openpyxl import Workbook
    return Workbook()


def _estilo_encabezado(ws, ancho):
    from openpyxl.styles import Alignment, Font, PatternFill
    for celda in ws[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='3730A3')
        celda.alignment = Alignment(vertical='center', wrap_text=True)
    for letra, w in ancho.items():
        ws.column_dimensions[letra].width = w
    ws.freeze_panes = 'A2'


def _bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def plantilla_catalogo(colegio=None):
    """Excel con las columnas del catálogo (y lo que el colegio ya tenga, para editarlo y volverlo a subir)."""
    wb = _libro()
    ws = wb.active
    ws.title = 'Catálogo'
    ws.append(['Tipo (I, II o III)', 'Numeral del manual', 'Descripción'])
    filas = FaltaCatalogo.objects.filter(colegio=colegio) if colegio else []
    for f in filas:
        ws.append([f.tipo_falta, f.numeral, f.descripcion])
    if not filas:
        ws.append(['I', 'Art. 40, num. 1', 'Llegar tarde a clase sin justificación.'])
    _estilo_encabezado(ws, {'A': 16, 'B': 22, 'C': 90})
    return _bytes(wb)


def _leer(archivo):
    from openpyxl import load_workbook
    try:
        return load_workbook(archivo, data_only=True, read_only=True)
    except Exception:
        return None


def importar_catalogo(colegio, archivo):
    """Agrega las faltas del Excel (sin repetir las que ya están). Devuelve (agregadas, errores)."""
    wb = _leer(archivo)
    if wb is None:
        return 0, ['El archivo no es un Excel (.xlsx) válido.']
    ws = wb.worksheets[0]
    ya = {d.strip().lower() for d in FaltaCatalogo.objects.filter(colegio=colegio).values_list('descripcion', flat=True)}
    nuevas, errores = [], []
    for n, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        fila = list(fila) + [None] * 3
        tipo, numeral, desc = tipo_de(fila[0]), str(fila[1] or '').strip()[:40], str(fila[2] or '').strip()
        if not any(fila[:3]):
            continue
        if not desc:
            errores.append(f'Fila {n}: falta la descripción.')
            continue
        if not tipo:
            errores.append(f'Fila {n}: el tipo «{fila[0] or ""}» no es I, II ni III.')
            continue
        if desc.lower() in ya:
            continue
        ya.add(desc.lower())
        nuevas.append(FaltaCatalogo(colegio=colegio, tipo_falta=tipo, numeral=numeral, descripcion=desc))
    FaltaCatalogo.objects.bulk_create(nuevas)
    return len(nuevas), errores


COLUMNAS = ['Documento', 'Estudiante', 'Fecha (dd/mm/aaaa)', 'Clase (Académica o Comportamental)',
            'Positiva o A mejorar', 'Tipo de falta (I, II, III)', 'N.º del catálogo', 'Descripción de lo sucedido',
            'Descargo del estudiante']


def plantilla_observaciones(colegio, estudiantes):
    """Excel para llenar sin internet: una fila por estudiante (se pueden copiar filas para
    anotar varias cosas al mismo estudiante) y una hoja con el catálogo de faltas."""
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = _libro()
    ws = wb.active
    ws.title = 'Observaciones'
    ws.append(COLUMNAS)
    for e in estudiantes:
        ficha = getattr(e, 'ficha', None)
        doc = getattr(ficha, 'numero_documento', '') or ''
        ws.append([doc or e.user.username, f'{e.user.last_name} {e.user.first_name}'.strip().upper()])
    _estilo_encabezado(ws, {'A': 14, 'B': 34, 'C': 14, 'D': 18, 'E': 14, 'F': 12, 'G': 12, 'H': 60, 'I': 45})
    ultima = max(ws.max_row, 2) + 200
    for col, opciones in (('D', 'Académica,Comportamental'), ('E', 'Positiva,A mejorar'), ('F', 'I,II,III')):
        dv = DataValidation(type='list', formula1=f'"{opciones}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f'{col}2:{col}{ultima}')

    cat = wb.create_sheet('Catálogo de faltas')
    cat.append(['N.º del catálogo', 'Tipo', 'Numeral', 'Descripción'])
    for f in FaltaCatalogo.objects.filter(colegio=colegio, activa=True):
        cat.append([f.id, f.tipo_falta, f.numeral, f.descripcion])
    _estilo_encabezado(cat, {'A': 16, 'B': 8, 'C': 18, 'D': 90})

    ayuda = wb.create_sheet('Instrucciones')
    for linea in (
        'Cómo llenar el observador sin internet',
        '',
        '1. En la hoja «Observaciones» escriba solo en las filas de los estudiantes a los que les va a anotar algo.',
        '   Las filas que deje sin descripción ni número del catálogo no se suben.',
        '2. Para anotar dos cosas al mismo estudiante copie su fila (con el documento) y péguela abajo.',
        '3. La fecha es la del día en que pasó (dd/mm/aaaa). No puede ser una fecha futura.',
        '4. Si la falta está en el catálogo, escriba su número (hoja «Catálogo de faltas»): el tipo y la',
        '   descripción se toman de allí. Puede agregar detalles en la descripción.',
        '5. En «Descargo del estudiante» escriba lo que el estudiante dijo en su defensa (si lo hubo).',
        '6. Guarde el archivo y súbalo en la plataforma: Observador → Subir desde Excel.',
    ):
        ayuda.append([linea])
    ayuda.column_dimensions['A'].width = 110
    return _bytes(wb)


def _fecha(valor):
    if isinstance(valor, datetime.datetime):
        return valor.date()
    if isinstance(valor, datetime.date):
        return valor
    t = str(valor or '').strip()
    for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d/%m/%y', '%d.%m.%Y'):
        try:
            return datetime.datetime.strptime(t, fmt).date()
        except ValueError:
            pass
    return None


def _clase(texto):
    t = str(texto or '').strip().lower()
    if t.startswith('acad'):
        return 'ACADEMICA'
    if t.startswith('comp') or t.startswith('conv'):
        return 'COMPORTAMENTAL'
    return ''


def _subtipo(texto):
    t = str(texto or '').strip().lower()
    if t.startswith('pos'):
        return 'POSITIVA'
    if t.startswith('a mej') or t.startswith('mej') or t.startswith('neg'):
        return 'NEGATIVA'
    return None


def _clave_doc(v):
    t = str(v if v is not None else '').strip()
    if re.fullmatch(r'\d+\.0', t):          # Excel guarda 1110555 como 1110555.0
        t = t[:-2]
    return t.replace('.', '').replace(' ', '').upper()


def importar_observaciones(colegio, docente, archivo, estudiantes):
    """Crea las observaciones del Excel para los `estudiantes` permitidos.

    Devuelve (registros creados, errores). Si alguna fila tiene error no se guarda
    ninguna: el docente corrige el archivo y lo vuelve a subir completo.
    """
    wb = _leer(archivo)
    if wb is None:
        return [], ['El archivo no es un Excel (.xlsx) válido.']
    ws = wb['Observaciones'] if 'Observaciones' in wb.sheetnames else wb.worksheets[0]
    por_doc = {}
    for e in estudiantes:
        ficha = getattr(e, 'ficha', None)
        for clave in (getattr(ficha, 'numero_documento', '') or '', e.user.username):
            if clave:
                por_doc.setdefault(_clave_doc(clave), e)
    catalogo = {f.id: f for f in FaltaCatalogo.objects.filter(colegio=colegio)}
    hoy = timezone.localdate()
    nuevos, errores = [], []
    for n, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        fila = list(fila) + [None] * len(COLUMNAS)
        doc, _, fecha, clase, subtipo, tipo_falta, num_cat, desc, descargo = fila[:9]
        desc = str(desc or '').strip()
        descargo = str(descargo or '').strip()
        if not desc and num_cat in (None, ''):
            continue                         # fila sin observación: el estudiante no tuvo nada
        e = por_doc.get(_clave_doc(doc))
        if e is None:
            errores.append(f'Fila {n}: el documento «{doc or ""}» no es de un estudiante de sus cursos.')
            continue
        f = _fecha(fecha)
        if f is None:
            errores.append(f'Fila {n}: escriba la fecha como dd/mm/aaaa.')
            continue
        if f > hoy:
            errores.append(f'Fila {n}: la fecha {f:%d/%m/%Y} todavía no ha llegado.')
            continue
        falta = None
        if num_cat not in (None, ''):
            try:
                falta = catalogo.get(int(float(str(num_cat).strip())))
            except ValueError:
                falta = None
            if falta is None:
                errores.append(f'Fila {n}: el número {num_cat} no está en el catálogo de faltas.')
                continue
        tf = tipo_de(tipo_falta) or (falta.tipo_falta if falta else '')
        if tipo_falta not in (None, '') and not tipo_de(tipo_falta):
            errores.append(f'Fila {n}: el tipo de falta «{tipo_falta}» no es I, II ni III.')
            continue
        texto = desc
        if falta:
            texto = falta.descripcion + (f'\n{desc}' if desc and desc != falta.descripcion else '')
        tipo = _clase(clase) or ('COMPORTAMENTAL' if (falta or tf) else '')
        if not tipo:
            errores.append(f'Fila {n}: escriba si es Académica o Comportamental.')
            continue
        st = _subtipo(subtipo) or ('NEGATIVA' if (falta or tf) else None)
        r = RegistroObservador(colegio=colegio, estudiante=e, docente_reporta=docente, fecha_suceso=f,
                               tipo=tipo, subtipo=st, descripcion=texto, falta=falta, tipo_falta=tf)
        if descargo:
            r.descargo_estudiante = descargo
            r.fecha_descargo = timezone.now()
        nuevos.append(r)
    if errores:
        return [], errores
    with transaction.atomic():
        for r in nuevos:
            r.save()
    return nuevos, []
