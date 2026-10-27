"""Sábana del reporte parcial: un curso y un periodo, estudiantes en filas y
asignaturas en columnas. Sin notas ni promedios: solo la marca de «va
perdiendo» que puso cada docente y los conteos.

La usan el director de grupo y la administración (vista, Excel y PDF)."""
from collections import OrderedDict

from .models import AsignacionDocente, Estudiante, ReporteParcial

# A partir de cuántas asignaturas reportadas el estudiante sale resaltado.
LIMITE_ALERTA = 3


def calcular(curso, periodo):
    asignaciones = list(AsignacionDocente.objects.filter(curso=curso, colegio=curso.colegio)
                        .select_related('materia', 'docente__user')
                        .order_by('materia__nombre', 'id'))
    # Una columna por asignatura, aunque la dicten dos docentes.
    materias = OrderedDict()
    for a in asignaciones:
        m = materias.setdefault(a.materia_id, {
            'id': a.materia_id, 'nombre': a.materia.nombre,
            'corto': (a.materia.abreviatura or a.materia.nombre[:4]).upper(),
            'docentes': [], 'reportada': False, 'total': 0})
        nombre = a.docente.user.get_full_name() or a.docente.user.username
        if nombre not in m['docentes']:
            m['docentes'].append(nombre)

    reportes = (ReporteParcial.objects.filter(periodo=periodo, asignacion__curso=curso)
                .values('estudiante_id', 'asignacion__materia_id', 'presenta_dificultades', 'sugerencia'))
    marcas = {}
    for r in reportes:
        m = materias.get(r['asignacion__materia_id'])
        if m is None:
            continue
        m['reportada'] = True       # el docente ya guardó su reporte de este curso
        if r['presenta_dificultades']:
            marcas[(r['estudiante_id'], m['id'])] = r['sugerencia'] or ''

    estudiantes = (Estudiante.objects.filter(curso=curso, is_active=True).select_related('user')
                   .order_by('user__last_name', 'user__first_name'))
    filas = []
    for n, est in enumerate(estudiantes, 1):
        celdas = []
        for m in materias.values():
            marcada = (est.id, m['id']) in marcas
            if marcada:
                m['total'] += 1
            celdas.append({'marcada': marcada, 'sugerencia': marcas.get((est.id, m['id']), '')})
        total = sum(1 for c in celdas if c['marcada'])
        filas.append({'n': n, 'est': est,
                      'nombre': f"{est.user.last_name} {est.user.first_name}".strip() or est.user.username,
                      'inclusion': bool(getattr(est, 'es_inclusion', False)),
                      'celdas': celdas, 'total': total, 'alerta': total >= LIMITE_ALERTA})

    materias = list(materias.values())
    return {
        'materias': materias,
        'filas': filas,
        'total_reportes': sum(f['total'] for f in filas),
        'estudiantes_reportados': sum(1 for f in filas if f['total']),
        'estudiantes_alerta': sum(1 for f in filas if f['alerta']),
        'estudiantes_sin_reporte': sum(1 for f in filas if not f['total']),
        'pendientes': [m for m in materias if not m['reportada']],
        'limite_alerta': LIMITE_ALERTA,
    }


def excel(curso, periodo, datos, colegio):
    """Libro de Excel con la sábana; devuelve los bytes."""
    from io import BytesIO

    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Reporte {curso.nombre}"[:31]
    fino = Side(style='thin', color='999999')
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    centro = Alignment(horizontal='center', vertical='center', wrap_text=True)
    vertical = Alignment(horizontal='center', vertical='bottom', text_rotation=90, wrap_text=True)
    cabeza = PatternFill('solid', start_color='DDEBF7')
    rojo = PatternFill('solid', start_color='F8D7DA')
    alerta = PatternFill('solid', start_color='FFE08A')
    gris = PatternFill('solid', start_color='EEEEEE')

    materias = datos['materias']
    ultima = 3 + len(materias) + 1
    ws.cell(row=1, column=1, value=colegio.nombre).font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"Sábana del reporte parcial · Curso {curso.nombre} · "
                                   f"{periodo.get_nombre_display()} {periodo.ano_lectivo}").font = Font(bold=True, size=12)
    ws.cell(row=3, column=1, value=f"X = el docente lo reporta con dificultades. Total de reportes: "
                                   f"{datos['total_reportes']} · Estudiantes reportados: "
                                   f"{datos['estudiantes_reportados']} de {len(datos['filas'])} · "
                                   f"Con {datos['limite_alerta']} o más asignaturas: {datos['estudiantes_alerta']}")
    for r in (1, 2, 3):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ultima)

    fila = 5
    titulos = ['N.º', 'Estudiante', 'Incl.'] + [m['nombre'] + ('' if m['reportada'] else ' (sin reporte)')
                                               for m in materias] + ['Total']
    for c, t in enumerate(titulos, 1):
        celda = ws.cell(row=fila, column=c, value=t)
        celda.font = Font(bold=True)
        celda.fill = cabeza
        celda.border = borde
        celda.alignment = vertical if 3 < c < ultima else centro
    ws.row_dimensions[fila].height = 120

    for f in datos['filas']:
        fila += 1
        valores = [f['n'], f['nombre'], 'Sí' if f['inclusion'] else ''] + \
                  ['X' if c['marcada'] else '' for c in f['celdas']] + [f['total']]
        for c, v in enumerate(valores, 1):
            celda = ws.cell(row=fila, column=c, value=v)
            celda.border = borde
            celda.alignment = Alignment(vertical='center') if c == 2 else centro
            if 3 < c < ultima and v == 'X':
                celda.fill = rojo
                celda.font = Font(bold=True, color='B42318')
            if c == ultima:
                celda.font = Font(bold=True)
                if f['alerta']:
                    celda.fill = alerta
        if f['alerta']:
            ws.cell(row=fila, column=2).fill = alerta

    fila += 1
    ws.cell(row=fila, column=2, value='Total por asignatura').font = Font(bold=True)
    for i, m in enumerate(materias):
        celda = ws.cell(row=fila, column=4 + i, value=m['total'])
        celda.font = Font(bold=True)
        celda.alignment = centro
    celda = ws.cell(row=fila, column=ultima, value=datos['total_reportes'])
    celda.font = Font(bold=True)
    celda.alignment = centro
    for c in range(1, ultima + 1):
        ws.cell(row=fila, column=c).border = borde
        ws.cell(row=fila, column=c).fill = gris

    if datos['pendientes']:
        fila += 2
        ws.cell(row=fila, column=1, value='Asignaturas sin reporte guardado: ' + ', '.join(
            f"{m['nombre']} ({', '.join(m['docentes'])})" for m in datos['pendientes'])).font = Font(italic=True)

    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 38
    ws.column_dimensions['C'].width = 6
    for c in range(4, ultima):
        ws.column_dimensions[get_column_letter(c)].width = 5
    ws.column_dimensions[get_column_letter(ultima)].width = 8
    ws.freeze_panes = 'D6'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
