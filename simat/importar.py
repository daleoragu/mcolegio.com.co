"""Leer archivos: el reporte de matrícula que se descarga del SIMAT y el Excel
de datos SIMAT que llena el colegio.

Los títulos se reconocen por su nombre (ver formato.py), así que sirve el
reporte de cualquier versión del SIMAT mientras traiga documento, nombres y
grado.
"""
import csv
import datetime
import io
import re

from django.contrib.auth.models import Group, User
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from notas.models import Curso, Estudiante, FichaEstudiante, Sede
from notas.models.perfiles import clave_subgrupo, normalizar_subgrupo

from . import codigos as C
from . import formato
from .models import DatosSimat


class ArchivoInvalido(Exception):
    pass


# ---------------------------------------------------------------------------
# Valores
# ---------------------------------------------------------------------------
def texto(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def documento(v):
    return re.sub(r'[^0-9A-Za-z]', '', texto(v)).upper()


def tipo_documento(v):
    t = texto(v)
    if t.isdigit():
        return t if t in dict(C.TIPO_DOCUMENTO) else ''
    n = C.normalizar(t)
    for clave, codigo in (('NUIP', '7'), ('NIP', '6'), ('TARJETA', '2'), ('TI', '2'), ('REGISTRO', '5'), ('RC', '5'),
                          ('EXTRANJER', '3'), ('CE', '3'), ('CEDULA', '1'), ('CC', '1'), ('CABILDO', '9')):
        if n == clave or (len(clave) > 3 and clave in n):
            return codigo
    return ''


def genero(v):
    n = C.normalizar(texto(v))
    if n in ('F', 'FEMENINO', 'MUJER'):
        return 'F'
    if n in ('M', 'MASCULINO', 'HOMBRE'):
        return 'M'
    return ''


def fecha(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    t = texto(v)
    for f in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%y', '%Y/%m/%d'):
        try:
            return datetime.datetime.strptime(t[:10], f).date()
        except ValueError:
            continue
    return None


def grado(v):
    t = texto(v)
    if re.fullmatch(r'-?\d{1,2}', t):
        g = int(t)
        return g if -2 <= g <= 11 else None
    return C.GRADO_POR_NOMBRE.get(C.normalizar(t))


def jornada(v):
    t = texto(v)
    if t in dict(C.JORNADA):
        return t
    n = C.normalizar(t)
    for clave, codigo in (('COMPLETA', '1'), ('MANANA', '2'), ('TARDE', '3'), ('NOCHE', '4'), ('NOCTURNA', '4'),
                          ('FIN', '5'), ('SABATINA', '5'), ('UNICA', '6')):
        if clave in n:
            return codigo
    return ''


def codigo_de(v, opciones):
    t = texto(v)
    if t in dict(opciones):
        return t
    n = C.normalizar(t)
    for codigo, nombre in opciones:
        if n and C.normalizar(nombre) == n:
            return codigo
    return ''


def si_no(v):
    n = C.normalizar(texto(v))
    if n in ('S', 'SI', '1', 'X', 'TRUE', 'VERDADERO'):
        return True
    if n in ('N', 'NO', '0', 'FALSE', 'FALSO'):
        return False
    return None


def zona(v):
    n = C.normalizar(texto(v))
    return {'1': '1', 'URBANA': '1', 'U': '1', '2': '2', 'RURAL': '2', 'R': '2'}.get(n, '')


def municipio(fila, base):
    """Código de 5 cifras: de una columna de 5, o de departamento (2) + municipio (3)."""
    entero = re.sub(r'\D', '', texto(fila.get(base)))
    if len(entero) == 5:
        return entero
    if len(entero) == 4:
        return entero.zfill(5)
    dpto = re.sub(r'\D', '', texto(fila.get('dpto_' + base[5:])))
    mpio = re.sub(r'\D', '', texto(fila.get(base + '_3')))
    if dpto and mpio:
        return dpto.zfill(2) + mpio.zfill(3)
    return ''


# ---------------------------------------------------------------------------
# Leer el archivo
# ---------------------------------------------------------------------------
def leer_tabla(archivo):
    """Filas del archivo (xlsx o csv) como listas de valores."""
    nombre = (getattr(archivo, 'name', '') or '').lower()
    datos = archivo.read()
    if nombre.endswith(('.xlsx', '.xlsm')):
        try:
            ws = load_workbook(io.BytesIO(datos), data_only=True, read_only=True).worksheets[0]
        except Exception:
            raise ArchivoInvalido('No se pudo abrir el Excel. Guárdelo como .xlsx e inténtelo de nuevo.')
        return [list(f) for f in ws.iter_rows(values_only=True)]
    if nombre.endswith(('.csv', '.txt')):
        for cod in ('utf-8-sig', 'cp1252'):
            try:
                contenido = datos.decode(cod)
                break
            except UnicodeDecodeError:
                continue
        primera = contenido.splitlines()[0] if contenido else ''
        sep = max([';', ',', '\t', '|'], key=primera.count)
        return list(csv.reader(io.StringIO(contenido), delimiter=sep))
    raise ArchivoInvalido('Suba el archivo en Excel (.xlsx) o CSV. Si el SIMAT lo entrega en .xls, ábralo y guárdelo como .xlsx.')


def con_titulos(filas):
    """[{clave: valor}] buscando la fila de títulos entre las primeras 15."""
    for i, fila in enumerate(filas[:15]):
        claves = [formato.clave_de(texto(v)) if v is not None else None for v in fila]
        reconocidas = {c for c in claves if c}
        if 'numero_documento' in reconocidas and len(reconocidas) >= 3:
            salida = []
            for f in filas[i + 1:]:
                if not any(texto(v) for v in f):
                    continue
                salida.append({c: v for c, v in zip(claves, f) if c})
            return salida, sorted(reconocidas), [texto(v) for v, c in zip(fila, claves) if v and not c]
    raise ArchivoInvalido('No encontré la fila de títulos: el archivo debe traer al menos la columna del número de '
                          'documento, los nombres y el grado.')


# ---------------------------------------------------------------------------
# Importar desde el SIMAT
# ---------------------------------------------------------------------------
def _nombres(f):
    a1, a2 = texto(f.get('apellido1')).upper(), texto(f.get('apellido2')).upper()
    n1, n2 = texto(f.get('nombre1')).upper(), texto(f.get('nombre2')).upper()
    if not (a1 or n1) and f.get('apellidos'):
        from .logica import partir
        a1, a2 = partir(texto(f.get('apellidos')))
        n1, n2 = partir(texto(f.get('nombres')))
    if not (a1 or n1) and f.get('nombre_completo'):
        p = texto(f.get('nombre_completo')).upper().split()
        a1, a2 = (p[0] if p else ''), (p[1] if len(p) > 3 else '')
        resto = p[2:] if len(p) > 3 else p[1:]
        n1, n2 = (resto[0] if resto else ''), ' '.join(resto[1:])
    return a1, a2, n1, n2


def _datos_simat(f):
    """Campos de DatosSimat que trae la fila (solo los que vienen con valor)."""
    d = {}
    a1, a2, n1, n2 = _nombres(f)
    for k, v in (('apellido1', a1), ('apellido2', a2), ('nombre1', n1), ('nombre2', n2)):
        if v:
            d[k] = v[:60]
    simples = {'genero': genero(f.get('genero')), 'estrato': codigo_de(f.get('estrato'), C.ESTRATO),
               'direccion': texto(f.get('direccion'))[:120], 'telefono': texto(f.get('telefono'))[:30],
               'zona': zona(f.get('zona')), 'vulnerabilidad': codigo_de(f.get('vulnerabilidad'), C.VULNERABILIDAD),
               'capacidad': codigo_de(f.get('capacidad'), C.CAPACIDAD),
               'etnia': re.sub(r'\D', '', texto(f.get('etnia')))[:3],
               'resguardo': re.sub(r'\D', '', texto(f.get('resguardo')))[:4],
               'jornada': jornada(f.get('jornada')),
               'situacion_anterior': codigo_de(f.get('situacion_anterior'), C.SITUACION_ANTERIOR[1:]),
               'mpio_residencia': municipio(f, 'mpio_residencia'), 'mpio_nacimiento': municipio(f, 'mpio_nacimiento'),
               'mpio_expedicion': municipio(f, 'mpio_expedicion'),
               'mpio_desplazamiento': municipio(f, 'mpio_desplazamiento')}
    disc = codigo_de(f.get('discapacidad'), C.DISCAPACIDAD)
    if disc and disc != '99':
        simples['discapacidad'] = disc
    d.update({k: v for k, v in simples.items() if v})
    for k in ('repitente', 'nuevo', 'subsidio', 'madre_cabeza', 'hijo_madre_cabeza', 'veterano', 'heroe'):
        v = si_no(f.get(k))
        if v is not None:
            d[k] = v
    return d


def _buscar_curso(cursos, g, grupo, sede_txt, sedes):
    candidatos = [c for c in cursos if c.grado == g]
    if grupo:
        mismos = [c for c in candidatos if clave_subgrupo(c.subgrupo) == clave_subgrupo(grupo)]
        if not mismos and len(candidatos) == 1 and not candidatos[0].subgrupo:
            mismos = candidatos
        candidatos = mismos
    if len(candidatos) > 1 and sede_txt:
        n = C.normalizar(sede_txt)
        de_sede = [c for c in candidatos if c.sede_id and (C.normalizar(sedes.get(c.sede_id, '')) in n
                                                           or n in C.normalizar(sedes.get(c.sede_id, '')))]
        candidatos = de_sede or candidatos
    return candidatos[0] if candidatos else None


def analizar(archivo, colegio):
    """Lee el reporte y dice qué haría con cada fila, sin cambiar nada."""
    filas, reconocidas, ignoradas = con_titulos(leer_tabla(archivo))
    cursos = list(Curso.objects.filter(colegio=colegio, grado__isnull=False))
    sedes = dict(Sede.objects.filter(colegio=colegio).values_list('id', 'nombre'))
    por_doc = {f.numero_documento: f.estudiante for f in FichaEstudiante.objects.filter(
        estudiante__colegio=colegio).exclude(numero_documento=None).select_related('estudiante__user')}
    otros = set(FichaEstudiante.objects.exclude(estudiante__colegio=colegio).exclude(numero_documento=None)
                .values_list('numero_documento', flat=True))
    sin_doc = {}
    for e in Estudiante.objects.filter(colegio=colegio).filter(ficha__numero_documento__isnull=True).select_related('user'):
        sin_doc.setdefault(C.normalizar(f'{e.user.last_name} {e.user.first_name}'), []).append(e)
    resultado, vistos = [], set()
    for i, f in enumerate(filas, start=2):
        doc = documento(f.get('numero_documento'))
        a1, a2, n1, n2 = _nombres(f)
        g = grado(f.get('grado'))
        grupo = normalizar_subgrupo(texto(f.get('grupo')))
        r = {'fila': i, 'doc': doc, 'tipo': tipo_documento(f.get('tipo_documento')) or '',
             'apellidos': f'{a1} {a2}'.strip(), 'nombres': f'{n1} {n2}'.strip(),
             'nacimiento': (fecha(f.get('fecha_nacimiento')) or '') and fecha(f.get('fecha_nacimiento')).isoformat(),
             'grado': g, 'grupo': grupo, 'datos': _datos_simat(f), 'estudiante': None, 'curso': None, 'error': ''}
        if not doc:
            r['error'] = 'Sin número de documento'
        elif not (a1 and n1):
            r['error'] = 'Sin primer apellido o primer nombre'
        elif doc in vistos:
            r['error'] = 'Documento repetido en el archivo'
        elif doc in otros and doc not in por_doc:
            r['error'] = 'Ese documento ya está registrado en otro colegio de la plataforma'
        if r['error']:
            resultado.append(r)
            continue
        vistos.add(doc)
        est = por_doc.get(doc)
        if est is None:
            iguales = sin_doc.get(C.normalizar(f'{r["apellidos"]} {r["nombres"]}'), [])
            if len(iguales) == 1:
                est = iguales[0]
        if est:
            r['estudiante'] = est.id
            r['actual'] = est.user.get_full_name()
            r['curso_actual'] = est.curso_id
        if g is not None:
            c = _buscar_curso(cursos, g, grupo, texto(f.get('sede')), sedes)
            r['curso'] = c.id if c else None
            r['curso_nombre'] = c.nombre if c else ''
        resultado.append(r)
    return {'filas': resultado, 'reconocidas': [formato.TITULO.get(c, c.upper()) for c in reconocidas], 'ignoradas': ignoradas}


def resumen(analisis):
    filas = analisis['filas']
    faltan = sorted({(r['grado'], r['grupo']) for r in filas if not r['error'] and r['grado'] is not None and not r['curso']},
                    key=lambda x: (x[0], x[1]))
    return {
        'nuevos': sum(1 for r in filas if not r['error'] and not r['estudiante']),
        'actualizar': sum(1 for r in filas if not r['error'] and r['estudiante']),
        'cambian_curso': sum(1 for r in filas if not r['error'] and r['estudiante'] and r['curso']
                             and r['curso'] != r.get('curso_actual')),
        'errores': [r for r in filas if r['error']],
        'cursos_faltan': faltan,
        'sin_grado': sum(1 for r in filas if not r['error'] and r['grado'] is None),
    }


def _usuario(nombres, apellidos, doc):
    from prematricula.logica import _usuario_nuevo
    try:
        return _usuario_nuevo(nombres or doc, apellidos or doc)
    except Exception:
        return f'est{doc}'


@transaction.atomic
def aplicar(analisis, colegio, crear_nuevos=True, mover_curso=True, crear_cursos=False, sede_cursos=None):
    """Aplica lo analizado. Devuelve un resumen con lo hecho y los usuarios nuevos."""
    hecho = {'creados': 0, 'actualizados': 0, 'movidos': 0, 'cursos_creados': 0, 'usuarios': []}
    cursos = {}
    if crear_cursos:
        for g, grupo in resumen(analisis)['cursos_faltan']:
            c = Curso.objects.create(colegio=colegio, grado=g, subgrupo=grupo, nombre='', sede=sede_cursos)
            cursos[(g, grupo)] = c.id
            hecho['cursos_creados'] += 1
    grupo_est, _ = Group.objects.get_or_create(name='Estudiantes')
    for r in analisis['filas']:
        if r['error']:
            continue
        curso_id = r['curso'] or cursos.get((r['grado'], r['grupo']))
        if r['estudiante']:
            est = Estudiante.objects.select_related('user').get(pk=r['estudiante'], colegio=colegio)
            u = est.user
            u.first_name, u.last_name = r['nombres'][:150], r['apellidos'][:150]
            u.save(update_fields=['first_name', 'last_name'])
            if mover_curso and curso_id and est.curso_id != curso_id:
                est.curso_id = curso_id
                est.save(update_fields=['curso'])
                hecho['movidos'] += 1
            hecho['actualizados'] += 1
        elif crear_nuevos:
            usuario = _usuario(r['nombres'], r['apellidos'], r['doc'])
            u = User.objects.create_user(username=usuario, password=usuario, first_name=r['nombres'][:150],
                                         last_name=r['apellidos'][:150])
            u.groups.add(grupo_est)
            est = Estudiante.objects.create(colegio=colegio, user=u, curso_id=curso_id)
            hecho['creados'] += 1
            hecho['usuarios'].append((u.get_full_name(), usuario))
        else:
            continue
        ficha, _ = FichaEstudiante.objects.get_or_create(estudiante=est)
        ficha.numero_documento = r['doc']
        if r['tipo']:
            ficha.tipo_documento = C.FICHA_DESDE_TIPO.get(r['tipo'], 'OT')
        if r['nacimiento']:
            ficha.fecha_nacimiento = datetime.date.fromisoformat(r['nacimiento'])
        ficha.save()
        d, _ = DatosSimat.objects.get_or_create(estudiante=est)
        for k, v in r['datos'].items():
            setattr(d, k, v)
        d.save()
    return hecho


# ---------------------------------------------------------------------------
# Excel de datos SIMAT para llenar en el colegio
# ---------------------------------------------------------------------------
# (título, campo, opciones o None)
COLUMNAS_DATOS = [
    ('TIPO_DOC', 'tipo_documento', C.TIPO_DOCUMENTO), ('DOCUMENTO', 'numero_documento', None),
    ('PRIMER_APELLIDO', 'apellido1', None), ('SEGUNDO_APELLIDO', 'apellido2', None),
    ('PRIMER_NOMBRE', 'nombre1', None), ('SEGUNDO_NOMBRE', 'nombre2', None),
    ('FECHA_NACIMIENTO', 'fecha_nacimiento', None), ('GENERO', 'genero', C.GENERO), ('ESTRATO', 'estrato', C.ESTRATO),
    ('DIRECCION', 'direccion', None), ('TELEFONO', 'telefono', None), ('MUNICIPIO_RESIDENCIA', 'mpio_residencia', None),
    ('ZONA', 'zona', C.ZONA), ('MUNICIPIO_NACIMIENTO', 'mpio_nacimiento', None),
    ('MUNICIPIO_EXPEDICION', 'mpio_expedicion', None), ('JORNADA', 'jornada', C.JORNADA),
    ('ETNIA', 'etnia', None), ('VICTIMA_CONFLICTO', 'vulnerabilidad', C.VULNERABILIDAD),
    ('REPITENTE', 'repitente', C.SI_NO), ('NUEVO', 'nuevo', C.SI_NO),
    ('SITUACION_ANO_ANTERIOR', 'situacion_anterior', C.SITUACION_ANTERIOR[1:]),
]


def plantilla_datos(colegio, ctx_filas):
    """Excel con un estudiante por fila y lo que ya tiene. Se llena y se vuelve a subir."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Datos SIMAT'
    ws.append(['ID', 'ESTUDIANTE', 'CURSO'] + [t for t, _, _ in COLUMNAS_DATOS])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1F4E79')
    for ef in ctx_filas:
        est = ef['est']
        d = getattr(est, 'simat', None)
        fila = [est.id, est.user.get_full_name(), ef['curso'].nombre if ef['curso'] else '']
        for _, campo, _ in COLUMNAS_DATOS:
            if campo in ('repitente', 'nuevo'):
                v = ef[campo]
            elif campo == 'jornada':
                v = (d.jornada if d else '') or ''
            elif campo == 'fecha_nacimiento':
                v = ef[campo].strftime('%d/%m/%Y') if ef[campo] else ''
            else:
                v = ef.get(campo, '')
            fila.append(v)
        ws.append(fila)
    ws.freeze_panes = 'D2'
    ws.column_dimensions['B'].width = 34
    ref = wb.create_sheet('Códigos')
    ref.append(['COLUMNA', 'CÓDIGO', 'SIGNIFICADO'])
    for c in ref[1]:
        c.font = Font(bold=True)
    for i, (titulo, _, opciones) in enumerate(COLUMNAS_DATOS, start=4):
        if not opciones:
            continue
        for codigo, nombre in opciones:
            ref.append([titulo, codigo, nombre])
        letra = ws.cell(row=1, column=i).column_letter
        dv = DataValidation(type='list', formula1='"' + ','.join(c for c, _ in opciones) + '"', allow_blank=True)
        dv.error, dv.errorTitle = 'Use uno de los códigos de la hoja «Códigos».', 'Código no válido'
        ws.add_data_validation(dv)
        dv.add(f'{letra}2:{letra}{max(ws.max_row, 2) + 200}')
    ref.append([])
    ref.append(['MUNICIPIOS', '73443', 'Código DANE de 5 cifras: departamento (2) + municipio (3). Ej.: 73443 = Mariquita, Tolima'])
    ref.column_dimensions['C'].width = 60
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def leer_datos(archivo, colegio):
    """Aplica el Excel de datos SIMAT llenado por el colegio. Devuelve (actualizados, errores)."""
    filas = leer_tabla(archivo)
    if not filas or 'ID' not in [texto(v).upper() for v in filas[0]]:
        raise ArchivoInvalido('Este no es el Excel de datos SIMAT de la plataforma: descárguelo de nuevo y llénelo.')
    titulos = [texto(v).upper() for v in filas[0]]
    mapa = {t: campo for t, campo, _ in COLUMNAS_DATOS}
    ids = dict(Estudiante.objects.filter(colegio=colegio).values_list('id', 'id'))
    errores, n = [], 0
    with transaction.atomic():
        for i, fila in enumerate(filas[1:], start=2):
            f = dict(zip(titulos, fila))
            try:
                est_id = int(texto(f.get('ID')))
            except ValueError:
                continue
            if est_id not in ids:
                errores.append(f'Fila {i}: el estudiante no es de este colegio.')
                continue
            est = Estudiante.objects.select_related('ficha').get(pk=est_id)
            d, _ = DatosSimat.objects.get_or_create(estudiante=est)
            ficha, _ = FichaEstudiante.objects.get_or_create(estudiante=est)
            malos = []
            for titulo, campo in mapa.items():
                if titulo not in f:
                    continue
                crudo = f[titulo]
                if campo == 'numero_documento':
                    doc = documento(crudo)
                    if doc and doc != (ficha.numero_documento or ''):
                        if FichaEstudiante.objects.filter(numero_documento=doc).exclude(estudiante=est).exists():
                            malos.append(f'el documento {doc} ya lo tiene otro estudiante')
                        else:
                            ficha.numero_documento = doc
                elif campo == 'tipo_documento':
                    t = tipo_documento(crudo)
                    if t:
                        ficha.tipo_documento = C.FICHA_DESDE_TIPO.get(t, 'OT')
                elif campo == 'fecha_nacimiento':
                    if texto(crudo):
                        v = fecha(crudo)
                        if v:
                            ficha.fecha_nacimiento = v
                        else:
                            malos.append(f'fecha «{texto(crudo)}» (use dd/mm/aaaa)')
                elif campo in ('repitente', 'nuevo'):
                    v = si_no(crudo)
                    if v is not None:
                        setattr(d, campo, v)
                elif campo.startswith('mpio_'):
                    v = re.sub(r'\D', '', texto(crudo))
                    if v and len(v) not in (4, 5):
                        malos.append(f'{titulo} «{texto(crudo)}» debe tener 5 cifras')
                    elif v:
                        setattr(d, campo, v.zfill(5))
                else:
                    opciones = dict(next(o for t, c, o in COLUMNAS_DATOS if c == campo) or [])
                    v = texto(crudo)
                    if campo == 'genero':
                        v = genero(crudo)
                    elif campo == 'jornada':
                        v = jornada(crudo)
                    elif campo == 'zona':
                        v = zona(crudo)
                    elif campo == 'etnia' and v.isdigit():
                        v = v.zfill(2)
                    elif campo in ('apellido1', 'apellido2', 'nombre1', 'nombre2'):
                        v = v.upper()[:60]
                    if opciones and texto(crudo) and (not v or v not in opciones):
                        malos.append(f'{titulo} «{texto(crudo)}» no es un código válido')
                        continue
                    if v or campo in ('apellido2', 'nombre2', 'telefono'):
                        setattr(d, campo, v)
            if malos:
                errores.append(f'Fila {i} ({texto(f.get("ESTUDIANTE"))}): ' + '; '.join(malos))
            ficha.save()
            d.save()
            n += 1
    return n, errores
