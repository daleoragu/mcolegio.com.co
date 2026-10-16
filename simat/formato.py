"""Columnas del archivo de cargue del SIMAT y cómo se arma.

El orden por defecto es el del manual del SIMAT. Como cada versión del SIMAT
puede traer otra plantilla, el colegio puede pegar la fila de títulos de la
suya en la configuración: cada título se reconoce por su nombre (con o sin
tildes, espacios o guiones bajos) y el archivo sale en ese orden. Una columna
que no se reconoce sale vacía.
"""
import csv
import io
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .codigos import normalizar

# (clave, título, otros nombres con que puede venir en una plantilla o reporte)
COLUMNAS = [
    ('ano', 'ANO', ['ANIO', 'ANO_INF', 'VIGENCIA', 'ANNO']),
    ('dpto_institucion', 'COD_DANE_MUNICIPIO', ['DPTO', 'COD_DPTO', 'CODIGO_DANE_MUNICIPIO', 'ETC']),
    ('dane_institucion', 'COD_DANE_INSTITUCION', ['DANE', 'CODIGO_DANE', 'DANE_ESTABLECIMIENTO', 'COD_DANE_ESTABLECIMIENTO',
                                                  'CODIGO_DANE_INSTITUCION']),
    ('dane_sede', 'COD_DANE_SEDE', ['DANE_SEDE', 'CODIGO_DANE_SEDE', 'DANE_ANTERIOR']),
    ('consecutivo_sede', 'CONS_SEDE', ['CONSECUTIVO', 'CONSECUTIVO_SEDE', 'CONSEC_SEDE', 'COD_SEDE']),
    ('tipo_documento', 'TIPO_DOCUMENTO', ['TIPO_DOC', 'TIPODOC', 'TIPO_IDENTIFICACION', 'TIPO_ID', 'TDOC']),
    ('numero_documento', 'NRO_DOCUMENTO', ['DOCUMENTO', 'NUMERO_DOCUMENTO', 'NUM_DOC', 'NRO_DOC', 'DOC',
                                           'IDENTIFICACION', 'NUMERO_IDENTIFICACION', 'NUM_DOCUMENTO']),
    ('dpto_expedicion', 'EXP_DEPTO', ['DPTO_EXPEDICION', 'DEPTO_EXPEDICION', 'EXP_DPTO', 'COD_DPTO_EXP']),
    ('mpio_expedicion_3', 'EXP_MUN', ['MUN_EXPEDICION', 'MPIO_EXPEDICION', 'MUNICIPIO_EXPEDICION', 'EXP_MPIO']),
    ('apellido1', 'APELLIDO1', ['PRIMER_APELLIDO', 'APELLIDO_PATERNO', 'PRI_APE', 'APELLIDO_1']),
    ('apellido2', 'APELLIDO2', ['SEGUNDO_APELLIDO', 'APELLIDO_MATERNO', 'SEG_APE', 'APELLIDO_2']),
    ('nombre1', 'NOMBRE1', ['PRIMER_NOMBRE', 'PRI_NOM', 'NOMBRE_1']),
    ('nombre2', 'NOMBRE2', ['SEGUNDO_NOMBRE', 'SEG_NOM', 'NOMBRE_2']),
    ('direccion', 'DIRECCION_RESIDENCIA', ['DIRECCION', 'DIR_RESIDENCIA', 'DIRECCION_RES']),
    ('telefono', 'TEL', ['TELEFONO', 'TELEFONO_RESIDENCIA', 'CELULAR']),
    ('dpto_residencia', 'RES_DEPTO', ['DPTO_RESIDENCIA', 'DEPTO_RESIDENCIA', 'RES_DPTO']),
    ('mpio_residencia_3', 'RES_MUN', ['MUN_RESIDENCIA', 'MPIO_RESIDENCIA', 'MUNICIPIO_RESIDENCIA', 'RES_MPIO']),
    ('estrato', 'ESTRATO', ['ESTRATO_SOCIOECONOMICO']),
    ('fecha_nacimiento', 'FECHA_NACIMIENTO', ['FEC_NACIMIENTO', 'FECHA_NAC', 'FEC_NAC', 'NACIMIENTO']),
    ('dpto_nacimiento', 'NAC_DEPTO', ['DPTO_NACIMIENTO', 'DEPTO_NACIMIENTO', 'NAC_DPTO']),
    ('mpio_nacimiento_3', 'NAC_MUN', ['MUN_NACIMIENTO', 'MPIO_NACIMIENTO', 'MUNICIPIO_NACIMIENTO', 'NAC_MPIO']),
    ('genero', 'GENERO', ['SEXO', 'GEN']),
    ('vulnerabilidad', 'POB_VICT_CONF', ['SITUACION_VULNERABILIDAD', 'VICTIMA', 'POBLACION_VICTIMA', 'DESPLAZADO']),
    ('dpto_desplazamiento', 'DPTO_EXP', ['DPTO_DESPLAZAMIENTO', 'DEPTO_DESPLAZAMIENTO']),
    ('mpio_desplazamiento_3', 'MUN_EXP', ['MUN_DESPLAZAMIENTO', 'MPIO_DESPLAZAMIENTO']),
    ('discapacidad', 'TIPO_DISCAPACIDAD', ['DISCAPACIDAD', 'DISC']),
    ('capacidad', 'CAP_EXC', ['CAPACIDADES_EXCEPCIONALES', 'CAPACIDAD_EXCEPCIONAL', 'TALENTO']),
    ('etnia', 'ETNIA', ['GRUPO_ETNICO', 'COD_ETNIA']),
    ('resguardo', 'RES', ['RESGUARDO', 'COD_RESGUARDO']),
    ('institucion_bienestar', 'INS_FAMILIAR', ['INSTITUCION_BIENESTAR', 'INST_BIENESTAR']),
    ('jornada', 'TIPO_JORNADA', ['JORNADA', 'COD_JORNADA']),
    ('caracter', 'CARACTER', []),
    ('especialidad', 'ESPECIALIDAD', []),
    ('grado', 'GRADO', ['GRADO_COD', 'COD_GRADO', 'CODIGO_GRADO']),
    ('grupo', 'GRUPO', ['COD_GRUPO']),
    ('metodologia', 'METODOLOGIA', ['MODELO', 'MODELO_EDUCATIVO']),
    ('subsidio', 'MATRICULA_CONTRATADA', ['SUBSIDIO', 'CONTRATADA']),
    ('repitente', 'REPITENTE', ['REPITE']),
    ('nuevo', 'NUEVO', ['ALUMNO_NUEVO', 'ESTUDIANTE_NUEVO']),
    ('situacion_anterior', 'SIT_ACAD_ANO_ANT', ['SITUACION_ANO_ANTERIOR', 'SITUACION_ANTERIOR', 'SIT_ANO_ANT']),
    ('motivo_retiro', 'CON_ALUM_ANO_ANT', ['MOTIVO_RETIRO', 'CONDICION_ANO_ANTERIOR']),
    ('fuente_recursos', 'FUE_RECU', ['FUENTE_RECURSOS', 'FUENTE']),
    ('zona', 'ZON_ALU', ['ZONA', 'ZONA_RESIDENCIA', 'ZONA_ALUMNO']),
    ('madre_cabeza', 'CAB_FAMILIA', ['MADRE_CABEZA_FAMILIA', 'MADRE_CABEZA']),
    ('hijo_madre_cabeza', 'BEN_MAD_FLIA', ['HIJO_MADRE_CABEZA', 'BENEFICIARIO_MADRE_CABEZA']),
    ('veterano', 'BEN_VET_FP', ['VETERANO', 'BENEFICIARIO_VETERANO']),
    ('heroe', 'BEN_HER_NAC', ['HEROE', 'BENEFICIARIO_HEROE']),
]
TITULO = {c: t for c, t, _ in COLUMNAS}
_POR_NOMBRE = {}
for _c, _t, _otros in COLUMNAS:
    for _n in [_t] + _otros:
        _POR_NOMBRE.setdefault(normalizar(_n), _c)
# Al importar también se reconocen estas (no van en el archivo de cargue).
_POR_NOMBRE.update({normalizar(n): c for n, c in [
    ('SEDE', 'sede'), ('NOMBRE_SEDE', 'sede'), ('APELLIDOS', 'apellidos'), ('NOMBRES', 'nombres'),
    ('NOMBRE_COMPLETO', 'nombre_completo'), ('ESTUDIANTE', 'nombre_completo'), ('ALUMNO', 'nombre_completo'),
    ('MUNICIPIO_RESIDENCIA_DANE', 'mpio_residencia'), ('ESTADO', 'estado'), ('CORREO', 'correo'), ('EMAIL', 'correo'),
]})


def clave_de(titulo):
    return _POR_NOMBRE.get(normalizar(titulo))


def columnas(conf):
    """[(clave o None, título)] en el orden del archivo."""
    texto = (conf.columnas or '').strip()
    if not texto:
        return [(c, t) for c, t, _ in COLUMNAS]
    sep = '\t' if '\t' in texto else (';' if ';' in texto else ',')
    titulos = [t.strip().strip('"') for t in texto.splitlines()[0].split(sep)]
    return [(clave_de(t), t) for t in titulos if t]


def valor(ef, clave):
    """Texto de una columna para un estudiante."""
    if clave is None:
        return ''
    for prefijo in ('dpto_', 'mpio_'):
        if clave.startswith(prefijo) and clave != 'dpto_institucion':
            base = 'mpio_' + clave[len(prefijo):].removesuffix('_3')
            codigo = ef.get(base, '') or ''
            return codigo[:2] if prefijo == 'dpto_' else codigo[2:]
    v = ef.get(clave, '')
    if clave == 'fecha_nacimiento':
        return v.strftime('%d/%m/%Y') if v else ''
    return '' if v is None else str(v)


def csv_bytes(conf, filas):
    cols = columnas(conf)
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=conf.separador, quoting=csv.QUOTE_MINIMAL, lineterminator='\r\n')
    w.writerow([t for _, t in cols])
    for ef in filas:
        w.writerow([valor(ef, c) for c, _ in cols])
    return buf.getvalue().encode(conf.codificacion, errors='replace')


def zip_bytes(conf, filas, nombre):
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr(f'{nombre}.csv', csv_bytes(conf, filas))
    return out.getvalue()


def excel_bytes(conf, filas):
    """El mismo archivo en Excel, con el estudiante y lo que le falta, para revisar."""
    cols = columnas(conf)
    wb = Workbook()
    ws = wb.active
    ws.title = 'SIMAT'
    ws.append(['ESTUDIANTE (plataforma)', 'CURSO', 'LE FALTA'] + [t for _, t in cols])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1F4E79')
    rojo = PatternFill('solid', fgColor='FDE2E1')
    for ef in filas:
        ws.append([ef['est'].user.get_full_name(), ef['curso'].nombre if ef['curso'] else '',
                   '; '.join(ef.get('faltan', []))] + [valor(ef, c) for c, _ in cols])
        if ef.get('faltan'):
            ws.cell(row=ws.max_row, column=3).fill = rojo
    ws.freeze_panes = 'D2'
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['C'].width = 40
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
