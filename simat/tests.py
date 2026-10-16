# -*- coding: utf-8 -*-
"""SIMAT: datos, revisión, archivo de cargue, Excel de datos e importación del reporte."""
import csv
import datetime
import io
import zipfile

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from notas.models import Curso, Estudiante, FichaEstudiante, HistorialMatricula
from notas.tests.base import ColegioDePrueba

from . import formato, importar, logica
from .models import ConfiguracionSimat, DatosSimat

HOY = datetime.date(2026, 10, 16)


def _xlsx(filas, nombre='reporte.xlsx'):
    wb = Workbook()
    for f in filas:
        wb.active.append(f)
    b = io.BytesIO()
    wb.save(b)
    return SimpleUploadedFile(nombre, b.getvalue())


class Base(ColegioDePrueba):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.a.dane = '173443000123'
        cls.a.save()
        cls.ana = cls.estudiantes[0]
        cls.ana.user.first_name, cls.ana.user.last_name = 'Ana María', 'de la Cruz Pérez'
        cls.ana.user.save()
        FichaEstudiante.objects.create(estudiante=cls.ana, tipo_documento='TI', numero_documento='1105031676',
                                       fecha_nacimiento=datetime.date(2014, 5, 3))
        conf = ConfiguracionSimat.de(cls.a)
        conf.municipio = '73443'
        conf.save()

    def completar(self, est):
        DatosSimat.objects.update_or_create(estudiante=est, defaults=dict(
            genero='F', estrato='2', direccion='Calle 5 # 3-20', mpio_residencia='73443', zona='1',
            mpio_nacimiento='73001', mpio_expedicion='73443'))


class Datos(Base):

    def test_partir_nombres_con_particulas(self):
        self.assertEqual(logica.partir('de la Cruz Pérez'), ('DE LA CRUZ', 'PÉREZ'))
        self.assertEqual(logica.partir('Ana María'), ('ANA', 'MARÍA'))
        self.assertEqual(logica.partir('Juan'), ('JUAN', ''))

    def test_lo_que_se_deduce_y_lo_que_falta(self):
        ctx, filas = logica.revision(self.a, hoy=HOY)
        ef = next(f for f in filas if f['est'] == self.ana)
        self.assertEqual((ef['apellido1'], ef['apellido2'], ef['nombre1'], ef['nombre2']),
                         ('DE LA CRUZ', 'PÉREZ', 'ANA', 'MARÍA'))
        self.assertEqual((ef['tipo_documento'], ef['grado'], ef['grupo'], ef['consecutivo_sede']),
                         ('2', '6', '01', '17344300012301'))
        self.assertEqual((ef['nuevo'], ef['situacion_anterior'], ef['discapacidad']), ('N', '1', '99'))
        self.assertIn('Género', ef['faltan'])
        self.completar(self.ana)
        ctx, filas = logica.revision(self.a, hoy=HOY)
        self.assertEqual(next(f for f in filas if f['est'] == self.ana)['faltan'], [])

    def test_validaciones_del_documento(self):
        ficha = self.ana.ficha
        ficha.numero_documento = '12345'
        ficha.save()
        self.completar(self.ana)
        ef = next(f for f in logica.revision(self.a, hoy=HOY)[1] if f['est'] == self.ana)
        self.assertIn('La tarjeta de identidad debe tener 10 u 11 números', ef['faltan'])
        ficha.tipo_documento, ficha.numero_documento = 'CC', '1105031676'
        ficha.save()
        ef = next(f for f in logica.revision(self.a, hoy=HOY)[1] if f['est'] == self.ana)
        self.assertTrue(any('Tiene cédula pero 12 años' in x for x in ef['faltan']))

    def test_historial_y_piar(self):
        HistorialMatricula.objects.create(colegio=self.a, estudiante=self.ana, ano_lectivo=HOY.year - 1,
                                          resultado='NO_PROMOVIDO', curso_nombre='601')
        from piar.models import Piar
        Piar.objects.create(colegio=self.a, estudiante=self.ana, ano_lectivo=HOY.year, categoria='TEA')
        self.ana.es_inclusion = True
        self.ana.save()
        ctx, filas = logica.revision(self.a, ano=HOY.year, hoy=HOY)
        ef = next(f for f in filas if f['est'] == self.ana)
        beto = next(f for f in filas if f['est'] == self.estudiantes[1])
        self.assertEqual((ef['repitente'], ef['situacion_anterior'], ef['discapacidad']), ('S', '2', '7'))
        self.assertEqual(beto['nuevo'], 'S')     # el colegio tiene historial y él no aparece

    def test_llenar_vacios(self):
        DatosSimat.objects.create(estudiante=self.ana, mpio_residencia='73001')
        n = logica.llenar_vacios(self.a, 'mpio_residencia', '73443')
        self.assertEqual(n, 2)
        self.assertEqual(DatosSimat.objects.get(estudiante=self.ana).mpio_residencia, '73001')


class Archivo(Base):

    def test_csv_con_las_columnas_del_colegio(self):
        self.completar(self.ana)
        conf = ConfiguracionSimat.de(self.a)
        conf.columnas = 'NRO_DOCUMENTO;Primer Apellido;APELLIDO2;NOMBRE1;FECHA_NACIMIENTO;RES_DEPTO;RES_MUN;GRADO;COLUMNA_RARA'
        conf.save()
        c = self.cliente(self.rectora)
        r = c.get('/simat/exportar/')
        self.assertEqual(r['Content-Type'], 'application/zip')
        z = zipfile.ZipFile(io.BytesIO(r.content))
        texto = z.read(z.namelist()[0]).decode('cp1252')
        filas = list(csv.reader(io.StringIO(texto), delimiter=';'))
        self.assertEqual(len(filas), 2)       # solo Ana está lista
        self.assertEqual(filas[1], ['1105031676', 'DE LA CRUZ', 'PÉREZ', 'ANA', '03/05/2014', '73', '443', '6', ''])
        self.assertIn('COLUMNA_RARA', c.get('/simat/configuracion/').content.decode())
        r = c.get('/simat/exportar/?formato=xlsx&todos=1')
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws.max_row, 4)       # títulos + 3 estudiantes

    def test_excel_de_datos_ida_y_vuelta(self):
        c = self.cliente(self.rectora)
        wb = load_workbook(io.BytesIO(c.get('/simat/datos/').content))
        ws = wb['Datos SIMAT']
        titulos = [x.value for x in ws[1]]
        fila = next(r for r in ws.iter_rows(min_row=2) if r[0].value == self.ana.id)
        valores = {'GENERO': 'F', 'ESTRATO': '2', 'DIRECCION': 'Cra 4 # 10-11', 'MUNICIPIO_RESIDENCIA': 73443,
                   'ZONA': 'Urbana', 'MUNICIPIO_NACIMIENTO': '73001', 'MUNICIPIO_EXPEDICION': '73443',
                   'SEGUNDO_NOMBRE': 'maria'}
        for t, v in valores.items():
            fila[titulos.index(t)].value = v
        beto = next(r for r in ws.iter_rows(min_row=2) if r[0].value == self.estudiantes[1].id)
        beto[titulos.index('GENERO')].value = 'X'
        b = io.BytesIO()
        wb.save(b)
        r = c.post('/simat/datos/', {'archivo': SimpleUploadedFile('datos_simat.xlsx', b.getvalue())}, follow=True)
        self.assertContains(r, 'GENERO «X» no es un código válido')
        d = DatosSimat.objects.get(estudiante=self.ana)
        self.assertEqual((d.direccion, d.zona, d.nombre2, d.mpio_residencia), ('Cra 4 # 10-11', '1', 'MARIA', '73443'))
        self.assertContains(c.get('/simat/?ver=todos'), 'Listo')


class ImportarSimat(Base):

    def _reporte(self):
        return _xlsx([
            ['REPORTE DE MATRÍCULA'],
            ['TIPO DOCUMENTO', 'NRO_DOCUMENTO', 'APELLIDO1', 'APELLIDO2', 'NOMBRE1', 'NOMBRE2', 'FECHA_NACIMIENTO',
             'GENERO', 'GRADO', 'GRUPO', 'JORNADA', 'ESTRATO', 'SEDE'],
            ['TI', 1105031676, 'DE LA CRUZ', 'PEREZ', 'ANA', 'MARIA', '03/05/2014', 'F', 6, '01', 'MAÑANA', 2, 'PRINCIPAL'],
            ['RC', '1104999888', 'GOMEZ', '', 'LUIS', '', datetime.datetime(2015, 2, 1), 'M', 'SEPTIMO', '2', 'TARDE', 1, ''],
            ['TI', '99001122', 'OTRO', '', 'COLEGIO', '', '', 'M', 6, '1', '', '', ''],
            ['TI', '', 'SIN', '', 'DOC', '', '', '', 6, '1', '', '', ''],
        ])

    def test_analizar_y_aplicar(self):
        otro = Estudiante.objects.create(colegio=self.b, user=User.objects.create_user('o1'))
        FichaEstudiante.objects.create(estudiante=otro, numero_documento='99001122')
        an = importar.analizar(self._reporte(), self.a)
        res = importar.resumen(an)
        self.assertEqual((res['nuevos'], res['actualizar'], len(res['errores'])), (1, 1, 2))
        self.assertEqual(res['cursos_faltan'], [(7, '2')])
        hecho = importar.aplicar(an, self.a, crear_cursos=True)
        self.assertEqual((hecho['creados'], hecho['actualizados'], hecho['cursos_creados']), (1, 1, 1))
        luis = Estudiante.objects.get(ficha__numero_documento='1104999888')
        self.assertEqual((luis.curso.grado, luis.curso.subgrupo, luis.ficha.tipo_documento), (7, '2', 'RC'))
        self.assertEqual((luis.simat.jornada, luis.simat.estrato, luis.simat.genero), ('3', '1', 'M'))
        self.assertEqual(luis.user.username, 'luis.gomez')
        self.ana.refresh_from_db()
        self.assertEqual((self.ana.user.first_name, self.ana.simat.apellido2), ('ANA MARIA', 'PEREZ'))

    def test_encuentra_por_nombre_al_que_no_tiene_documento(self):
        beto = self.estudiantes[1]
        an = importar.analizar(_xlsx([
            ['NRO_DOCUMENTO', 'APELLIDO1', 'NOMBRE1', 'GRADO', 'GRUPO'],
            ['1104000111', beto.user.last_name, beto.user.first_name, 6, '1']]), self.a)
        self.assertEqual(an['filas'][0]['estudiante'], beto.id)

    def test_pantallas(self):
        c = self.cliente(self.rectora)
        r = c.post('/simat/importar/', {'accion': 'analizar', 'archivo': self._reporte()}, follow=True)
        self.assertContains(r, 'Crear los cursos que no existen')
        self.assertContains(r, 'Séptimo 2')
        r = c.post('/simat/importar/', {'accion': 'aplicar', 'crear_nuevos': 'on', 'mover_curso': 'on'})
        self.assertContains(r, 'luis.gomez')
        self.assertTrue(Estudiante.objects.filter(ficha__numero_documento='1104999888', curso=None).exists())
        r = c.post('/simat/importar/', {'accion': 'analizar', 'archivo': _xlsx([['A', 'B'], [1, 2]])}, follow=True)
        self.assertContains(r, 'No encontré la fila de títulos')

    def test_solo_administracion_y_modulo(self):
        self.assertEqual(self.cliente(self.u_docente).get('/simat/').status_code, 403)
        self.assertContains(self.cliente(self.rectora).get(f'/simat/estudiante/{self.ana.id}/'), 'Le falta')
        self.assertContains(self.cliente(self.rectora).get('/panel-administrador/'), '/simat/')
        self.a.modulos_activos = ['circulares']
        self.a.save()
        self.assertEqual(self.cliente(self.rectora).get('/simat/').status_code, 403)

    def test_editar_estudiante(self):
        c = self.cliente(self.rectora)
        r = c.post(f'/simat/estudiante/{self.ana.id}/', {
            'tipo_documento': '2', 'numero_documento': '1105031676', 'fecha_nacimiento': '2014-05-03', 'genero': 'F',
            'estrato': '2', 'direccion': 'Calle 1', 'mpio_residencia': '73443', 'zona': '1', 'mpio_nacimiento': '73001',
            'mpio_expedicion': '7344', 'vulnerabilidad': '9', 'capacidad': '9', 'etnia': '00', 'resguardo': '000',
            'nuevo': 'unknown'})
        self.assertContains(r, 'código DANE de 5 cifras')
