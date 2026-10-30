"""Importar notas al aula: Excel o CSV de otras herramientas (Forms, Quizizz, Kahoot, Wordwall…),
planillas completas de otros sistemas y rejillas de rúbrica.

La idea es la misma en todo: leer la tabla, encontrar la columna de los nombres (o del documento
o del código del aula), emparejar cada fila con un estudiante y convertir lo que haya en las
columnas de notas a la escala del aula.
"""
import csv
import io
import re
import unicodedata
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from . import logica

UN_DECIMAL = Decimal('0.1')
MAX_FILAS = 400
MAX_COLUMNAS = 60

PALABRAS_NOMBRE = ('nombre', 'estudiante', 'alumno', 'apellido', 'player', 'name', 'participante', 'jugador',
                   'nickname', 'usuario', 'integrante')
PALABRAS_DOCUMENTO = ('documento', 'identificacion', 'cedula', 'tarjeta', 'doc', 'ti', 'cc', 'id')
PALABRAS_IGNORAR = ('n', 'no', 'n.o', 'n o', 'num', 'numero', '#', 'item', 'orden', 'puesto', 'rank', 'posicion',
                    'curso', 'grado', 'grupo', 'correo', 'email', 'fecha', 'hora', 'marca temporal', 'timestamp')


# ---------------------------------------------------------------------------
# Texto
# ---------------------------------------------------------------------------
def normalizar(texto):
    texto = unicodedata.normalize('NFD', str(texto or '').lower())
    texto = ''.join(ch for ch in texto if unicodedata.category(ch) != 'Mn')
    texto = re.sub(r'[^a-z0-9ñ#]+', ' ', texto)
    return ' '.join(texto.split())


def _es_encabezado_de(texto, palabras):
    t = normalizar(texto)
    return any(t == p or t.startswith(p + ' ') or f' {p} ' in f' {t} ' for p in palabras)


# ---------------------------------------------------------------------------
# Leer el archivo
# ---------------------------------------------------------------------------
def leer_archivo(archivo=None, texto=None, hoja=0):
    """(filas, nombres_de_hojas). Acepta .xlsx, .csv/.txt o una tabla pegada."""
    if archivo is None:
        return logica.leer_tabla(texto=texto)[:MAX_FILAS], []
    nombre = (getattr(archivo, 'name', '') or '').lower()
    if nombre.endswith(('.csv', '.txt', '.tsv')):
        crudo = archivo.read()
        for codificacion in ('utf-8-sig', 'cp1252', 'latin-1'):
            try:
                contenido = crudo.decode(codificacion)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialecto = csv.Sniffer().sniff(contenido[:4000], delimiters=',;\t|')
            delimitador = dialecto.delimiter
        except csv.Error:
            delimitador = ';' if contenido.count(';') > contenido.count(',') else ','
        filas = []
        for fila in csv.reader(io.StringIO(contenido), delimiter=delimitador):
            celdas = [c.strip() for c in fila]
            if any(celdas):
                filas.append(celdas)
        return _parejas(filas)[:MAX_FILAS], []
    import openpyxl
    try:
        libro = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
    except Exception:
        raise ValueError('No se pudo abrir el archivo. Súbalo como Excel (.xlsx) o CSV.')
    nombres = libro.sheetnames
    hoja = hoja if 0 <= hoja < len(nombres) else 0
    filas = []
    for fila in libro.worksheets[hoja].iter_rows(values_only=True):
        celdas = ['' if v is None else _texto_celda(v) for v in fila][:MAX_COLUMNAS]
        if any(celdas):
            filas.append(celdas)
        if len(filas) >= MAX_FILAS:
            break
    return _parejas(filas), nombres


def _texto_celda(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _parejas(filas):
    ancho = max((len(f) for f in filas), default=0)
    return [(f + [''] * ancho)[:ancho] for f in filas]


# ---------------------------------------------------------------------------
# Números
# ---------------------------------------------------------------------------
def numero(texto):
    """'4,5' -> 4.5 · '80%' -> 80 · '8/10' -> (8, 10). None si no es un número."""
    t = str(texto or '').strip().replace(' ', '')
    if not t:
        return None
    m = re.fullmatch(r'(-?\d+(?:[.,]\d+)?)/(\d+(?:[.,]\d+)?)', t)
    if m:
        try:
            a, b = Decimal(m.group(1).replace(',', '.')), Decimal(m.group(2).replace(',', '.'))
            return (a, b) if b > 0 else None
        except InvalidOperation:
            return None
    t = t.rstrip('%')
    if t.count(',') == 1 and t.count('.') == 0:
        t = t.replace(',', '.')
    elif t.count(',') and t.count('.'):
        t = t.replace('.', '').replace(',', '.') if t.rfind(',') > t.rfind('.') else t.replace(',', '')
    try:
        return Decimal(t)
    except InvalidOperation:
        return None


def a_la_escala(valor, maximo_archivo, minima, maxima):
    """Lleva un valor del archivo a la escala del aula.

    Sin máximo del archivo: el valor ya está en la escala (se ajusta a los bordes).
    Con máximo (10, 100, los puntos de Kahoot…): se reparte proporcional entre la mínima y la máxima.
    Una fracción «8/10» trae su propio máximo.
    """
    if valor is None:
        return None
    if isinstance(valor, tuple):
        valor, maximo_archivo = valor
    if maximo_archivo:
        fraccion = max(Decimal(0), min(Decimal(1), Decimal(valor) / Decimal(maximo_archivo)))
        nota = minima + (maxima - minima) * fraccion
    else:
        nota = max(minima, min(maxima, Decimal(valor)))
    return nota.quantize(UN_DECIMAL, rounding=ROUND_HALF_UP)


# ---------------------------------------------------------------------------
# Entender la tabla
# ---------------------------------------------------------------------------
def analizar(filas, maxima_escala):
    """Dónde está el encabezado, qué columnas son nombres/documento/código y cuáles traen notas."""
    if not filas:
        return None
    # El encabezado: la primera fila (de las 15 primeras) con una celda que parezca «nombre».
    encabezado = 0
    for i, f in enumerate(filas[:15]):
        if any(_es_encabezado_de(c, PALABRAS_NOMBRE) for c in f):
            encabezado = i
            break
    titulos = filas[encabezado]
    datos = [f for f in filas[encabezado + 1:] if any(f)]
    col_nombres, col_documento, col_codigo = [], None, None
    for j, t in enumerate(titulos):
        nt = normalizar(t)
        if nt in ('codigo aula', 'codigo del aula'):
            col_codigo = j
        elif _es_encabezado_de(t, PALABRAS_NOMBRE):
            col_nombres.append(j)
        elif col_documento is None and _es_encabezado_de(t, PALABRAS_DOCUMENTO):
            col_documento = j
    if not col_nombres:
        # Sin encabezado reconocible: la columna con más texto.
        puntajes = [(sum(1 for f in datos if f[j] and numero(f[j]) is None), j) for j in range(len(titulos))]
        mejor = max(puntajes, default=(0, None))
        if mejor[0]:
            col_nombres = [mejor[1]]
    especiales = set(col_nombres) | {col_documento, col_codigo}
    columnas = []
    for j, t in enumerate(titulos):
        if j in especiales or _es_encabezado_de(t, PALABRAS_IGNORAR):
            continue
        valores = [numero(f[j]) for f in datos if f[j]]
        numericos = [v for v in valores if v is not None]
        if not numericos or len(numericos) < max(1, len(valores) // 2):
            continue
        simples = [v for v in numericos if not isinstance(v, tuple)]
        mayor = max(simples, default=Decimal(0))
        columnas.append({'j': j, 'titulo': (t or f'Columna {j + 1}')[:60],
                         'maximo_sugerido': _maximo_sugerido(mayor, maxima_escala, t),
                         'mayor': mayor, 'cuantos': len(numericos)})
    registros = []
    for f in datos:
        nombre = ' '.join(f[j] for j in col_nombres if f[j]).strip()
        documento = f[col_documento].strip() if col_documento is not None else ''
        codigo = f[col_codigo].strip() if col_codigo is not None else ''
        if not (nombre or documento or codigo):
            continue
        registros.append({'nombre': nombre, 'documento': documento, 'codigo': codigo,
                          'valores': [f[c['j']] for c in columnas]})
    return {'encabezado': encabezado, 'titulos': titulos, 'columnas': columnas, 'registros': registros,
            'col_nombres': col_nombres, 'col_documento': col_documento, 'col_codigo': col_codigo}


def _maximo_sugerido(mayor, maxima_escala, titulo):
    """Si las notas del archivo pasan la escala del aula, de cuánto es la escala del archivo."""
    if '%' in (titulo or '') or 'porcentaje' in normalizar(titulo):
        return Decimal(100)
    if mayor <= maxima_escala:
        return None
    for tope in (Decimal(10), Decimal(20), Decimal(50), Decimal(100)):
        if mayor <= tope:
            return tope
    return mayor        # puntos (Kahoot, Quizizz): el mejor puntaje es el máximo


# ---------------------------------------------------------------------------
# Emparejar con los estudiantes
# ---------------------------------------------------------------------------
def _fichas(personas):
    docs = {}
    for p in personas:
        ficha = getattr(p, 'ficha', None) if not hasattr(p, 'completo') else None
        if ficha is not None and ficha.numero_documento:
            docs[re.sub(r'\D', '', ficha.numero_documento)] = p
    return docs


def emparejar(registros, personas):
    """Para cada registro, el estudiante que le corresponde (o None). Usa, en ese orden:
    el código del aula (e12/i7), el documento y el nombre (sin tildes, en cualquier orden)."""
    por_clave = {logica.clave_de(p): p for p in personas}
    try:
        por_documento = _fichas(personas)
    except Exception:
        por_documento = {}
    tokens = [(p, set(normalizar(logica.nombre_de(p)).split())) for p in personas]
    usados = set()
    salida = []
    for r in registros:
        p = por_clave.get(r.get('codigo', '').strip().lower())
        if p is None and r.get('documento'):
            p = por_documento.get(re.sub(r'\D', '', r['documento']))
        if p is None and r.get('nombre'):
            p = _por_nombre(r['nombre'], tokens, usados)
        if p is not None:
            usados.add(p.pk)
        salida.append(p)
    return salida


def _por_nombre(nombre, tokens, usados):
    mios = set(normalizar(nombre).split())
    if not mios:
        return None
    mejores, mejor = [], 0.0
    for p, suyos in tokens:
        comunes = len(mios & suyos)
        if not comunes:
            continue
        # Todos los nombres del archivo deben estar en el estudiante (o casi), y al menos dos si hay dos.
        necesarios = min(2, len(mios), len(suyos))
        if comunes < necesarios:
            continue
        puntaje = comunes / max(len(mios), 1) + comunes / max(len(suyos), 1) / 10
        if puntaje > mejor + 1e-9:
            mejores, mejor = [p], puntaje
        elif abs(puntaje - mejor) < 1e-9:
            mejores.append(p)
    if len(mejores) == 1:
        return mejores[0]
    return None


# ---------------------------------------------------------------------------
# Rejilla de rúbrica (estudiantes × criterios)
# ---------------------------------------------------------------------------
SI = {'si', 's', 'x', '1', 'true', 'verdadero', 'cumple', 'ok', 'yes', 'v', '✓', '✔'}
NO = {'no', 'n', '0', 'false', 'falso', 'no cumple', 'f', '✗', '-'}


def nivel_de(texto, niveles):
    """El nivel de la rúbrica que dice la celda: por nombre, por su valor (4,0) o por su posición (1…n)."""
    t = normalizar(texto)
    if not t:
        return None
    for n in niveles:
        if normalizar(n.nombre) == t:
            return n
    if len(t) >= 3:
        parecidos = [n for n in niveles if normalizar(n.nombre).startswith(t) or t.startswith(normalizar(n.nombre))]
        if len(parecidos) == 1:
            return parecidos[0]
    v = numero(texto)
    if isinstance(v, Decimal):
        for n in niveles:
            if n.valor == v:
                return n
        if v == v.to_integral_value() and 1 <= int(v) <= len(niveles):
            return sorted(niveles, key=lambda n: n.valor)[int(v) - 1]
    return None


def cumple(texto):
    t = normalizar(texto) or str(texto or '').strip()
    if t in SI:
        return True
    if t in NO:
        return False
    return None


def leer_rejilla(filas, instrumento, personas):
    """Lee una rejilla y devuelve {'calificados': [(persona, seleccion)], 'sin_estudiante': [nombres],
    'errores': [textos], 'criterios_sin_columna': [nombres]}."""
    criterios = list(instrumento.criterios.all())
    niveles = list(instrumento.niveles.all())
    encabezado = 0
    for i, f in enumerate(filas[:15]):
        if any(_es_encabezado_de(c, PALABRAS_NOMBRE) for c in f) or any(normalizar(c) == 'codigo aula' for c in f):
            encabezado = i
            break
    titulos = [normalizar(t) for t in filas[encabezado]] if filas else []
    columna_de = {}
    for cr in criterios:
        nombre = normalizar(cr.nombre)
        for j, t in enumerate(titulos):
            if t and (t == nombre or t.startswith(nombre[:25]) or nombre.startswith(t[:25])) and j not in columna_de.values():
                columna_de[cr.id] = j
                break
    analisis = analizar(filas, Decimal(0))
    registros = []
    for f in filas[encabezado + 1:]:
        nombre = ' '.join(f[j] for j in (analisis['col_nombres'] if analisis else []) if f[j]).strip()
        codigo = f[analisis['col_codigo']].strip() if analisis and analisis['col_codigo'] is not None else ''
        documento = f[analisis['col_documento']].strip() if analisis and analisis['col_documento'] is not None else ''
        if nombre or codigo or documento:
            registros.append({'nombre': nombre, 'codigo': codigo, 'documento': documento, 'fila': f})
    quienes = emparejar(registros, personas)
    salida = {'calificados': [], 'sin_estudiante': [], 'errores': [],
              'criterios_sin_columna': [cr.nombre for cr in criterios if cr.id not in columna_de]}
    for r, p in zip(registros, quienes):
        if p is None:
            salida['sin_estudiante'].append(r['nombre'] or r['codigo'] or r['documento'])
            continue
        sel = {}
        for cr in criterios:
            j = columna_de.get(cr.id)
            if j is None or not r['fila'][j]:
                continue
            celda = r['fila'][j]
            if instrumento.tipo == 'lista':
                v = cumple(celda)
            else:
                nivel = nivel_de(celda, niveles)
                v = nivel.id if nivel else None
            if v is None:
                salida['errores'].append(f'{logica.nombre_de(p)} · {cr.nombre}: «{celda}» no corresponde a ningún nivel.')
                continue
            sel[str(cr.id)] = v
        if sel:
            salida['calificados'].append((p, sel))
    return salida


def plantilla_rejilla(tarea, personas):
    """Excel para llenar la rejilla sin internet: una fila por estudiante y una columna por criterio."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    inst = tarea.instrumento
    criterios = list(inst.criterios.all())
    niveles = list(inst.niveles.all())
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Rejilla'
    ws.append([f'{tarea.titulo} · {inst.nombre}'])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append(['Escriba en cada casilla el nivel (o Sí / No en una lista de chequeo). No cambie la columna «Código aula».'])
    ws['A2'].font = Font(italic=True, color='5B6474')
    ws.append(['Código aula', 'Estudiante'] + [cr.nombre for cr in criterios])
    for celda in ws[3]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', start_color='1F3A68')
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for p in personas:
        ws.append([logica.clave_de(p), logica.nombre_de(p)] + [''] * len(criterios))
    if criterios and personas:
        opciones = '"Sí,No"' if inst.tipo == 'lista' else '"' + ','.join(n.nombre for n in niveles)[:250] + '"'
        dv = DataValidation(type='list', formula1=opciones, allow_blank=True)
        ws.add_data_validation(dv)
        ultima = openpyxl.utils.get_column_letter(2 + len(criterios))
        dv.add(f'C4:{ultima}{3 + len(personas)}')
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 34
    for i in range(len(criterios)):
        ws.column_dimensions[openpyxl.utils.get_column_letter(3 + i)].width = 22
    ws.row_dimensions[3].height = 42
    ws.freeze_panes = 'C4'
    if inst.tipo != 'lista':
        ws2 = wb.create_sheet('Niveles')
        ws2.append(['Nivel', 'Valor'])
        for n in niveles:
            ws2.append([n.nombre, float(n.valor)])
        ws2.append([])
        ws2.append(['Criterio'] + [n.nombre for n in niveles])
        for cr in criterios:
            ws2.append([cr.nombre] + [cr.descriptor(n) for n in niveles])
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
