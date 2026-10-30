"""Reglas del aula virtual: quién entra, cómo se calcula la nota y cómo pasa a la planilla."""
import io
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from typing import Optional

from django.db import transaction
from django.utils import timezone

from .models import Aula, Criterio, Entrega, Instrumento, Nivel

UN_DECIMAL = Decimal('0.1')
MAX_ARCHIVO_MB = 10
MAX_FOTOS = 12
NIVELES_POR_DEFECTO = [('BAJO', Decimal('2.0')), ('BÁSICO', Decimal('3.5')), ('ALTO', Decimal('4.3')),
                       ('SUPERIOR', Decimal('5.0'))]


# ---------------------------------------------------------------------------
# Quién está mirando
# ---------------------------------------------------------------------------
@dataclass
class Contexto:
    aula: Aula
    rol: str                       # 'docente' | 'estudiante'
    usuario: object = None
    estudiante: object = None      # notas.Estudiante
    integrante: object = None      # salon_digital.Integrante

    @property
    def es_docente(self):
        return self.rol == 'docente'

    @property
    def nombre(self):
        if self.integrante is not None:
            return self.integrante.completo
        if self.usuario is not None:
            return self.usuario.get_full_name() or self.usuario.username
        return ''

    @property
    def clave(self):
        """'e12' o 'i7': cómo se guarda este estudiante en «para» de tareas, material y actividades."""
        persona = self.estudiante or self.integrante
        return clave_de(persona) if persona is not None else ''

    @property
    def base(self):
        return 'aula/base_salon.html' if self.aula.es_salon else 'notas/base.html'


def contexto(request, aula):
    """El Contexto de quien mira, o None si no tiene que ver con esta aula."""
    user = request.user
    if aula.asignacion_id:
        if getattr(request, 'colegio', None) is None or aula.colegio_id != request.colegio.id:
            return None
        if not user.is_authenticated:
            return None
        from notas.permisos import es_admin
        asig = aula.asignacion
        if (asig.docente_id and asig.docente.user_id == user.id) or es_admin(request):
            return Contexto(aula, 'docente', usuario=user)
        est = getattr(user, 'estudiante', None)
        if est is not None and est.colegio_id == aula.colegio_id and est.curso_id == asig.curso_id and est.is_active:
            return Contexto(aula, 'estudiante', usuario=user, estudiante=est)
        return None
    if aula.grupo_id:
        if user.is_authenticated and aula.grupo.propietario.usuario_id == user.id:
            return Contexto(aula, 'docente', usuario=user)
        from salon_digital import identidad
        from salon_digital.models import Integrante
        datos = identidad.leer(request)
        if datos and datos.get('grupo') == aula.grupo_id:
            integrante = Integrante.objects.filter(pk=datos.get('integrante'), grupo=aula.grupo,
                                                   activo=True).first()
            if integrante:
                return Contexto(aula, 'estudiante', integrante=integrante)
    return None


def aula_de_asignacion(asignacion):
    aula, _ = Aula.objects.get_or_create(asignacion=asignacion, defaults={'colegio': asignacion.colegio})
    return aula


def aula_de_grupo(grupo):
    aula, _ = Aula.objects.get_or_create(grupo=grupo)
    return aula


def estudiantes(aula):
    """Los de la clase: estudiantes del curso o integrantes del grupo, ordenados por apellido."""
    if aula.asignacion_id:
        from notas.models import Estudiante
        return list(Estudiante.objects.filter(curso=aula.asignacion.curso, is_active=True)
                    .select_related('user').order_by('user__last_name', 'user__first_name'))
    if aula.grupo_id:
        return list(aula.grupo.integrantes.filter(activo=True))
    return []


def nombre_de(persona):
    if hasattr(persona, 'completo'):
        return persona.completo
    u = persona.user
    return f'{u.last_name} {u.first_name}'.strip() or u.username


def clave_de(persona):
    return f'i{persona.pk}' if hasattr(persona, 'completo') else f'e{persona.pk}'


# ---------------------------------------------------------------------------
# Para quién: tareas, material y actividades para toda la clase o para algunos
# ---------------------------------------------------------------------------
def es_para(obj, clave):
    """¿Le toca a este estudiante? (vacío = toda la clase; el docente lo ve todo)."""
    return not obj.para or clave in obj.para


def para_c(obj, c):
    return c.es_docente or es_para(obj, c.clave)


def personas_de(obj, aula, todas=None):
    """Los estudiantes a quienes va dirigido."""
    todas = estudiantes(aula) if todas is None else todas
    return [p for p in todas if es_para(obj, clave_de(p))]


def opciones_para(aula):
    """Para el formulario: [{'clave', 'nombre', 'inclusion'}]."""
    return [{'clave': clave_de(p), 'nombre': nombre_de(p), 'inclusion': bool(getattr(p, 'es_inclusion', False))}
            for p in estudiantes(aula)]


def leer_para(request, aula):
    """Lee el «¿para quién?» del formulario. Devuelve (lista, error)."""
    if request.POST.get('para_modo') != 'algunos':
        return [], ''
    validas = {clave_de(p) for p in estudiantes(aula)}
    escogidas = sorted(set(request.POST.getlist('para')) & validas)
    if not escogidas:
        return [], 'Escoja al menos un estudiante (o deje «Toda la clase»).'
    return escogidas, ''


def nombres_para(obj, aula):
    return [nombre_de(p) for p in personas_de(obj, aula)] if obj.para else []


# ---------------------------------------------------------------------------
# Novedades: el globo rojo con lo que el estudiante no ha visto
# ---------------------------------------------------------------------------
SECCIONES = ('muro', 'material', 'tareas', 'actividades', 'cuestionarios')


def _visto(aula, persona, crear=False):
    from .models import Visto
    filtro = filtro_persona(persona)
    v = Visto.objects.filter(aula=aula, **filtro).first()
    if v is None and crear:
        v = Visto.objects.create(aula=aula, **filtro)
    return v


def marcar_visto(aula, persona, seccion):
    if seccion not in SECCIONES or persona is None:
        return
    v = _visto(aula, persona, crear=True)
    v.secciones[seccion] = timezone.now().isoformat()
    v.save(update_fields=['secciones'])


def novedades(aula, persona, secciones_vistas=None):
    """{sección: cuántas cosas nuevas hay desde la última vez que la abrió}."""
    from django.utils.dateparse import parse_datetime
    if secciones_vistas is None:
        v = _visto(aula, persona)
        secciones_vistas = v.secciones if v else {}
    clave = clave_de(persona)

    def nuevas(qs, campo, seccion, filtrar_para=True):
        desde = parse_datetime(secciones_vistas.get(seccion) or '')
        if desde is not None:
            qs = qs.filter(**{f'{campo}__gt': desde})
        return sum(1 for o in qs if not filtrar_para or es_para(o, clave))

    publicaciones = aula.publicaciones.filter(tarea__isnull=True)
    return {
        'muro': nuevas(publicaciones, 'creada', 'muro', filtrar_para=False),
        'material': nuevas(aula.materiales.all(), 'creado', 'material'),
        'tareas': nuevas(aula.tareas.filter(publicada=True), 'creada', 'tareas'),
        'actividades': nuevas(aula.actividades.filter(publicada=True), 'creada', 'actividades'),
        'cuestionarios': nuevas(aula.cuestionarios.filter(a_su_ritmo=True, preguntas__isnull=False).distinct(),
                                'creado', 'cuestionarios', filtrar_para=False),
    }


def filtro_persona(persona):
    """{'estudiante': x} o {'integrante': x}, para buscar su entrega."""
    return {'integrante': persona} if hasattr(persona, 'completo') else {'estudiante': persona}


# ---------------------------------------------------------------------------
# Escala de notas
# ---------------------------------------------------------------------------
def escala_docente(usuario):
    """[(nombre, min, max)] de la escala del docente de Salón Digital (la nacional si no ha puesto la suya)."""
    from .models import EscalaDocente
    escala = EscalaDocente.objects.filter(usuario=usuario).first() if usuario is not None else None
    return (escala or EscalaDocente()).lista()


def validar_escala(niveles):
    """Errores en palabras de una escala [(nombre, min, max)]; vacía si está bien."""
    errores = []
    if len(niveles) < 2:
        errores.append('La escala necesita al menos dos desempeños.')
    nombres = [n for n, _, _ in niveles]
    if len(set(nombres)) != len(nombres):
        errores.append('Hay desempeños con el mismo nombre.')
    for n, a, b in niveles:
        if a is None or b is None:
            errores.append(f'«{n}» necesita nota mínima y máxima.')
        elif a >= b:
            errores.append(f'En «{n}» la nota mínima debe ser menor que la máxima.')
    buenos = sorted([x for x in niveles if x[1] is not None and x[2] is not None], key=lambda x: x[1])
    for (n1, _, b1), (n2, a2, _) in zip(buenos, buenos[1:]):
        if a2 <= b1:
            errores.append(f'Los rangos de «{n1}» y «{n2}» se cruzan.')
    return errores


def rango(aula):
    """(mínima, máxima) de las notas: la escala del colegio, o la del docente en Salón Digital."""
    if aula.colegio_id:
        from notas.boletin.ponderacion import rango_notas
        return rango_notas(aula.colegio)
    escala = escala_docente(aula.docente_usuario)
    return escala[0][1], escala[-1][2]


def niveles_sugeridos(colegio, usuario=None):
    """Los niveles de la escala (del colegio o del docente) con un valor representativo; si no hay, los de siempre."""
    if colegio is None and usuario is not None:
        escala = escala_docente(usuario)
        return [(n, b if i == len(escala) - 1 else ((a + b) / 2).quantize(UN_DECIMAL, rounding=ROUND_HALF_UP))
                for i, (n, a, b) in enumerate(escala)]
    if colegio is not None:
        from notas.models import EscalaValoracion
        escala = list(EscalaValoracion.objects.filter(colegio=colegio).order_by('valor_minimo'))
        if len(escala) >= 2:
            salida = []
            for i, e in enumerate(escala):
                if i == len(escala) - 1:
                    valor = e.valor_maximo
                else:
                    valor = ((e.valor_minimo + e.valor_maximo) / 2).quantize(UN_DECIMAL, rounding=ROUND_HALF_UP)
                salida.append((e.nombre_desempeno.upper(), valor))
            return salida
    return list(NIVELES_POR_DEFECTO)


def calcular_nota(instrumento, seleccion, minima, maxima):
    """(nota, completa). Rúbrica: promedio ponderado de los niveles. Lista: % cumplido llevado a la escala."""
    criterios = list(instrumento.criterios.all())
    if not criterios:
        return None, False
    seleccion = seleccion or {}
    if instrumento.tipo == 'lista':
        total = sum(c.peso for c in criterios)
        cumplido = sum(c.peso for c in criterios if seleccion.get(str(c.id)) is True)
        completa = all(str(c.id) in seleccion for c in criterios)
        if total <= 0:
            return None, False
        nota = minima + (maxima - minima) * cumplido / total
        return nota.quantize(UN_DECIMAL, rounding=ROUND_HALF_UP), completa
    niveles = {n.id: n for n in instrumento.niveles.all()}
    suma, pesos = Decimal('0'), Decimal('0')
    for c in criterios:
        nivel = niveles.get(_entero(seleccion.get(str(c.id))))
        if nivel is not None:
            suma += nivel.valor * c.peso
            pesos += c.peso
    if pesos <= 0:
        return None, False
    nota = (suma / pesos).quantize(UN_DECIMAL, rounding=ROUND_HALF_UP)
    nota = max(minima, min(maxima, nota))
    return nota, pesos == sum(c.peso for c in criterios)


def _entero(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def desempeno(aula, nota):
    if nota is None:
        return ''
    if not aula.colegio_id:
        nota = Decimal(nota)
        return next((n for n, a, b in escala_docente(aula.docente_usuario) if a <= nota <= b), '')
    from notas import escalas
    asig = aula.asignacion
    return escalas.de(aula.colegio).desempeno(nota, asig.materia_id, asig.curso.grado)


# ---------------------------------------------------------------------------
# Instrumentos: crear, plantillas e importar
# ---------------------------------------------------------------------------
PLANTILLAS = {
    'guia': ('Guía de trabajo', 'rubrica', [
        'Desarrolla todos los puntos de la guía',
        'Procedimientos claros y correctos',
        'Presentación, orden y letra legible',
        'Entrega a tiempo']),
    'exposicion': ('Exposición oral', 'rubrica', [
        'Dominio del tema',
        'Claridad y orden de las ideas',
        'Material de apoyo',
        'Manejo del tiempo',
        'Responde las preguntas']),
    'laboratorio': ('Informe de laboratorio', 'rubrica', [
        'Planteamiento y objetivos',
        'Procedimiento y toma de datos',
        'Análisis de resultados',
        'Conclusiones',
        'Presentación del informe']),
    'proyecto': ('Proyecto', 'rubrica', [
        'Planeación y cumplimiento de etapas',
        'Calidad del producto final',
        'Trabajo en equipo',
        'Socialización']),
    'escrito': ('Texto escrito o ensayo', 'rubrica', [
        'Idea principal y argumentos',
        'Coherencia y cohesión',
        'Ortografía y redacción',
        'Uso de fuentes']),
    'cuaderno': ('Revisión de cuaderno', 'lista', [
        'Tiene todas las clases',
        'Fecha y título en cada clase',
        'Ejercicios resueltos',
        'Correcciones hechas',
        'Orden y presentación']),
    'grupo': ('Trabajo en grupo (coevaluación)', 'lista', [
        'Aportó ideas',
        'Cumplió con su parte',
        'Escuchó a los compañeros',
        'Llegó a tiempo a los encuentros']),
}

# Descriptores escalonados a partir del criterio, del nivel más bajo al más alto.
GRADACION = ['No lo logra o no lo presenta', 'Lo logra de forma parcial',
             'Lo logra en su mayoría', 'Lo logra por completo y con calidad']


def gradacion(n):
    """n frases de menor a mayor para rellenar los descriptores."""
    if n <= 1:
        return GRADACION[-1:]
    if n == len(GRADACION):
        return list(GRADACION)
    idx = [round(i * (len(GRADACION) - 1) / (n - 1)) for i in range(n)]
    return [GRADACION[i] for i in idx]


@transaction.atomic
def crear_instrumento(propietario, colegio, nombre, tipo, criterios, niveles=None, descriptores=None):
    """criterios: [(nombre, peso)]; niveles: [(nombre, valor)] de menor a mayor;
    descriptores: [[texto por nivel] por criterio] (opcional)."""
    inst = Instrumento.objects.create(propietario=propietario, colegio=colegio, nombre=nombre[:120], tipo=tipo)
    objs_niveles = []
    if tipo == 'rubrica':
        for i, (n, v) in enumerate(niveles or niveles_sugeridos(colegio, propietario)):
            objs_niveles.append(Nivel.objects.create(instrumento=inst, nombre=str(n)[:40], valor=Decimal(v), orden=i))
    for i, (n, p) in enumerate(criterios):
        textos = (descriptores[i] if descriptores and i < len(descriptores) else None) or \
            (gradacion(len(objs_niveles)) if objs_niveles else [])
        Criterio.objects.create(
            instrumento=inst, nombre=str(n)[:200], peso=Decimal(p or 1), orden=i,
            descriptores={str(nv.id): (textos[k] if k < len(textos) else '') for k, nv in enumerate(objs_niveles)})
    return inst


def desde_plantilla(propietario, colegio, clave):
    nombre, tipo, criterios = PLANTILLAS[clave]
    return crear_instrumento(propietario, colegio, nombre, tipo, [(c, 1) for c in criterios])


def leer_tabla(texto=None, archivo=None):
    """Filas (lista de listas de texto) de una tabla pegada (RubiStar, Word, Excel) o de un .xlsx."""
    filas = []
    if archivo is not None:
        import openpyxl
        libro = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
        hoja = libro.worksheets[0]
        for fila in hoja.iter_rows(values_only=True):
            celdas = ['' if v is None else str(v).strip() for v in fila]
            if any(celdas):
                filas.append(celdas)
    else:
        for linea in (texto or '').splitlines():
            if not linea.strip():
                continue
            if '\t' in linea:
                celdas = linea.split('\t')
            elif '|' in linea:
                celdas = linea.strip().strip('|').split('|')
            elif ';' in linea:
                celdas = linea.split(';')
            else:
                celdas = [linea]
            filas.append([c.strip() for c in celdas])
    # quitar columnas vacías al final
    ancho = max((len(f) for f in filas), default=0)
    while ancho and all((len(f) < ancho or not f[ancho - 1]) for f in filas):
        ancho -= 1
    return [(f + [''] * ancho)[:ancho] for f in filas]


def _numero(texto):
    try:
        return Decimal(str(texto).replace(',', '.').replace('%', '').strip())
    except Exception:
        return None


def interpretar_tabla(filas, colegio=None, usuario=None):
    """De una tabla a (criterios, niveles, descriptores, tipo) listos para crear_instrumento.

    Primera fila = encabezados: «Criterio» (o «Category» de RubiStar), una columna
    opcional «Peso» o «%», y un encabezado por nivel («4», «3»… o «Superior», «Alto»…).
    Sin columnas de nivel, es una lista de chequeo.
    """
    if not filas:
        raise ValueError('La tabla está vacía.')
    encabezado = [c.strip() for c in filas[0]]
    cuerpo = [f for f in filas[1:] if f and f[0].strip()]
    if not cuerpo:
        raise ValueError('La tabla necesita una fila de encabezados y al menos un criterio.')
    col_peso = next((i for i, h in enumerate(encabezado) if i > 0 and (('peso' in h.lower()) or h.strip() == '%')),
                    None)
    cols_nivel = [i for i in range(1, len(encabezado)) if i != col_peso and encabezado[i]]
    criterios = []
    for f in cuerpo:
        peso = _numero(f[col_peso]) if col_peso is not None else None
        criterios.append((f[0].strip()[:200], peso if peso and peso > 0 else Decimal('1')))
    if not cols_nivel:
        return criterios, None, None, 'lista'

    nombres = [encabezado[i] for i in cols_nivel]
    numeros = [_numero(n) for n in nombres]
    # RubiStar va de mayor a menor (4, 3, 2, 1); la plataforma guarda de menor a mayor.
    if all(n is not None for n in numeros) and numeros[0] > numeros[-1]:
        cols_nivel, nombres, numeros = cols_nivel[::-1], nombres[::-1], numeros[::-1]
    elif all(n is None for n in numeros) and _parece_descendente(nombres):
        cols_nivel, nombres = cols_nivel[::-1], nombres[::-1]
        numeros = [None] * len(nombres)

    sugeridos = niveles_sugeridos(colegio, usuario)
    niveles = []
    for k, nombre in enumerate(nombres):
        # Valor: el del nivel de la escala en el mismo puesto (el más bajo con el más bajo).
        pos = round(k * (len(sugeridos) - 1) / max(1, len(nombres) - 1)) if len(nombres) > 1 else len(sugeridos) - 1
        nombre_nivel, valor = sugeridos[pos]
        etiqueta = nombre if numeros[k] is None else nombre_nivel
        niveles.append((etiqueta.upper()[:40], valor))
    descriptores = [[f[i].strip() for i in cols_nivel] for f in cuerpo]
    return criterios, niveles, descriptores, 'rubrica'


def _parece_descendente(nombres):
    orden = ['superior', 'excelente', 'alto', 'bueno', 'básico', 'basico', 'aceptable', 'bajo', 'insuficiente',
             'deficiente']
    pos = [next((i for i, o in enumerate(orden) if o in n.lower()), None) for n in nombres]
    pos = [p for p in pos if p is not None]
    return len(pos) >= 2 and pos[0] < pos[-1]


# ---------------------------------------------------------------------------
# Archivos
# ---------------------------------------------------------------------------
def imagenes_a_pdf(archivos):
    """Las fotos del cuaderno, achicadas, en un solo PDF. Devuelve bytes."""
    from PIL import Image, ImageOps
    paginas = []
    for f in archivos:
        img = Image.open(f)
        img = ImageOps.exif_transpose(img).convert('RGB')
        img.thumbnail((1600, 1600))
        paginas.append(img)
    if not paginas:
        return b''
    salida = io.BytesIO()
    paginas[0].save(salida, format='PDF', save_all=True, append_images=paginas[1:], resolution=150, quality=75)
    return salida.getvalue()


def es_imagen(nombre):
    return (nombre or '').lower().endswith(('.jpg', '.jpeg', '.png', '.webp', '.heic', '.gif'))


def muy_grande(archivo):
    return archivo is not None and archivo.size > MAX_ARCHIVO_MB * 1024 * 1024


# ---------------------------------------------------------------------------
# Horario de la clase en vivo
# ---------------------------------------------------------------------------
def horario_de(aula):
    """[{'dia', 'texto', 'hoy', 'ahora'}] de las clases de la asignación en el horario publicado."""
    if not aula.asignacion_id:
        return []
    try:
        from horarios.models import DIAS, Sesion
    except Exception:
        return []
    sesiones = (Sesion.objects.filter(horario__colegio=aula.colegio, horario__estado='PUBLICADO',
                                      asignacion=aula.asignacion)
                .select_related('franja').order_by('dia', 'franja__inicio'))
    nombres = dict(DIAS)
    ahora = timezone.localtime()
    salida = []
    for s in sesiones:
        hoy = s.dia == ahora.isoweekday()
        en_curso = False
        if hoy:
            inicio = ahora.replace(hour=s.franja.inicio.hour, minute=s.franja.inicio.minute, second=0)
            fin = ahora.replace(hour=s.franja.fin.hour, minute=s.franja.fin.minute, second=0)
            en_curso = (inicio - timezone.timedelta(minutes=10)) <= ahora <= fin
        salida.append({'dia': nombres.get(s.dia, ''), 'texto': s.franja.texto, 'hoy': hoy, 'ahora': en_curso})
    return salida


# ---------------------------------------------------------------------------
# Avisos
# ---------------------------------------------------------------------------
def avisar_estudiantes(aula, mensaje, url, para=None):
    """Notificación (y push al celular) a los estudiantes del curso, o solo a los de «para».
    En Salón Digital no hay cuentas."""
    if not aula.asignacion_id:
        return 0
    from notas.models import Notificacion
    avisos = [Notificacion(colegio=aula.colegio, destinatario_id=e.user_id, mensaje=mensaje[:255], url=url,
                           tipo='GENERAL') for e in estudiantes(aula) if not para or clave_de(e) in para]
    Notificacion.objects.bulk_create(avisos)
    return len(avisos)


# ---------------------------------------------------------------------------
# Paso a la planilla del colegio
# ---------------------------------------------------------------------------
def componentes_de(aula):
    """[(código, nombre)] de los componentes que cuentan en la asignación."""
    if not aula.asignacion_id:
        return []
    from notas.planillas.columnas import componentes_activos
    return [(c, n) for c, n, _ in componentes_activos(aula.asignacion)]


@transaction.atomic
def pasar_notas_a_planilla(aula, periodo, componente, titulo, pares):
    """Pone las notas [(estudiante, nota)] en la columna «titulo» del componente. (pasadas, error)."""
    if not aula.asignacion_id:
        return 0, 'Esta aula no tiene planilla del colegio.'
    if periodo is None or not componente:
        return 0, 'Escoja el periodo y el componente antes de pasar las notas a la planilla.'
    if not periodo.esta_activo:
        return 0, f'El {periodo} está cerrado: la planilla no recibe notas.'
    from notas.planillas.columnas import (columnas_del_plan, componentes_activos, guardar_plan,
                                          limpiar_columnas, notas_guardadas)
    from notas.planillas.guardar import guardar_estudiante
    asig = aula.asignacion
    activos = [c for c, _, _ in componentes_activos(asig)]
    if componente not in activos:
        return 0, 'Ese componente no cuenta en esta asignatura.'
    pares = [(e, n) for e, n in pares if e is not None and n is not None]
    if not pares:
        return 0, 'Todavía no hay notas para pasar.'
    columna = limpiar_columnas([titulo])[0]
    columnas = columnas_del_plan(asig, periodo, componente)
    if columna not in columnas:
        guardar_plan(asig, periodo, componente, columnas + [columna])
    minima, maxima = rango(aula)
    for estudiante, valor in pares:
        por_componente = {}
        for codigo in activos:
            notas = notas_guardadas(asig, periodo, codigo, [estudiante]).get(estudiante.id, [])
            if codigo == componente:
                notas = [(d, v) for d, v in notas if d != columna]
                notas.append((columna, max(minima, min(maxima, valor))))
            por_componente[codigo] = notas
        guardar_estudiante(aula.colegio, asig, periodo, estudiante, por_componente)
    return len(pares), ''


def pasar_a_planilla(tarea):
    """Copia las notas calificadas de la tarea a su columna en la planilla. Devuelve (pasadas, error)."""
    pares = [(e.estudiante, e.nota) for e in tarea.entregas.filter(estudiante__isnull=False, nota__isnull=False)
             .select_related('estudiante')]
    n, error = pasar_notas_a_planilla(tarea.aula, tarea.periodo, tarea.componente, tarea.titulo, pares)
    if n:
        tarea.pasada_planilla = timezone.now()
        tarea.save(update_fields=['pasada_planilla'])
    return n, error


# ---------------------------------------------------------------------------
# Planilla propia (Salón Digital) y resumen de calificaciones
# ---------------------------------------------------------------------------
def examenes_de(aula):
    """Los exámenes PuntoExacto de esta clase (misma asignatura y curso, o el mismo grupo)."""
    from puntoexacto.models import Examen
    if aula.asignacion_id:
        return Examen.objects.filter(asignacion=aula.asignacion, archivado=False).order_by('fecha', 'id')
    if aula.grupo_id:
        return Examen.objects.filter(grupo=aula.grupo, archivado=False).order_by('fecha', 'id')
    return Examen.objects.none()


def columnas_calificaciones(aula):
    """Todo lo que da nota en el aula: tareas, actividades, cuestionarios, exámenes PuntoExacto
    y (en el plan individual) las columnas escritas a mano.

    [{'clave', 'titulo', 'tipo', 'url', 'notas': {persona_id: nota}, 'peso', 'manual'}]
    """
    from django.urls import reverse
    from . import juegos
    clave = 'integrante_id' if aula.es_salon else 'estudiante_id'
    pesos = (aula.planilla or {}).get('pesos', {})
    columnas = []

    def agregar(k, titulo, tipo, url, notas, manual=None):
        try:
            peso = Decimal(str(pesos.get(k, 1)))
        except Exception:
            peso = Decimal(1)
        columnas.append({'clave': k, 'titulo': titulo, 'tipo': tipo, 'url': url, 'notas': notas,
                         'peso': peso, 'manual': manual})

    for t in aula.tareas.filter(publicada=True).order_by('creada'):
        agregar(f'tarea:{t.pk}', t.titulo, 'Tarea', reverse('aula:tarea', args=[aula.pk, t.pk]),
                {getattr(e, clave): e.nota for e in t.entregas.filter(nota__isnull=False)})
    for a in aula.actividades.filter(publicada=True, califica=True).order_by('creada'):
        agregar(f'act:{a.pk}', a.titulo, 'Actividad', reverse('aula:actividad', args=[aula.pk, a.pk]),
                {getattr(r, clave): r.nota for r in a.resultados.all()})
    for q in aula.cuestionarios.order_by('creado'):
        mejores = juegos.mejores_notas(q)
        if not mejores and not q.a_su_ritmo:
            continue
        agregar(f'quiz:{q.pk}', q.titulo, 'Cuestionario', reverse('aula:cuestionario', args=[aula.pk, q.pk]),
                {(k[1] if aula.es_salon else k[0]): it.nota for k, it in mejores.items()})
    for ex in examenes_de(aula):
        notas = {getattr(h, clave): h.nota for h in ex.hojas.filter(estado='calificada', nota__isnull=False)
                 if getattr(h, clave)}
        agregar(f'examen:{ex.pk}', ex.titulo, 'Examen', reverse('aula:examen', args=[aula.pk, ex.pk]), notas)
    if aula.es_salon:
        for col in aula.columnas_manuales.prefetch_related('notas'):
            agregar(f'manual:{col.pk}', col.titulo, 'A mano', '', {getattr(n, clave): n.nota for n in col.notas.all()},
                    manual=col)
    return columnas


def tabla_calificaciones(aula):
    """{'columnas': [...], 'filas': [{'nombre', 'notas': [nota|None…], 'promedio', 'desempeno'}]}

    El promedio es ponderado con el peso de cada columna (1 si no se ha cambiado; 0 = no cuenta).
    Lo que falta no cuenta, salvo que el docente pida contarlo con la nota mínima.
    """
    columnas = columnas_calificaciones(aula)
    faltantes_minima = bool((aula.planilla or {}).get('faltantes_minima'))
    minima, _ = rango(aula)
    total_pesos = sum(c['peso'] for c in columnas)
    for c in columnas:
        c['porcentaje'] = (c['peso'] * 100 / total_pesos).quantize(Decimal('1')) if total_pesos else 0
    filas = []
    for p in estudiantes(aula):
        fila = [c['notas'].get(p.id) for c in columnas]
        suma, pesos = Decimal(0), Decimal(0)
        for c, n in zip(columnas, fila):
            if c['peso'] <= 0:
                continue
            if n is None and faltantes_minima:
                n = minima
            if n is not None:
                suma += Decimal(n) * c['peso']
                pesos += c['peso']
        prom = (suma / pesos).quantize(UN_DECIMAL, rounding=ROUND_HALF_UP) if pesos else None
        filas.append({'persona': p, 'clave': clave_de(p), 'nombre': nombre_de(p), 'notas': fila,
                      'celdas': list(zip(columnas, fila)), 'promedio': prom,
                      'desempeno': desempeno(aula, prom) if prom is not None else ''})
    return {'columnas': columnas, 'filas': filas, 'faltantes_minima': faltantes_minima}


def _numero_nota(texto):
    texto = (str(texto or '')).strip().replace(',', '.')
    if not texto:
        return None
    try:
        return Decimal(texto)
    except Exception:
        raise ValueError(texto)


@transaction.atomic
def guardar_planilla(aula, datos):
    """Guarda pesos, «contar lo que falta» y las notas de las columnas a mano. Devuelve [errores]."""
    from .models import NotaManual
    errores = []
    columnas = columnas_calificaciones(aula)
    pesos = {}
    for c in columnas:
        crudo = datos.get(f'peso-{c["clave"]}')
        if crudo is None:
            continue
        try:
            peso = _numero_nota(crudo)
        except ValueError:
            peso = None
        if peso is None or peso < 0 or peso > 100:
            errores.append(f'El peso de «{c["titulo"]}» debe ser un número entre 0 y 100.')
            continue
        pesos[c['clave']] = float(peso)
    aula.planilla = dict(aula.planilla or {}, pesos=dict((aula.planilla or {}).get('pesos', {}), **pesos),
                         faltantes_minima=datos.get('faltantes_minima') == 'on')
    aula.save(update_fields=['planilla'])
    minima, maxima = rango(aula)
    personas = estudiantes(aula)
    for c in columnas:
        col = c['manual']
        if col is None:
            continue
        actuales = {(n.integrante_id or n.estudiante_id): n for n in col.notas.all()}
        for p in personas:
            nombre = f'n-{col.pk}-{clave_de(p)}'
            if nombre not in datos:
                continue
            try:
                valor = _numero_nota(datos.get(nombre))
            except ValueError:
                errores.append(f'«{datos.get(nombre)}» no es una nota ({nombre_de(p)}, {col.titulo}).')
                continue
            existente = actuales.get(p.pk)
            if valor is None:
                if existente:
                    existente.delete()
                continue
            if valor < minima or valor > maxima:
                errores.append(f'{nombre_de(p)}, {col.titulo}: la nota debe estar entre {minima} y {maxima}.')
                continue
            if existente:
                if existente.nota != valor:
                    existente.nota = valor
                    existente.save(update_fields=['nota'])
            else:
                NotaManual.objects.create(columna=col, nota=valor, **filtro_persona(p))
    return errores


def excel_calificaciones(aula):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    datos = tabla_calificaciones(aula)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Calificaciones'
    ws.append([f'{aula.titulo} · {aula.subtitulo}'])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(['N.º', 'Estudiante'] + [c['titulo'] for c in datos['columnas']] + ['Promedio', 'Desempeño'])
    for c in ws[3]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', start_color='1F3A68')
        c.alignment = Alignment(horizontal='center', wrap_text=True)
    ws.append(['', 'Peso'] + [f"{c['porcentaje']} %" for c in datos['columnas']] + ['', ''])
    for c in ws[4]:
        c.font = Font(italic=True, color='5B6474')
        c.alignment = Alignment(horizontal='center')
    for i, f in enumerate(datos['filas'], 1):
        ws.append([i, f['nombre']] + [float(n) if n is not None else None for n in f['notas']]
                  + [float(f['promedio']) if f['promedio'] is not None else None, f['desempeno']])
    ws.column_dimensions['B'].width = 36
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def entrega_de(tarea, persona) -> Optional[Entrega]:
    return tarea.entregas.filter(**filtro_persona(persona)).first()
