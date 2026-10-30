"""Aula virtual: muro, tareas con rúbrica, entrega con fotos, paso a la planilla y Salón Digital."""
import io
import shutil
import tempfile
from decimal import Decimal

import openpyxl
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, override_settings

from aula import logica
from aula.models import Aula, Comentario, Entrega, Instrumento, Publicacion, Tarea
from notas.models import Curso, Notificacion, NotaDetallada, Calificacion
from notas.tests.base import CLAVE, ColegioDePrueba

MEDIA = tempfile.mkdtemp(prefix='aula-pruebas-')


def foto(nombre='hoja.png'):
    from PIL import Image
    salida = io.BytesIO()
    Image.new('RGB', (400, 300), (240, 240, 240)).save(salida, format='PNG')
    return SimpleUploadedFile(nombre, salida.getvalue(), content_type='image/png')


@override_settings(MEDIA_ROOT=MEDIA)
class AulaDelColegio(ColegioDePrueba):

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        # El aula es del plan Premium: se activa para el colegio de prueba.
        cls.a.modulos_activos = list(cls.a.modulos_activos) + ['aula']
        cls.a.save(update_fields=['modulos_activos'])

    def setUp(self):
        self.ana = self.estudiantes[0]
        self.docente_c = self.cliente(self.u_docente)

    def abrir(self):
        r = self.docente_c.get(f'/aula/asignacion/{self.asig.id}/')
        self.assertEqual(r.status_code, 302)
        return Aula.objects.get(asignacion=self.asig)

    def test_inicio_y_permisos(self):
        r = self.docente_c.get('/aula/')
        self.assertContains(r, 'Matemáticas')
        aula = self.abrir()
        self.assertEqual(self.cliente(self.ana.user).get(f'/aula/{aula.pk}/').status_code, 200)
        # Un estudiante de otro curso no entra
        c702 = Curso.objects.create(colegio=self.a, nombre='702', grado=7, subgrupo='2')
        otro = self.crear_estudiante('otro', curso=c702)
        self.assertEqual(self.cliente(otro.user).get(f'/aula/{aula.pk}/').status_code, 403)
        # Ni alguien de otro colegio
        self.assertNotEqual(self.cliente(self.ana.user, host='b.localhost').get(f'/aula/{aula.pk}/').status_code, 200)
        # El estudiante no publica
        self.cliente(self.ana.user).post(f'/aula/{aula.pk}/', {'texto': 'Hola'})
        self.assertFalse(Publicacion.objects.exists())

    def test_muro_avisa_y_comentarios(self):
        aula = self.abrir()
        r = self.docente_c.post(f'/aula/{aula.pk}/', {'texto': 'Bienvenidos a la clase'})
        self.assertEqual(r.status_code, 302)
        pub = Publicacion.objects.get()
        self.assertEqual(Notificacion.objects.filter(mensaje__contains='nueva publicación').count(), 3)
        est = self.cliente(self.ana.user)
        est.post(f'/aula/{aula.pk}/publicacion/{pub.pk}/comentar/', {'texto': '¡Gracias profe!'})
        self.assertEqual(Comentario.objects.get().nombre, 'Ana Prueba')
        # El docente apaga los comentarios
        self.docente_c.post(f'/aula/{aula.pk}/clase/', {'enlace_clase': 'meet.google.com/abc-defg-hij'})
        aula.refresh_from_db()
        self.assertFalse(aula.comentarios_activos)
        self.assertEqual(aula.enlace_clase, 'https://meet.google.com/abc-defg-hij')
        self.assertEqual(aula.plataforma, 'meet')
        est.post(f'/aula/{aula.pk}/publicacion/{pub.pk}/comentar/', {'texto': 'otro'})
        self.assertEqual(Comentario.objects.count(), 1)
        self.assertContains(est.get(f'/aula/{aula.pk}/'), 'Unirse a la clase en vivo')

    def test_material_con_archivo_y_video(self):
        aula = self.abrir()
        self.docente_c.post(f'/aula/{aula.pk}/material/', {
            'titulo': 'Guía 1', 'unidad': 'Funciones', 'archivo': SimpleUploadedFile('guia.pdf', b'%PDF-1.4 x')})
        self.docente_c.post(f'/aula/{aula.pk}/material/', {
            'titulo': 'Video', 'enlace': 'https://youtu.be/dQw4w9WgXcQ'})
        r = self.cliente(self.ana.user).get(f'/aula/{aula.pk}/material/')
        self.assertContains(r, 'Guía 1')
        self.assertContains(r, 'youtube-nocookie.com/embed/dQw4w9WgXcQ')
        m = aula.materiales.get(titulo='Guía 1')
        r = self.cliente(self.ana.user).get(f'/aula/{aula.pk}/archivo/mat/{m.pk}/')
        self.assertEqual(r.status_code, 302)

    def test_tarea_con_rubrica_entrega_calificacion_y_planilla(self):
        aula = self.abrir()
        inst = logica.desde_plantilla(self.u_docente, self.a, 'guia')
        self.assertEqual([n.nombre for n in inst.niveles.all()], ['BAJO', 'BÁSICO', 'ALTO', 'SUPERIOR'])
        r = self.docente_c.post(f'/aula/{aula.pk}/tareas/nueva/', {
            'titulo': 'Guía 3', 'instrucciones': 'Resolver', 'instrumento': inst.pk, 'periodo': self.p1.pk,
            'componente': 'SABER', 'publicada': 'on', 'acepta_tardias': 'on'})
        self.assertEqual(r.status_code, 302)
        tarea = Tarea.objects.get()
        self.assertTrue(Publicacion.objects.filter(tarea=tarea).exists())

        # Ana entrega dos fotos: quedan en un solo PDF
        est = self.cliente(self.ana.user)
        est.post(f'/aula/{aula.pk}/tareas/{tarea.pk}/entregar/', {'fotos': [foto('a.png'), foto('b.png')]})
        e = Entrega.objects.get()
        self.assertTrue(e.nombre_archivo.endswith('.pdf'))
        with e.archivo.open('rb') as f:
            self.assertTrue(f.read(5).startswith(b'%PDF'))

        # El docente califica: Superior en todo menos un criterio en Básico
        criterios = list(inst.criterios.all())
        niveles = list(inst.niveles.all())
        datos = {f'c{cr.id}': niveles[3].id for cr in criterios}
        datos[f'c{criterios[0].id}'] = niveles[1].id
        datos['comentario'] = 'Muy bien'
        r = self.docente_c.post(f'/aula/{aula.pk}/tareas/{tarea.pk}/calificar/e{self.ana.id}/', datos)
        e.refresh_from_db()
        # (3,5 + 5 + 5 + 5) / 4 = 4,6
        self.assertEqual(e.nota, Decimal('4.6'))
        self.assertEqual(e.estado, 'calificada')
        self.assertContains(est.get(f'/aula/{aula.pk}/tareas/{tarea.pk}/'), 'Muy bien')

        # A la planilla
        self.docente_c.post(f'/aula/{aula.pk}/tareas/{tarea.pk}/planilla/')
        nd = NotaDetallada.objects.get(calificacion_promedio__estudiante=self.ana, descripcion='Guía 3')
        self.assertEqual(nd.valor_nota, Decimal('4.6'))
        self.assertEqual(nd.calificacion_promedio.tipo_nota, 'SABER')
        definitiva = Calificacion.objects.get(estudiante=self.ana, periodo=self.p1, tipo_nota='PROM_PERIODO')
        self.assertEqual(definitiva.valor_nota, Decimal('2.76'))      # 4,6 × 60 % (HACER sin notas)
        tarea.refresh_from_db()
        self.assertIsNotNone(tarea.pasada_planilla)

        # Pasarla otra vez no duplica la nota
        self.docente_c.post(f'/aula/{aula.pk}/tareas/{tarea.pk}/planilla/')
        self.assertEqual(NotaDetallada.objects.filter(calificacion_promedio__estudiante=self.ana,
                                                      descripcion='Guía 3').count(), 1)

    def test_grilla_y_nota_directa(self):
        aula = self.abrir()
        lista = logica.desde_plantilla(self.u_docente, self.a, 'cuaderno')
        tarea = Tarea.objects.create(aula=aula, titulo='Cuaderno', instrumento=lista)
        crits = list(lista.criterios.all())
        datos = {}
        for k, cr in enumerate(crits):        # beto cumple 3 de 5; ana todo
            datos[f'e{self.ana.id}-{cr.id}'] = 'si'
            datos[f'e{self.estudiantes[1].id}-{cr.id}'] = 'si' if k < 3 else 'no'
        self.docente_c.post(f'/aula/{aula.pk}/tareas/{tarea.pk}/grilla/', datos)
        notas = {e.estudiante_id: e.nota for e in Entrega.objects.all()}
        self.assertEqual(notas[self.ana.id], Decimal('5.0'))
        self.assertEqual(notas[self.estudiantes[1].id], Decimal('3.4'))     # 1 + 4 × 3/5
        # Nota directa
        t2 = Tarea.objects.create(aula=aula, titulo='Quiz')
        self.docente_c.post(f'/aula/{aula.pk}/tareas/{t2.pk}/calificar/e{self.ana.id}/', {'nota': '4,2'})
        self.assertEqual(Entrega.objects.get(tarea=t2).nota, Decimal('4.2'))
        r = self.docente_c.get(f'/aula/{aula.pk}/calificaciones/?formato=excel')
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws['C3'].value, 'Cuaderno')

    def test_importar_rubistar_y_editar(self):
        pegado = ('CATEGORY\tPeso\t4\t3\t2\t1\n'
                  'Contenido\t2\tCompleto\tCasi\tPoco\tNada\n'
                  'Ortografía\t1\tSin errores\tPocos\tVarios\tMuchos\n')
        r = self.docente_c.post('/aula/instrumentos/', {'accion': 'importar', 'nombre': 'Ensayo', 'tabla': pegado})
        self.assertEqual(r.status_code, 302)
        inst = Instrumento.objects.get(nombre='Ensayo')
        niveles = list(inst.niveles.all())
        self.assertEqual(len(niveles), 4)
        self.assertEqual(niveles[0].valor, Decimal('2.0'))       # el «1» de RubiStar = el nivel más bajo
        contenido = inst.criterios.get(nombre='Contenido')
        self.assertEqual(contenido.peso, Decimal('2'))
        self.assertEqual(contenido.descriptor(niveles[0]), 'Nada')
        self.assertEqual(contenido.descriptor(niveles[3]), 'Completo')
        # Excel
        libro = openpyxl.Workbook()
        hoja = libro.active
        hoja.append(['Criterio', 'Superior', 'Alto', 'Básico', 'Bajo'])
        hoja.append(['Dominio', 'a', 'b', 'c', 'd'])
        salida = io.BytesIO()
        libro.save(salida)
        archivo = SimpleUploadedFile('r.xlsx', salida.getvalue())
        self.docente_c.post('/aula/instrumentos/', {'accion': 'importar', 'nombre': 'Expo', 'archivo': archivo})
        expo = Instrumento.objects.get(nombre='Expo')
        self.assertEqual([n.nombre for n in expo.niveles.all()], ['BAJO', 'BÁSICO', 'ALTO', 'SUPERIOR'])
        # Editar: cambiar un valor y agregar un criterio
        url = f'/aula/instrumentos/{expo.pk}/'
        self.docente_c.post(url, {
            'nombre': 'Exposición', 'nivel_nombre': ['BAJO', 'BÁSICO', 'ALTO', 'SUPERIOR'],
            'nivel_valor': ['1,5', '3.5', '4.3', '5'], 'crit_nombre': ['Dominio', 'Voz'], 'crit_peso': ['1', '1'],
            'desc_1_3': 'Se escucha muy bien'})
        expo.refresh_from_db()
        self.assertEqual(expo.nombre, 'Exposición')
        self.assertEqual(expo.niveles.first().valor, Decimal('1.5'))
        self.assertEqual(expo.criterios.count(), 2)
        self.assertEqual(expo.criterios.get(nombre='Voz').descriptor(expo.niveles.last()), 'Se escucha muy bien')

    def test_lista_de_chequeo_calculo(self):
        inst = logica.crear_instrumento(self.u_docente, self.a, 'L', 'lista', [('a', 1), ('b', 1), ('c', 2)])
        crs = list(inst.criterios.all())
        sel = {str(crs[0].id): True, str(crs[1].id): False, str(crs[2].id): True}
        nota, completa = logica.calcular_nota(inst, sel, Decimal('1'), Decimal('5'))
        self.assertTrue(completa)
        self.assertEqual(nota, Decimal('4.0'))      # 3 de 4 → 1 + 4 × 0,75


class AulaSalonDigital(ColegioDePrueba):

    def test_grupo_codigo_entrega_y_planilla(self):
        from salon_digital.models import Cuenta, Grupo, Integrante
        u = User.objects.create_user('libre', password=CLAVE, first_name='Lina', last_name='Libre')
        cuenta = Cuenta.objects.create(usuario=u, slug='lina', nombre_publico='Lina Libre', estado='verificada')
        grupo = Grupo.objects.create(propietario=cuenta, nombre='7-01')
        luis = Integrante.objects.create(grupo=grupo, apellidos='Pérez', nombres='Luis')
        Integrante.objects.create(grupo=grupo, apellidos='Gómez', nombres='Sara')

        profe = Client(HTTP_HOST='localhost')
        profe.force_login(u)
        self.assertContains(profe.get('/aula/'), '7-01')
        r = profe.get(f'/aula/grupo/{grupo.pk}/')
        aula = Aula.objects.get(grupo=grupo)
        self.assertEqual(r.url, f'/aula/{aula.pk}/')
        tarea = Tarea.objects.create(aula=aula, titulo='Lectura')

        est = Client(HTTP_HOST='localhost')
        r = est.get(f'/aula/{aula.pk}/')
        self.assertEqual(r.url, f'/aula/{aula.pk}/entrar/')
        r = est.get(f'/aula/{aula.pk}/entrar/', {'codigo': 'MAL123'})
        self.assertContains(r, 'no corresponde')
        r = est.get(f'/aula/{aula.pk}/entrar/', {'codigo': grupo.codigo.lower()})
        self.assertContains(r, 'Pérez Luis')
        est.post(f'/aula/{aula.pk}/entrar/', {'codigo': grupo.codigo, 'integrante': luis.pk})
        self.assertEqual(est.get(f'/aula/{aula.pk}/').status_code, 200)
        est.post(f'/aula/{aula.pk}/tareas/{tarea.pk}/entregar/', {'texto': 'Mi resumen'})
        e = Entrega.objects.get()
        self.assertEqual(e.integrante, luis)

        profe.post(f'/aula/{aula.pk}/tareas/{tarea.pk}/calificar/i{luis.pk}/', {'nota': '4.5'})
        e.refresh_from_db()
        self.assertEqual(e.nota, Decimal('4.5'))
        r = profe.get(f'/aula/{aula.pk}/calificaciones/')
        self.assertContains(r, 'Pérez Luis')
        # El estudiante no ve las calificaciones de todos
        self.assertEqual(est.get(f'/aula/{aula.pk}/calificaciones/').status_code, 403)
        # Pasar a planilla no aplica en Salón Digital
        self.assertEqual(logica.pasar_a_planilla(tarea)[0], 0)
