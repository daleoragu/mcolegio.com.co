"""Planilla con pesos del plan individual y exámenes PuntoExacto dentro del aula."""
import datetime
import io
from decimal import Decimal

import openpyxl
from django.contrib.auth.models import User
from django.test import Client

from aula import logica
from aula.models import Aula, ColumnaManual, Entrega, NotaManual, Tarea
from notas.tests.base import CLAVE, ColegioDePrueba
from puntoexacto.models import Examen, Hoja, Pregunta, Respuesta


def salon():
    from salon_digital.models import Cuenta, Grupo, Integrante
    u = User.objects.create_user('libre', password=CLAVE, first_name='Lina', last_name='Libre')
    cuenta = Cuenta.objects.create(usuario=u, slug='lina', nombre_publico='Lina Libre', estado='verificada')
    grupo = Grupo.objects.create(propietario=cuenta, nombre='7-01')
    luis = Integrante.objects.create(grupo=grupo, apellidos='Pérez', nombres='Luis')
    sara = Integrante.objects.create(grupo=grupo, apellidos='Gómez', nombres='Sara')
    profe = Client(HTTP_HOST='localhost')
    profe.force_login(u)
    profe.get(f'/aula/grupo/{grupo.pk}/')
    return cuenta, grupo, luis, sara, profe, Aula.objects.get(grupo=grupo)


class PlanillaIndividual(ColegioDePrueba):

    def test_columnas_a_mano_pesos_y_promedio(self):
        cuenta, grupo, luis, sara, profe, aula = salon()
        url = f'/aula/{aula.pk}/calificaciones/'
        t = Tarea.objects.create(aula=aula, titulo='Taller')
        Entrega.objects.create(tarea=t, integrante=luis, nota=Decimal('4.0'))
        Entrega.objects.create(tarea=t, integrante=sara, nota=Decimal('2.0'))
        r = profe.get(url)
        self.assertContains(r, 'Planilla del grupo')

        profe.post(url, {'accion': 'agregar', 'titulo': 'Quiz en papel'})
        col = ColumnaManual.objects.get()
        datos = {'accion': 'guardar', f'peso-tarea:{t.pk}': '1', f'peso-manual:{col.pk}': '3',
                 f'n-{col.pk}-i{luis.pk}': '5', f'n-{col.pk}-i{sara.pk}': ''}
        r = profe.post(url, datos)
        self.assertEqual(r.status_code, 302)
        self.assertEqual(NotaManual.objects.get().nota, Decimal('5'))
        filas = {f['nombre']: f for f in logica.tabla_calificaciones(Aula.objects.get(pk=aula.pk))['filas']}
        self.assertEqual(filas['Pérez Luis']['promedio'], Decimal('4.8'))     # (4 × 1 + 5 × 3) / 4
        self.assertEqual(filas['Gómez Sara']['promedio'], Decimal('2.0'))     # lo que falta no cuenta

        # Contar lo que falta con la mínima (1,0)
        datos['faltantes_minima'] = 'on'
        profe.post(url, datos)
        filas = {f['nombre']: f for f in logica.tabla_calificaciones(Aula.objects.get(pk=aula.pk))['filas']}
        self.assertEqual(filas['Gómez Sara']['promedio'], Decimal('1.3'))     # (2 + 1 × 3) / 4

        # Notas fuera de la escala o que no son números: no se guardan
        r = profe.post(url, dict(datos, **{f'n-{col.pk}-i{sara.pk}': '9'}), follow=True)
        self.assertContains(r, 'debe estar entre')
        profe.post(url, dict(datos, **{f'n-{col.pk}-i{sara.pk}': 'abc'}))
        self.assertEqual(NotaManual.objects.count(), 1)
        # Peso 0 = no cuenta
        profe.post(url, dict(datos, **{f'peso-manual:{col.pk}': '0'}))
        filas = {f['nombre']: f for f in logica.tabla_calificaciones(Aula.objects.get(pk=aula.pk))['filas']}
        self.assertEqual(filas['Pérez Luis']['promedio'], Decimal('4.0'))

        r = profe.get(url + '?formato=excel')
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws['D3'].value, 'Quiz en papel')
        self.assertEqual(ws['D4'].value, '0 %')

        # Borrar la columna
        profe.post(url, {'accion': 'eliminar', 'columna': col.pk})
        self.assertFalse(ColumnaManual.objects.exists())

    def test_en_el_colegio_no_se_edita(self):
        self.a.modulos_activos = list(self.a.modulos_activos) + ['salon']
        self.a.save(update_fields=['modulos_activos'])
        profe = self.cliente(self.u_docente)
        profe.get(f'/aula/asignacion/{self.asig.id}/')
        aula = Aula.objects.get(asignacion=self.asig)
        self.assertEqual(profe.post(f'/aula/{aula.pk}/calificaciones/', {'accion': 'agregar', 'titulo': 'x'})
                         .status_code, 403)


def examen_con_hoja(hoja_kw, **examen_kw):
    ex = Examen.objects.create(titulo='Parcial fracciones', fecha=datetime.date.today(), numero_preguntas=3,
                               **examen_kw)
    preguntas = [Pregunta.objects.create(examen=ex, numero=i, correcta='A') for i in (1, 2, 3)]
    hoja = Hoja.objects.create(examen=ex, identificador=f'PE-{ex.id:05d}-0001', estado='calificada',
                               nota=Decimal('3.67'), **hoja_kw)
    for p, marcada in zip(preguntas, 'AAB'):
        Respuesta.objects.create(hoja=hoja, pregunta=p, marcada=marcada)
    return ex, hoja


class ExamenesEnElAula(ColegioDePrueba):

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.a.modulos_activos = list(cls.a.modulos_activos) + ['salon']
        cls.a.save(update_fields=['modulos_activos'])

    def test_examen_del_colegio(self):
        ana, beto = self.estudiantes[0], self.estudiantes[1]
        profe = self.cliente(self.u_docente)
        profe.get(f'/aula/asignacion/{self.asig.id}/')
        aula = Aula.objects.get(asignacion=self.asig)
        ex, hoja = examen_con_hoja({'estudiante': ana}, colegio=self.a, docente=self.asig.docente,
                                   asignacion=self.asig)
        base = f'/aula/{aula.pk}'
        r = profe.get(f'{base}/examenes/')
        self.assertContains(r, 'Parcial fracciones')
        self.assertContains(r, f'/puntoexacto/nuevo/?asignacion={self.asig.id}')
        r = profe.get(f'{base}/examenes/{ex.pk}/')
        self.assertEqual([h.nota for h in r.context['hojas']], [Decimal('3.67')])
        r = profe.get(f'{base}/calificaciones/')
        self.assertContains(r, 'Parcial fracciones')

        est = self.cliente(ana.user)
        self.assertContains(est.get(f'{base}/examenes/'), 'Parcial fracciones')
        r = est.get(f'{base}/examenes/{ex.pk}/')
        self.assertEqual(r.context['aciertos'], 2)
        self.assertEqual([f['estado'] for f in r.context['filas']], ['bien', 'bien', 'mal'])
        # Beto no tiene hoja: no lo ve
        otro = self.cliente(beto.user)
        self.assertNotContains(otro.get(f'{base}/examenes/'), 'Parcial fracciones')
        self.assertEqual(otro.get(f'{base}/examenes/{ex.pk}/').status_code, 404)

        # Crear desde el aula: llega con la asignatura escogida
        r = profe.get(f'/puntoexacto/nuevo/?asignacion={self.asig.id}')
        self.assertContains(r, f'value="{self.asig.id}" selected')

        # Sin el módulo PuntoExacto: no hay pestaña
        self.a.modulos_activos = [m for m in self.a.modulos_activos if m != 'puntoexacto']
        self.a.save(update_fields=['modulos_activos'])
        self.assertNotContains(profe.get(f'{base}/'), 'Exámenes</a>')
        self.assertEqual(profe.get(f'{base}/examenes/').status_code, 404)

    def test_puntoexacto_en_el_plan_individual(self):
        cuenta, grupo, luis, sara, profe, aula = salon()
        self.assertEqual(profe.get('/puntoexacto/').status_code, 200)
        r = profe.get(f'/puntoexacto/nuevo/?grupo={grupo.pk}')
        self.assertContains(r, '7-01')
        self.assertNotContains(r, 'name="asignacion"')
        r = profe.post('/puntoexacto/nuevo/', {
            'titulo': 'Quiz 1', 'grupo': grupo.pk, 'fecha': '2026-10-30', 'numero_preguntas': 5,
            'numero_opciones': 4, 'hojas_por_pagina': 1, 'metodo': 'con_piso', 'nota_maxima': '5.00',
            'nota_minima': '1.00', 'nota_todo_mal': '1.00', 'nota_nada_marcado': '2.00'})
        self.assertEqual(r.status_code, 302, r.context and r.context['form'].errors)
        ex = Examen.objects.get()
        self.assertEqual((ex.cuenta, ex.grupo, ex.colegio), (cuenta, grupo, None))
        self.assertEqual(ex.preguntas.count(), 5)

        profe.post(f'/puntoexacto/{ex.pk}/hojas/', {'grupo_ids': [grupo.pk], 'sueltas': 1})
        self.assertEqual(ex.hojas.count(), 3)
        h = ex.hojas.get(integrante=luis)
        self.assertEqual(h.nombre, 'Pérez Luis')
        # Preparar otra vez no repite
        profe.post(f'/puntoexacto/{ex.pk}/hojas/', {'grupo_ids': [grupo.pk]})
        self.assertEqual(ex.hojas.count(), 3)
        self.assertEqual(profe.get(f'/puntoexacto/{ex.pk}/hojas/imprimir/').status_code, 200)

        Hoja.objects.filter(pk=h.pk).update(estado='calificada', nota=Decimal('4.5'))
        self.assertContains(profe.get(f'/aula/{aula.pk}/calificaciones/'), 'Quiz 1')
        self.assertContains(profe.get(f'/puntoexacto/{ex.pk}/resultados/'), 'planilla del grupo')

        # Otro docente individual no ve el examen
        otro = User.objects.create_user('otro', password=CLAVE)
        from salon_digital.models import Cuenta
        Cuenta.objects.create(usuario=otro, slug='otro', nombre_publico='Otro', estado='verificada')
        c2 = Client(HTTP_HOST='localhost')
        c2.force_login(otro)
        self.assertEqual(c2.get(f'/puntoexacto/{ex.pk}/resultados/').status_code, 404)

        # Luis ve su nota en el aula
        est = Client(HTTP_HOST='localhost')
        est.post(f'/aula/{aula.pk}/entrar/', {'codigo': grupo.codigo, 'integrante': luis.pk})
        r = est.get(f'/aula/{aula.pk}/examenes/')
        self.assertContains(r, 'Quiz 1')
        self.assertEqual(r.context['examenes'][0].mia.nota, Decimal('4.5'))
