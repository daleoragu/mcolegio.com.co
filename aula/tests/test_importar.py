"""Importar notas (Excel, CSV, tabla pegada) y rejillas de rúbrica."""
import io
from decimal import Decimal

import openpyxl
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client

from aula import importar, logica
from aula.models import Aula, ColumnaManual, Entrega, NotaManual, Tarea
from notas.models import NotaDetallada
from notas.tests.base import CLAVE, ColegioDePrueba


def excel(filas, nombre='notas.xlsx', hojas=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Resumen'
    for f in filas:
        ws.append(f)
    for titulo, otras in (hojas or []):
        w2 = wb.create_sheet(titulo)
        for f in otras:
            w2.append(f)
    salida = io.BytesIO()
    wb.save(salida)
    return SimpleUploadedFile(nombre, salida.getvalue())


def enviar_revision(cliente, url, r, extra=None):
    """Manda el formulario de revisión tal como viene (todas las columnas y las parejas sugeridas)."""
    datos = {'accion': 'importar', 'hoja': r.context['hoja']}
    for k, col in enumerate(r.context['analisis']['columnas']):
        datos[f'usar-{k}'] = 'on'
        datos[f'titulo-{k}'] = col['titulo']
        datos[f'max-{k}'] = str(col['maximo_sugerido'] or '')
    for i, reg in enumerate(r.context['analisis']['registros']):
        datos[f'quien-{i}'] = reg['clave']
    datos.update(extra or {})
    return cliente.post(url + '?paso=revisar', datos)


class Utilidades(ColegioDePrueba):

    def test_numeros_y_nombres(self):
        self.assertEqual(importar.numero('4,5'), Decimal('4.5'))
        self.assertEqual(importar.numero('80%'), Decimal('80'))
        self.assertEqual(importar.numero('8/10'), (Decimal('8'), Decimal('10')))
        self.assertIsNone(importar.numero('ausente'))
        self.assertEqual(importar.a_la_escala(Decimal(8), Decimal(10), Decimal(1), Decimal(5)), Decimal('4.2'))
        self.assertEqual(importar.a_la_escala((Decimal(7), Decimal(10)), None, Decimal(1), Decimal(5)), Decimal('3.8'))
        self.assertEqual(importar.a_la_escala(Decimal(9), None, Decimal(1), Decimal(5)), Decimal('5.0'))
        ana, beto, caro = self.estudiantes
        registros = [{'nombre': 'ANA PRUEBA'}, {'nombre': 'Prueba, Béto'}, {'nombre': 'Prueba'}, {'nombre': 'Zoe'},
                     {'codigo': f'e{caro.id}', 'nombre': 'otra cosa'}]
        self.assertEqual(importar.emparejar(registros, self.estudiantes), [ana, beto, None, None, caro])


class ImportarNotas(ColegioDePrueba):

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.a.modulos_activos = list(cls.a.modulos_activos) + ['salon']
        cls.a.save(update_fields=['modulos_activos'])

    def test_planilla_individual_desde_csv_pegado(self):
        from salon_digital.models import Cuenta, Grupo, Integrante
        u = User.objects.create_user('libre', password=CLAVE)
        cuenta = Cuenta.objects.create(usuario=u, slug='lina', nombre_publico='Lina', estado='verificada')
        grupo = Grupo.objects.create(propietario=cuenta, nombre='7-01')
        luis = Integrante.objects.create(grupo=grupo, apellidos='Pérez', nombres='Luis')
        sara = Integrante.objects.create(grupo=grupo, apellidos='Gómez', nombres='Sara')
        profe = Client(HTTP_HOST='localhost')
        profe.force_login(u)
        profe.get(f'/aula/grupo/{grupo.pk}/')
        aula = Aula.objects.get(grupo=grupo)
        url = f'/aula/{aula.pk}/importar/'
        self.assertContains(profe.get(url), 'Suba el archivo')
        texto = 'Marca temporal;Nombre;Quiz;Puntos\n1/10;Luis Perez;4,5;80\n1/10;GOMEZ SARA;3;40\n1/10;Desconocido;5;100'
        r = profe.post(url, {'accion': 'leer', 'texto': texto})
        r = profe.get(r.url)
        analisis = r.context['analisis']
        self.assertEqual([c['titulo'] for c in analisis['columnas']], ['Quiz', 'Puntos'])
        self.assertEqual(analisis['columnas'][1]['maximo_sugerido'], Decimal(100))
        self.assertEqual([x['clave'] for x in analisis['registros']], [f'i{luis.pk}', f'i{sara.pk}', ''])
        self.assertEqual(r.context['sin_pareja'], 1)
        r = enviar_revision(profe, url, r)
        self.assertRedirects(r, f'/aula/{aula.pk}/calificaciones/', fetch_redirect_response=False)
        notas = {(n.columna.titulo, n.integrante_id): n.nota for n in NotaManual.objects.select_related('columna')}
        self.assertEqual(notas[('Quiz', luis.pk)], Decimal('4.5'))
        self.assertEqual(notas[('Puntos', luis.pk)], Decimal('4.2'))       # 80 de 100
        self.assertEqual(notas[('Puntos', sara.pk)], Decimal('2.6'))       # 40 de 100
        self.assertEqual(ColumnaManual.objects.count(), 2)
        # Importar otra vez la misma columna la actualiza, no la repite
        profe.post(url, {'accion': 'leer', 'texto': 'Nombre;Quiz\nLuis Pérez;3'})
        r = profe.get(url + '?paso=revisar')
        enviar_revision(profe, url, r)
        self.assertEqual(ColumnaManual.objects.count(), 2)
        self.assertEqual(NotaManual.objects.get(columna__titulo='Quiz', integrante=luis).nota, Decimal('3.0'))

    def test_planilla_oficial_desde_excel_de_varias_hojas(self):
        ana, beto, caro = self.estudiantes
        profe = self.cliente(self.u_docente)
        profe.get(f'/aula/asignacion/{self.asig.id}/')
        aula = Aula.objects.get(asignacion=self.asig)
        url = f'/aula/{aula.pk}/importar/'
        archivo = excel([['Kahoot! Resumen'], ['Total de jugadores', 3]], hojas=[('Final Scores', [
            ['Rank', 'Player', 'Total Score (points)', 'Correct Answers'],
            [1, 'ana prueba', 4000, 4], [2, 'Beto P.', 2000, 2]])])
        r = profe.post(url, {'accion': 'leer', 'archivo': archivo})
        r = profe.get(url + '?paso=revisar&hoja=1')
        self.assertEqual([c['titulo'] for c in r.context['analisis']['columnas']],
                         ['Total Score (points)', 'Correct Answers'])
        # Sin periodo ni componente: no pasa
        r2 = enviar_revision(profe, url, r, {'usar-1': ''})
        self.assertEqual(NotaDetallada.objects.count(), 0)
        # Beto no se emparejó solo («Beto P.»): el docente lo escoge
        self.assertEqual(r.context['analisis']['registros'][1]['clave'], '')
        enviar_revision(profe, url, r, {'periodo': self.p1.pk, 'componente': 'SABER', 'usar-1': '',
                                        'titulo-0': 'Kahoot fracciones', 'quien-1': f'e{beto.id}'})
        nd = {n.calificacion_promedio.estudiante_id: n.valor_nota
              for n in NotaDetallada.objects.filter(descripcion='Kahoot fracciones')}
        self.assertEqual(nd, {ana.id: Decimal('5.0'), beto.id: Decimal('3.0')})


class ImportarRejilla(ColegioDePrueba):

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.a.modulos_activos = list(cls.a.modulos_activos) + ['salon']
        cls.a.save(update_fields=['modulos_activos'])

    def test_rejilla_descargada_y_rejilla_propia(self):
        ana, beto, caro = self.estudiantes
        profe = self.cliente(self.u_docente)
        profe.get(f'/aula/asignacion/{self.asig.id}/')
        aula = Aula.objects.get(asignacion=self.asig)
        inst = logica.desde_plantilla(self.u_docente, self.a, 'guia')
        t = Tarea.objects.create(aula=aula, titulo='Guía 3', instrumento=inst)
        criterios = list(inst.criterios.all())
        base = f'/aula/{aula.pk}/tareas/{t.pk}'
        self.assertContains(profe.get(f'{base}/'), 'Rejilla en Excel')

        r = profe.get(f'{base}/rejilla.xlsx')
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        ws = wb['Rejilla']
        self.assertEqual(ws['A3'].value, 'Código aula')
        self.assertEqual([ws.cell(row=3, column=3 + i).value for i in range(len(criterios))],
                         [cr.nombre for cr in criterios])
        filas = {ws.cell(row=i, column=1).value: i for i in range(4, 7)}
        fila_ana = filas[f'e{ana.id}']
        for i in range(len(criterios)):
            ws.cell(row=fila_ana, column=3 + i, value='Superior')
        ws.cell(row=fila_ana, column=3, value='básico')
        ws.cell(row=filas[f'e{beto.id}'], column=3, value='Excelente')       # no existe ese nivel
        salida = io.BytesIO()
        wb.save(salida)
        r = profe.post(f'{base}/rejilla/importar/', {'archivo': SimpleUploadedFile('r.xlsx', salida.getvalue())},
                       follow=True)
        self.assertContains(r, 'Rejilla importada: 1 estudiante(s) calificados')
        self.assertContains(r, 'Excelente')
        self.assertEqual(Entrega.objects.get(estudiante=ana).nota, Decimal('4.6'))   # (3,5 + 5 + 5 + 5) / 4

        # Una rejilla propia, por nombres y con niveles en números (4 = el más alto, como en RubiStar)
        propia = [['Estudiante'] + [cr.nombre for cr in criterios], ['Caro Prueba'] + ['4'] * len(criterios)]
        r = profe.post(f'{base}/rejilla/importar/', {'archivo': excel(propia)}, follow=True)
        self.assertEqual(Entrega.objects.get(estudiante=caro).nota, Decimal('5.0'))
        # El estudiante no importa rejillas
        est = self.cliente(ana.user)
        self.assertEqual(est.get(f'{base}/rejilla.xlsx').status_code, 403)
