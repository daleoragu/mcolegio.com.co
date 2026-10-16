"""Qué dice cada certificado, de dónde salen los datos y cómo se cargan años anteriores."""
import base64
import re
import unicodedata
from collections import OrderedDict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO

from django.contrib.auth.models import User
from django.db import transaction
from django.utils import timezone

from notas.models import (Curso, EscalaValoracion, Estudiante, FichaEstudiante, HistorialMatricula,
                          PeriodoAcademico)
from notas.models.perfiles import NOMBRE_GRADO

from .models import CalificacionHistorica, CertificadoEmitido, ConfiguracionCertificados

MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre',
         'octubre', 'noviembre', 'diciembre']


class SinDatos(Exception):
    """No hay con qué armar el certificado pedido."""


# ---------------------------------------------------------------------------
# Años y datos
# ---------------------------------------------------------------------------
def ano_actual(colegio):
    ano = (PeriodoAcademico.objects.filter(colegio=colegio)
           .order_by('-ano_lectivo').values_list('ano_lectivo', flat=True).first())
    return ano or timezone.localdate().year


def anos_disponibles(estudiante):
    """Años de los que se puede sacar certificado de calificaciones, del más reciente al más viejo."""
    colegio = estudiante.colegio
    anos = set(CalificacionHistorica.objects.filter(estudiante=estudiante).values_list('ano_lectivo', flat=True))
    anos |= set(HistorialMatricula.objects.filter(estudiante=estudiante, curso__isnull=False)
                .values_list('ano_lectivo', flat=True))
    if estudiante.is_active and estudiante.curso_id:
        anos.add(ano_actual(colegio))
    return sorted(anos, reverse=True)


def _texto_nota(valor):
    if valor is None:
        return ''
    return f'{Decimal(valor).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)}'.replace('.', ',')


def _desempeno(colegio, nota):
    if nota is None:
        return ''
    escala = EscalaValoracion.objects.filter(colegio=colegio, valor_minimo__lte=nota, valor_maximo__gte=nota).first()
    return escala.nombre_desempeno.upper() if escala else ''


def escala_del_colegio(colegio):
    return [{'nombre': e.nombre_desempeno.upper(), 'min': _texto_nota(e.valor_minimo), 'max': _texto_nota(e.valor_maximo)}
            for e in EscalaValoracion.objects.filter(colegio=colegio).order_by('-valor_maximo')]


def _nombre_grado(curso, nombre_guardado=''):
    if curso is not None:
        if curso.grado is not None:
            return f'{NOMBRE_GRADO.get(curso.grado, "")} ({curso.nombre})'.strip()
        return curso.nombre
    return nombre_guardado


def _datos_estudiante(estudiante):
    ficha = FichaEstudiante.objects.filter(estudiante=estudiante).first()
    return {
        'nombre': (estudiante.user.get_full_name() or estudiante.user.username).upper(),
        'tipo_documento': ficha.get_tipo_documento_display() if ficha else 'documento',
        'documento': (ficha.numero_documento or '') if ficha else '',
    }


RESULTADO_TEXTO = {
    'PROMOVIDO': 'APROBÓ el grado y fue PROMOVIDO(A)',
    'NO_PROMOVIDO': 'NO FUE PROMOVIDO(A)',
    'GRADUADO': 'APROBÓ el grado y se GRADUÓ',
    'RETIRADO': 'SE RETIRÓ antes de terminar el año',
}


def contenido_calificaciones(estudiante, ano):
    """Lo que dice el certificado de calificaciones de un año.

    Primero mira si hay notas cargadas desde Excel para ese año (años antes de
    usar la plataforma); si no, las calcula como el boletín final.
    """
    colegio = estudiante.colegio
    conf = ConfiguracionCertificados.de(colegio)
    hist = HistorialMatricula.objects.filter(estudiante=estudiante, ano_lectivo=ano).select_related('curso').first()
    filas = list(CalificacionHistorica.objects.filter(estudiante=estudiante, ano_lectivo=ano))
    actual = ano_actual(colegio)
    parcial = False

    if filas:
        grado = filas[0].grado
        areas = OrderedDict()
        for f in filas:
            a = areas.setdefault(f.area.strip().upper(), {'nombre': f.area.strip().upper(), 'nota': None,
                                                          'desempeno': '', 'ih': None, 'asignaturas': []})
            if f.asignatura.strip():
                a['asignaturas'].append({'nombre': f.asignatura.strip(), 'ih': f.intensidad, 'nota': f.nota,
                                         'desempeno': f.desempeno.upper() or _desempeno(colegio, f.nota)})
            else:
                a['nota'], a['ih'] = f.nota, f.intensidad
                a['desempeno'] = f.desempeno.upper() or _desempeno(colegio, f.nota)
        for a in areas.values():
            if a['nota'] is None and a['asignaturas']:
                # Sin nota del área en el Excel: promedio de sus asignaturas
                # ponderado por la intensidad horaria (1 si no la trae).
                con_nota = [s for s in a['asignaturas'] if s['nota'] is not None]
                if con_nota:
                    pesos = [Decimal(s['ih'] or 1) for s in con_nota]
                    a['nota'] = sum(Decimal(s['nota']) * p for s, p in zip(con_nota, pesos)) / sum(pesos)
                    a['desempeno'] = _desempeno(colegio, a['nota'].quantize(Decimal('0.1'), rounding=ROUND_HALF_UP))
            if a['ih'] is None and a['asignaturas'] and any(s['ih'] for s in a['asignaturas']):
                a['ih'] = sum(s['ih'] or 0 for s in a['asignaturas'])
        areas = list(areas.values())
        resultado = RESULTADO_TEXTO.get(hist.resultado, '') if hist else ''
    else:
        from notas.boletin.logic import get_datos_boletin_final
        curso = hist.curso if hist and hist.curso_id else None
        if curso is None and ano == actual and estudiante.is_active:
            curso = estudiante.curso
        if curso is None:
            raise SinDatos(f'No hay notas de {ano} para este estudiante.')
        datos, _ = get_datos_boletin_final(colegio, curso, ano, estudiante)
        if not datos:
            raise SinDatos(f'No hay notas de {ano} para este estudiante.')
        d = datos[0]
        grado = _nombre_grado(curso)
        areas = []
        for ar in d['areas']:
            if ar.get('es_informativa'):
                continue
            asignaturas = [{'nombre': m['nombre'], 'ih': m['ih'],
                            'nota': m['definitiva_recuperada'] if m['definitiva_recuperada'] is not None
                            else m['definitiva_original'],
                            'desempeno': m['valoracion']}
                           for m in ar['materias'] if m['promedia']]
            ih = sum(s['ih'] or 0 for s in asignaturas) or None
            nota_area = ar['nota_final_area']
            if nota_area is None and len(asignaturas) == 1:
                nota_area = asignaturas[0]['nota']
            areas.append({'nombre': ar['nombre'].upper(), 'nota': nota_area,
                          'desempeno': ar['desempeno_area'] or _desempeno(colegio, nota_area),
                          'ih': ih, 'asignaturas': asignaturas})
        if hist:
            resultado = RESULTADO_TEXTO.get(hist.resultado, '')
        elif ano >= actual:
            parcial = True
            resultado = ''
        else:
            resultado = ('APROBÓ el grado y fue PROMOVIDO(A)' if d['estado_promocion'] == 'PROMOVIDO'
                         else 'NO FUE PROMOVIDO(A)')

    # Listo para el papel y para guardar en JSON: notas como texto.
    for a in areas:
        a['nota'] = _texto_nota(a['nota'])
        if len(a['asignaturas']) == 1 and a['asignaturas'][0]['nombre'].strip().upper() == a['nombre']:
            a['asignaturas'] = []          # área de una sola asignatura con el mismo nombre
        for s in a['asignaturas']:
            s['nota'] = _texto_nota(s['nota'])
    return {
        'tipo': 'CALIFICACIONES', 'ano': ano, 'grado': grado, 'areas': areas, 'resultado': resultado,
        'parcial': parcial, 'escala': escala_del_colegio(colegio),
        'mostrar_asignaturas': conf.mostrar_asignaturas, 'mostrar_intensidad': conf.mostrar_intensidad,
        **_datos_estudiante(estudiante),
    }


def contenido_estudio(estudiante):
    if not estudiante.is_active or not estudiante.curso_id:
        raise SinDatos('El certificado de estudio es para estudiantes matriculados. '
                       'Para un exalumno, saque el certificado de calificaciones del año que necesite.')
    return {'tipo': 'ESTUDIO', 'ano': ano_actual(estudiante.colegio), 'grado': _nombre_grado(estudiante.curso),
            **_datos_estudiante(estudiante)}


def _datos_colegio(colegio):
    conf = ConfiguracionCertificados.de(colegio)
    rector = conf.firma_nombre.strip() or (colegio.admin_general.get_full_name() if colegio.admin_general_id else '')
    return {
        'colegio': colegio.nombre, 'resolucion': colegio.resolucion_aprobacion or '', 'dane': colegio.dane or '',
        'nit': colegio.nit or '', 'ciudad': colegio.ciudad or '', 'departamento': colegio.departamento or '',
        'firma_nombre': rector.upper(), 'firma_cargo': conf.firma_cargo,
        'firma2_nombre': conf.firma2_nombre.strip().upper(), 'firma2_cargo': conf.firma2_cargo,
        'texto_final': conf.texto_final.strip(),
    }


def emitir(estudiante, tipo, ano=None, usuario=None):
    """Arma el certificado, lo guarda con su código y lo devuelve."""
    if tipo == 'ESTUDIO':
        contenido = contenido_estudio(estudiante)
        ano = contenido['ano']
    elif tipo == 'CALIFICACIONES':
        if ano is None:
            raise SinDatos('Falta el año.')
        contenido = contenido_calificaciones(estudiante, int(ano))
    else:
        raise SinDatos('Tipo de certificado desconocido.')
    hoy = timezone.localdate()
    contenido.update(_datos_colegio(estudiante.colegio))
    contenido['fecha'] = f'{hoy.day} de {MESES[hoy.month - 1]} de {hoy.year}'
    return CertificadoEmitido.objects.create(
        colegio=estudiante.colegio, estudiante=estudiante, tipo=tipo, ano_lectivo=ano,
        contenido=contenido, emitido_por=usuario if usuario and usuario.is_authenticated else None)


def qr_base64(texto):
    import qrcode
    qr = qrcode.QRCode(version=None, error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=6, border=2)
    qr.add_data(texto)
    qr.make(fit=True)
    buf = BytesIO()
    qr.make_image(fill_color='black', back_color='white').save(buf, format='PNG')
    return base64.b64encode(buf.getvalue()).decode('ascii')


# ---------------------------------------------------------------------------
# Importar años anteriores desde Excel
# ---------------------------------------------------------------------------
COLUMNAS = [
    ('documento', 'Documento', True),
    ('apellidos', 'Apellidos', False),
    ('nombres', 'Nombres', False),
    ('ano', 'Año', True),
    ('grado', 'Grado', True),
    ('area', 'Área', True),
    ('asignatura', 'Asignatura', False),
    ('ih', 'IH', False),
    ('nota', 'Nota', False),
    ('desempeno', 'Desempeño', False),
    ('resultado', 'Resultado del año', False),
]

RESULTADOS = {
    'PROMOVIDO': 'PROMOVIDO', 'APROBO': 'PROMOVIDO', 'APROBADO': 'PROMOVIDO', 'SI': 'PROMOVIDO',
    'NO PROMOVIDO': 'NO_PROMOVIDO', 'REPROBO': 'NO_PROMOVIDO', 'REPROBADO': 'NO_PROMOVIDO', 'NO': 'NO_PROMOVIDO',
    'GRADUADO': 'GRADUADO', 'GRADUO': 'GRADUADO', 'RETIRADO': 'RETIRADO', 'RETIRO': 'RETIRADO',
}


def _llano(texto):
    t = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', t).strip().upper()


def plantilla_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Notas'
    ws.append([c[1] for c in COLUMNAS])
    for celda in ws[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='193661')
    ejemplos = [
        ['1110123456', 'Pérez Gómez', 'Ana María', 2019, 'Sexto', 'Matemáticas', '', 5, 4.2, '', 'Promovido'],
        ['1110123456', 'Pérez Gómez', 'Ana María', 2019, 'Sexto', 'Humanidades', 'Lengua castellana', 4, 3.8, '', ''],
        ['1110123456', 'Pérez Gómez', 'Ana María', 2019, 'Sexto', 'Humanidades', 'Inglés', 3, 4.0, '', ''],
    ]
    for fila in ejemplos:
        ws.append(fila)
    for col, ancho in zip('ABCDEFGHIJK', [14, 22, 20, 7, 12, 26, 24, 5, 7, 12, 16]):
        ws.column_dimensions[col].width = ancho
    ws2 = wb.create_sheet('Instrucciones')
    for linea in [
        'Una fila por cada área o asignatura de cada estudiante en cada año.',
        'Documento, Año, Grado y Área son obligatorios. Apellidos y Nombres solo hacen falta si el '
        'estudiante no está en la plataforma (exalumnos): con ellos se crea su registro.',
        'Asignatura: déjela vacía en la fila que lleva la nota del ÁREA. Si el área no trae fila propia, '
        'su nota se calcula con las asignaturas, según la intensidad horaria (IH).',
        'Nota: con punto o coma (4.2 o 4,2). Desempeño: si lo deja vacío, sale según la escala del colegio.',
        'Resultado del año (en cualquier fila del estudiante): Promovido, No promovido, Graduado o Retirado.',
        'Si vuelve a subir un estudiante y un año que ya estaban, se reemplazan esas notas.',
        'Borre las filas de ejemplo antes de subir el archivo.',
    ]:
        ws2.append([linea])
    ws2.column_dimensions['A'].width = 120
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _leer_filas(archivo):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(archivo, read_only=True, data_only=True)
    except Exception:
        raise SinDatos('No se pudo abrir el archivo. Debe ser un Excel (.xlsx).')
    ws = wb['Notas'] if 'Notas' in wb.sheetnames else wb.worksheets[0]
    filas = ws.iter_rows(values_only=True)
    try:
        cabecera = next(filas)
    except StopIteration:
        raise SinDatos('El archivo está vacío.')
    nombres = {_llano(c[1]): c[0] for c in COLUMNAS}
    nombres.update({'ANO LECTIVO': 'ano', 'ANIO': 'ano', 'INTENSIDAD': 'ih', 'INTENSIDAD HORARIA': 'ih',
                    'RESULTADO': 'resultado', 'MATERIA': 'asignatura', 'VALORACION': 'nota',
                    'DOCUMENTO DE IDENTIDAD': 'documento', 'IDENTIFICACION': 'documento'})
    indice = {}
    for i, c in enumerate(cabecera):
        clave = nombres.get(_llano(c))
        if clave and clave not in indice:
            indice[clave] = i
    faltan = [c[1] for c in COLUMNAS if c[2] and c[0] not in indice]
    if faltan:
        raise SinDatos('Al archivo le faltan las columnas: ' + ', '.join(faltan) + '. Use la plantilla.')
    for n, fila in enumerate(filas, start=2):
        if not fila or all(v in (None, '') for v in fila):
            continue
        yield n, {k: (fila[i] if i < len(fila) else None) for k, i in indice.items()}


def _numero(valor):
    if valor in (None, ''):
        return None
    try:
        return Decimal(str(valor).replace(',', '.').strip())
    except InvalidOperation:
        raise ValueError


def importar_excel(archivo, colegio, usuario):
    """Revisa todo el archivo; si no hay errores, guarda. Devuelve (resumen, errores)."""
    actual = ano_actual(colegio)
    escala = EscalaValoracion.objects.filter(colegio=colegio)
    nota_min = min((e.valor_minimo for e in escala), default=Decimal('0'))
    nota_max = max((e.valor_maximo for e in escala), default=Decimal('10'))
    errores, grupos = [], OrderedDict()
    for n, f in _leer_filas(archivo):
        crudo = f.get('documento')
        if isinstance(crudo, float) and crudo.is_integer():
            crudo = int(crudo)              # Excel guarda 1110123456 como 1110123456.0
        doc = re.sub(r'[\s.]', '', str(crudo or '')).strip()
        try:
            ano = int(str(f.get('ano') or '').strip().split('.')[0])
        except ValueError:
            ano = None
        area = str(f.get('area') or '').strip()
        grado = str(f.get('grado') or '').strip()
        if grado.endswith('.0'):
            grado = grado[:-2]
        if not doc:
            errores.append(f'Fila {n}: falta el documento.')
            continue
        if ano is None or not (1950 <= ano <= actual):
            errores.append(f'Fila {n}: el año «{f.get("ano")}» no es válido.')
            continue
        if not area or not grado:
            errores.append(f'Fila {n}: faltan el grado o el área.')
            continue
        try:
            nota = _numero(f.get('nota'))
            ih = _numero(f.get('ih'))
        except ValueError:
            errores.append(f'Fila {n}: la nota o la IH no son números.')
            continue
        if nota is not None and not (nota_min <= nota <= nota_max):
            errores.append(f'Fila {n}: la nota {nota} está fuera de la escala ({nota_min} a {nota_max}).')
            continue
        resultado = ''
        if f.get('resultado'):
            resultado = RESULTADOS.get(_llano(f.get('resultado')), '')
            if not resultado:
                errores.append(f'Fila {n}: no se entiende el resultado «{f.get("resultado")}» '
                               f'(use Promovido, No promovido, Graduado o Retirado).')
                continue
        g = grupos.setdefault((doc, ano), {'apellidos': '', 'nombres': '', 'grado': grado, 'resultado': '',
                                           'filas': []})
        g['apellidos'] = g['apellidos'] or str(f.get('apellidos') or '').strip()
        g['nombres'] = g['nombres'] or str(f.get('nombres') or '').strip()
        g['resultado'] = g['resultado'] or resultado
        g['filas'].append({'area': area, 'asignatura': str(f.get('asignatura') or '').strip(), 'nota': nota,
                           'ih': int(ih) if ih is not None else None,
                           'desempeno': str(f.get('desempeno') or '').strip()})

    # Estudiantes: por documento; los que no están se crean como exalumnos.
    fichas = {re.sub(r'[\s.]', '', fi.numero_documento): fi.estudiante
              for fi in FichaEstudiante.objects.filter(estudiante__colegio=colegio, numero_documento__isnull=False)
              .select_related('estudiante__user')}
    por_crear = {}
    for (doc, ano), g in grupos.items():
        if doc not in fichas and doc not in por_crear:
            if not (g['apellidos'] and g['nombres']):
                errores.append(f'El documento {doc} no está en la plataforma: escriba sus apellidos y nombres '
                               f'para crearlo como exalumno.')
                continue
            if FichaEstudiante.objects.filter(numero_documento=doc).exclude(estudiante__colegio=colegio).exists():
                errores.append(f'El documento {doc} ya está registrado en otro colegio de la plataforma.')
                continue
            por_crear[doc] = (g['apellidos'], g['nombres'])
    if errores:
        return None, errores
    if not grupos:
        return None, ['El archivo no trae filas con datos.']

    with transaction.atomic():
        creados = 0
        for doc, (apellidos, nombres) in por_crear.items():
            base = f'ex{doc}'[:140]
            usuario_nuevo, k = base, 1
            while User.objects.filter(username=usuario_nuevo).exists():
                k += 1
                usuario_nuevo = f'{base}-{k}'
            u = User.objects.create(username=usuario_nuevo, first_name=nombres[:150], last_name=apellidos[:150],
                                    is_active=False)
            u.set_unusable_password()
            u.save(update_fields=['password'])
            est = Estudiante.objects.create(colegio=colegio, user=u, curso=None, is_active=False)
            FichaEstudiante.objects.create(estudiante=est, numero_documento=doc, tipo_documento='OT')
            fichas[doc] = est
            creados += 1
        filas_guardadas = 0
        for (doc, ano), g in grupos.items():
            est = fichas[doc]
            CalificacionHistorica.objects.filter(estudiante=est, ano_lectivo=ano).delete()
            CalificacionHistorica.objects.bulk_create([
                CalificacionHistorica(colegio=colegio, estudiante=est, ano_lectivo=ano, grado=g['grado'],
                                      area=r['area'], asignatura=r['asignatura'], intensidad=r['ih'],
                                      nota=r['nota'], desempeno=r['desempeno'], orden=i, cargado_por=usuario)
                for i, r in enumerate(g['filas'])])
            filas_guardadas += len(g['filas'])
            # El resultado del año va al historial, solo si ese año aún no lo
            # tiene (lo que registró la promoción de la plataforma no se toca).
            if g['resultado'] and ano < actual and not HistorialMatricula.objects.filter(
                    estudiante=est, ano_lectivo=ano).exists():
                HistorialMatricula.objects.create(colegio=colegio, estudiante=est, ano_lectivo=ano,
                                                  curso_nombre=g['grado'][:100], resultado=g['resultado'],
                                                  estaba_activo=est.is_active, registrado_por=usuario)
    estudiantes = len({doc for doc, _ in grupos})
    return {'estudiantes': estudiantes, 'anos': len(grupos), 'filas': filas_guardadas,
            'exalumnos_creados': creados}, []
